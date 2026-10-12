# -*- coding: utf-8 -*-
"""Estilos e a linha do tempo do modelo mensal.

Convencao de colunas em TODA aba mensal:

    A  rotulo
    B  valor inicial          (azul: voce digita)
    C  evolucao por mes       (azul: voce digita)
    D  mes de inicio (1..60)  (azul: voce digita)
    E..BL  os 60 meses
    BM TOTAL
    BN nota

Nao existe coluna de ano em lugar nenhum. Premissa anual foi o erro da versao
anterior: obrigava a pensar o negocio em blocos de doze meses quando a decisao
e mensal.
"""
import datetime as dt
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as gcl

FONT = "Arial"
BRL   = '"R$ "#,##0;("R$ "#,##0);"–"'
BRL2  = '"R$ "#,##0.00;("R$ "#,##0.00);"–"'
BRL4  = '"R$ "#,##0.0000;("R$ "#,##0.0000);"–"'
INT   = '#,##0;(#,##0);"–"'
PCT   = '0.0%;(0.0%);"–"'
PCT2  = '0.00%;(0.00%);"–"'
MULT  = '0.00"x";(0.00"x");"–"'
MES   = 'mmm/yy'
TXT   = '@'
NUM1  = '0.0;(0.0);"–"'

AZUL  = "0000FF"      # voce digita
PRETO = "000000"      # formula na propria aba
VERDE = "008000"      # vem de outra aba
CINZA = "595959"
BRANCO = "FFFFFF"
VERM  = "C00000"
LARANJA = "833C00"

F_TIT    = PatternFill("solid", fgColor="1F3864")
F_SEC    = PatternFill("solid", fgColor="2E5C8A")
F_SUB    = PatternFill("solid", fgColor="D9E2F3")
F_TOT    = PatternFill("solid", fgColor="BDD7EE")
F_IN     = PatternFill("solid", fgColor="FFF2CC")   # celula de entrada
F_IN_KEY = PatternFill("solid", fgColor="FFFF00")   # entrada que move muito
F_HDR    = PatternFill("solid", fgColor="404040")
F_OK     = PatternFill("solid", fgColor="E2EFDA")
F_ALERTA = PatternFill("solid", fgColor="FCE4D6")
F_DESTAQ = PatternFill("solid", fgColor="1F3864")

_t = Side(style="thin", color="BFBFBF")
BORDA = Border(left=_t, right=_t, top=_t, bottom=_t)

# ---- linha do tempo ------------------------------------------------------
INICIO  = dt.date(2026, 1, 1)
N       = 60
C_INI   = 5                     # coluna E = mes 1
C_FIM   = C_INI + N - 1         # 64 = BL
C_TOT   = C_FIM + 1             # 65 = BM
C_NOTA  = C_TOT + 1             # 66 = BN
L_INI   = gcl(C_INI)
L_FIM   = gcl(C_FIM)
L_TOT   = gcl(C_TOT)
L_NOTA  = gcl(C_NOTA)

R_DATA, R_IDX, R_ANO = 4, 5, 6


def col(i):
    """Letra da coluna do mes i (1..60)."""
    return gcl(C_INI + i - 1)


def ant(i):
    """Letra da coluna do mes anterior."""
    return gcl(C_INI + i - 2)


def data_mes(i):
    a = 2026 + (i - 1) // 12
    m = (i - 1) % 12 + 1
    return dt.datetime(a, m, 1)


def f(size=10, bold=False, color=PRETO, italic=False):
    return Font(name=FONT, size=size, bold=bold, color=color, italic=italic)


def titulo(ws, texto, sub=None):
    ws["A1"] = texto
    ws["A1"].font = f(14, True, BRANCO)
    ws["A1"].alignment = Alignment(vertical="center", indent=1)
    for c in range(1, C_NOTA + 1):
        ws.cell(1, c).fill = F_TIT
    ws.row_dimensions[1].height = 26
    if sub:
        ws["A2"] = sub
        ws["A2"].font = f(9, False, CINZA, italic=True)
        ws["A2"].alignment = Alignment(vertical="center", indent=1)
        ws.row_dimensions[2].height = 15


def secao(ws, r, texto, largura=None):
    ws.cell(r, 1, texto).font = f(10, True, BRANCO)
    for c in range(1, (largura or C_NOTA) + 1):
        ws.cell(r, c).fill = F_SEC
    ws.row_dimensions[r].height = 18


def nota(ws, r, texto, col_ini=C_NOTA, ate=None):
    c = ws.cell(r, col_ini, texto)
    c.font = f(8, False, CINZA, italic=True)
    c.alignment = Alignment(wrap_text=True, vertical="center", indent=1)
    if ate:
        ws.merge_cells(start_row=r, start_column=col_ini, end_row=r, end_column=ate)


def tempo(ws, congelar="E8"):
    """Datas, numero do mes e ano nas linhas 4, 5 e 6."""
    ws.cell(R_DATA, 1, "Linha do tempo").font = f(9, True, CINZA)
    ws.cell(R_IDX, 1, "Mês nº").font = f(8, False, CINZA)
    ws.cell(R_ANO, 1, "Ano").font = f(8, False, CINZA)
    for i in range(1, N + 1):
        L = col(i)
        c = ws[f"{L}{R_DATA}"]
        c.value = data_mes(i)
        c.number_format = MES
        c.font = f(9, True, BRANCO)
        c.fill = F_HDR
        c.alignment = Alignment(horizontal="center")
        ci = ws[f"{L}{R_IDX}"]
        ci.value = i
        ci.number_format = INT
        ci.font = f(8, False, CINZA)
        ci.alignment = Alignment(horizontal="center")
        ca = ws[f"{L}{R_ANO}"]
        ca.value = f"=YEAR({L}{R_DATA})"
        ca.number_format = "0"
        ca.font = f(8, False, "BFBFBF")
        ca.alignment = Alignment(horizontal="center")
    for L, t in ((L_TOT, "TOTAL"), (L_NOTA, "O que é / como mexer")):
        c = ws[f"{L}{R_DATA}"]
        c.value = t
        c.font = f(9, True, BRANCO)
        c.fill = F_TIT
        c.alignment = Alignment(horizontal="center")
    for j, t in enumerate(("valor inicial", "evolução/mês", "mês de início")):
        c = ws.cell(R_DATA, 2 + j, t)
        c.font = f(8, True, BRANCO)
        c.fill = F_HDR
        c.alignment = Alignment(horizontal="center", wrap_text=True)
    ws.row_dimensions[R_DATA].height = 24
    ws.freeze_panes = congelar


def larguras(ws, a=50):
    ws.column_dimensions["A"].width = a
    ws.column_dimensions["B"].width = 13
    ws.column_dimensions["C"].width = 12
    ws.column_dimensions["D"].width = 12
    for i in range(1, N + 1):
        ws.column_dimensions[col(i)].width = 11
    ws.column_dimensions[L_TOT].width = 15
    ws.column_dimensions[L_NOTA].width = 62


def entrada(cel, valor, fmt, chave=False):
    cel.value = valor
    cel.number_format = fmt
    cel.font = f(10, False, AZUL)
    cel.fill = F_IN_KEY if chave else F_IN
    cel.border = BORDA
    cel.alignment = Alignment(horizontal="center")
    return cel
