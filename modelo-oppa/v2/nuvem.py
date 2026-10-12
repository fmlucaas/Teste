# -*- coding: utf-8 -*-
"""Aba 'Nuvem e Escala' — o custo de nuvem calculado de tras pra frente.

Fornecedor de nuvem nao cobra por usuario: cobra uma conta mensal. Entao a
entrada aqui e a CONTA, nao o custo unitario — e o custo por usuario sai
dividido. E a conta ao contrario.

O custo por usuario cai com a escala porque a conta nao cresce na mesma
proporcao: dobrar a base nao dobra o banco de dados nem a transferencia. A
tabela abaixo ancora a conta em alguns tamanhos de base e o modelo INTERPOLA
entre eles, mes a mes, pela base daquele mes.

Por que interpolar em vez de usar faixas em degrau: com degrau, sair de 999
para 1.001 usuarios derrubaria a conta pela metade — o custo total CAIRIA ao
ganhar um usuario. Interpolando, a conta so sobe e o custo por usuario so cai.
"""
from openpyxl.styles import Alignment
from base import *

ABA = "Nuvem e Escala"

# usuarios na base, conta mensal de nuvem, o que compoe
ANCORAS = [
    (0,         300,     "Ambiente mínimo: banco de dados pequeno, um servidor, backup e "
                         "monitoramento. Existe mesmo com zero usuário."),
    (100,       640,     "Piloto. Daqui sai o R$ 6,40 por usuário que apareceu na conta do sócio."),
    (1_000,     3_200,   "Primeiro degrau de verdade: banco maior, fila de notificações, CDN."),
    (10_000,    16_000,  "Começa a valer reserva de capacidade e desconto de compromisso."),
    (100_000,   90_000,  "Escala em que se negocia contrato anual com o provedor."),
    (1_000_000, 550_000, "Escala de operação nacional."),
]


def build(wb):
    ws = wb.create_sheet(ABA)
    larguras(ws, a=46)
    titulo(ws, "NUVEM E ESCALA — A CONTA DE NUVEM, DE TRÁS PRA FRENTE",
           "Você digita a CONTA mensal esperada em cada tamanho de base. O custo por usuário é "
           "calculado, e cai conforme a base cresce. Mês a mês, o modelo interpola pela base "
           "daquele mês.")
    tempo(ws)

    N_ANC = len(ANCORAS)
    r = 8
    secao(ws, r, "0. QUEM CONSOME NUVEM")
    r += 1
    c = ws.cell(r, 1, "Peso de um usuário do plano GRATUITO no consumo de nuvem")
    c.font = f(10, True, PRETO)
    entrada(ws.cell(r, 2), 0.10, PCT, chave=True)
    cd = ws.cell(r, 4, "Um assinante pagante tem monitoramento contínuo: vale 100%. Um usuário "
                       "do plano gratuito abre o app de vez em quando, recebe notificação e "
                       "guarda pouco dado — aqui vale 10% de um pagante. É a premissa que mais "
                       "muda o resultado do modelo, porque a base gratuita é 30 vezes maior que "
                       "a pagante. Com 100% você está dizendo que o gratuito custa igual ao "
                       "pagante; com 0%, que não custa nada.")
    cd.font = f(9, False, VERM, italic=True)
    cd.alignment = Alignment(wrap_text=True, vertical="center", indent=1)
    ws.merge_cells(start_row=r, start_column=4, end_row=r, end_column=14)
    ws.row_dimensions[r].height = 44
    PESO = r
    r += 1
    c = ws.cell(r, 1, "Mês em que a nuvem começa a custar")
    c.font = f(10, True, PRETO)
    entrada(ws.cell(r, 2), 9, INT, chave=False)
    cd = ws.cell(r, 4, "1 = jan/2026 · 9 = set/2026 · 60 = dez/2030. Antes deste mês a conta é "
                       "zero: não existe ambiente para pagar. Padrão 9, quando começa o "
                       "desenvolvimento do MVP.")
    cd.font = f(9, False, CINZA, italic=True)
    cd.alignment = Alignment(wrap_text=True, vertical="center", indent=1)
    ws.merge_cells(start_row=r, start_column=4, end_row=r, end_column=14)
    ws.row_dimensions[r].height = 26
    MES_INI = r
    r += 2
    secao(ws, r, "1. ÂNCORAS — a conta de nuvem em cada tamanho de base  (você digita as colunas B e C)")
    r += 1
    hdr = r
    for j, t in enumerate(["Usuários na base", "Conta de nuvem no mês (R$)",
                           "R$ por usuário (calculado)", "O que entra nessa conta"]):
        c = ws.cell(r, 1 + j, t)
        c.font = f(9, True, BRANCO)
        c.fill = F_HDR
        c.alignment = Alignment(horizontal="center" if j < 3 else "left", wrap_text=True, indent=1)
    ws.row_dimensions[r].height = 26
    r += 1
    A_INI = r
    for usuarios, conta, desc in ANCORAS:
        entrada(ws.cell(r, 1), usuarios, INT, chave=False)
        entrada(ws.cell(r, 2), conta, BRL, chave=True)
        c = ws.cell(r, 3, f'=IFERROR($B{r}/$A{r},"–")')
        c.number_format = BRL4
        c.font = f(10, True, PRETO)
        c.alignment = Alignment(horizontal="center")
        c.border = BORDA
        cd = ws.cell(r, 4, desc)
        cd.font = f(8, False, CINZA, italic=True)
        cd.alignment = Alignment(wrap_text=True, vertical="center", indent=1)
        ws.merge_cells(start_row=r, start_column=4, end_row=r, end_column=12)
        ws.row_dimensions[r].height = 26
        r += 1
    A_FIM = r - 1
    rng_u = f"$A${A_INI}:$A${A_FIM}"
    rng_c = f"$B${A_INI}:$B${A_FIM}"
    r += 1

    c = ws.cell(r, 1, "Como ler: a terceira coluna é a divisão simples da conta pelo número de "
                      "usuários. É ela que mostra o custo por usuário caindo de R$ 6,40 para "
                      "R$ 0,55 conforme a base vai de 100 para 1 milhão. Se o provedor der uma "
                      "cotação diferente, troque a conta na coluna B e todo o modelo se ajusta.")
    c.font = f(9, False, CINZA, italic=True)
    c.alignment = Alignment(wrap_text=True, vertical="top", indent=1)
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=12)
    ws.row_dimensions[r].height = 30
    r += 2

    # ------------------------------------------------------------------
    secao(ws, r, "2. O QUE ISSO DÁ MÊS A MÊS — interpolado pela base de cada mês")
    r += 1
    C = {}

    def linha(rot, gen, key, fmt=BRL, bold=False, cor=PRETO, obs=None, destaque=False):
        nonlocal r
        cl = ws.cell(r, 1, rot)
        cl.font = f(10, bold, cor)
        cl.alignment = Alignment(indent=0 if bold else 1, vertical="center")
        for i in range(1, N + 1):
            L = col(i)
            cc = ws[f"{L}{r}"]
            cc.value = gen(i, L)
            cc.number_format = fmt
            cc.font = f(9, bold, cor)
            if destaque:
                cc.fill = F_TOT
        ct = ws[f"{L_TOT}{r}"]
        ct.value = f"=SUM({L_INI}{r}:{L_FIM}{r})" if fmt in (BRL, BRL2) else f"={L_FIM}{r}"
        ct.number_format = fmt
        ct.font = f(9, True, PRETO)
        ct.fill = F_TOT
        if obs:
            nota(ws, r, obs)
        C[key] = r
        r += 1
        return C[key]

    # A base vem da aba DRE Mensal, que ainda nao existe nesta altura da
    # construcao. Fica zero aqui e liga() escreve a referencia depois.
    linha("Assinantes pagantes (vem da DRE Mensal)", lambda i, L: "=0", "pag", INT, cor=VERDE)
    linha("Usuários do plano gratuito (vem da DRE Mensal)", lambda i, L: "=0", "free", INT,
          cor=VERDE)
    linha("Usuários equivalentes que consomem nuvem",
          lambda i, L: f"=ROUND({L}{C['pag']}+{L}{C['free']}*$B${PESO},0)", "base", INT,
          bold=True, obs="Pagantes contam 100%; gratuitos contam pelo peso da seção 0. "
                         "É este número que entra na tabela de âncoras.")
    linha("Âncora imediatamente abaixo — usuários",
          lambda i, L: (f"=INDEX({rng_u},MATCH(MIN({L}{C['base']},MAX({rng_u})),{rng_u},1))"),
          "u0", INT, cor=CINZA)
    linha("Âncora imediatamente abaixo — conta",
          lambda i, L: (f"=INDEX({rng_c},MATCH(MIN({L}{C['base']},MAX({rng_u})),{rng_u},1))"),
          "b0", BRL, cor=CINZA)
    linha("Âncora de cima — usuários",
          lambda i, L: (f"=INDEX({rng_u},MIN({A_FIM-A_INI+1},"
                        f"MATCH(MIN({L}{C['base']},MAX({rng_u})),{rng_u},1)+1))"),
          "u1", INT, cor=CINZA)
    linha("Âncora de cima — conta",
          lambda i, L: (f"=INDEX({rng_c},MIN({A_FIM-A_INI+1},"
                        f"MATCH(MIN({L}{C['base']},MAX({rng_u})),{rng_u},1)+1))"),
          "b1", BRL, cor=CINZA)
    linha("CONTA DE NUVEM DO MÊS",
          lambda i, L: (f"=IF({L}${R_IDX}<$B${MES_INI},0,"
                        f"IF({L}{C['base']}<=0,{L}{C['b0']},"
                        f"IF({L}{C['u1']}<={L}{C['u0']},"
                        f"{L}{C['b1']}/MAX(1,{L}{C['u1']})*{L}{C['base']},"
                        f"{L}{C['b0']}+({L}{C['base']}-{L}{C['u0']})"
                        f"*({L}{C['b1']}-{L}{C['b0']})/({L}{C['u1']}-{L}{C['u0']}))))"),
          "conta", BRL, bold=True, destaque=True,
          obs="Interpolação linear entre as duas âncoras. Acima da última âncora, mantém o "
              "custo por usuário da última faixa.")
    linha("Custo de nuvem por usuário equivalente",
          lambda i, L: f"=IFERROR({L}{C['conta']}/{L}{C['base']},0)", "unit", BRL4,
          bold=True, obs="É esta linha que cai com a escala. Compare o mês do lançamento com "
                         "dezembro de 2030.")
    r += 1
    linha("Parcela da conta que vem dos assinantes pagantes",
          lambda i, L: f"=IFERROR({L}{C['conta']}*{L}{C['pag']}/{L}{C['base']},0)", "parte_pag",
          BRL)
    linha("Parcela da conta que vem do plano gratuito",
          lambda i, L: f"=IFERROR({L}{C['conta']}*({L}{C['base']}-{L}{C['pag']})/{L}{C['base']},0)",
          "parte_free", BRL,
          obs="O que o plano gratuito custa em nuvem. Não gera receita nenhuma.")
    linha("Subsídio do plano gratuito, por assinante pagante",
          lambda i, L: f"=IFERROR({L}{C['parte_free']}/{L}{C['pag']},0)", "subsidio", BRL2,
          bold=True,
          obs="Quanto cada assinante paga, sem saber, para sustentar os usuários gratuitos. "
              "Compare com o preço da assinatura: é a conta que decide se o plano gratuito "
              "cabe no modelo.")
    linha("Custo de nuvem por assinante pagante (total da conta ÷ pagantes)",
          lambda i, L: f"=IFERROR({L}{C['conta']}/{L}{C['pag']},0)", "por_pag", BRL2, bold=True,
          obs="É este o número que se compara com o preço da assinatura — não o custo por "
              "usuário equivalente.")
    r += 1

    c = ws.cell(r, 1, "✔ Conferência: a conta nunca pode cair quando a base sobe")
    c.font = f(10, True, PRETO)
    for i in range(1, N + 1):
        L, P = col(i), ant(i) if i > 1 else None
        cc = ws[f"{L}{r}"]
        cc.value = ("=1" if i == 1 else
                    f'=IF(AND({L}{C["base"]}>={P}{C["base"]},{L}{C["conta"]}<{P}{C["conta"]}-0.01),0,1)')
        cc.number_format = INT
        cc.font = f(9, False, CINZA)
    ct = ws[f"{L_TOT}{r}"]
    ct.value = f'=IF(MIN({L_INI}{r}:{L_FIM}{r})=1,"✔ monotônica","⚠ a conta caiu com a base subindo")'
    ct.number_format = TXT
    ct.font = f(9, True, PRETO)
    ct.fill = F_TOT
    nota(ws, r, "Com faixas em degrau esta linha acusaria erro. Com interpolação, não.")
    C["chk"] = r
    r += 2

    c = ws.cell(r, 1, "Mesma mecânica serve para qualquer custo que escale: se o suporte, as "
                      "licenças de software ou a taxa de um parceiro tiverem desconto por volume, "
                      "copie este bloco. A aba DRE Mensal trata suporte como R$ por pagante "
                      "porque, no tamanho da Oppa, suporte escala quase linear.")
    c.font = f(9, False, CINZA, italic=True)
    c.alignment = Alignment(wrap_text=True, vertical="top", indent=1)
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=12)
    ws.row_dimensions[r].height = 30

    ws.sheet_properties.tabColor = "7030A0"
    C["aba"] = ABA
    return ws, C


def liga(ws, C, D):
    """Fecha o circuito: pagantes e gratuitos passam a vir da aba DRE Mensal."""
    for i in range(1, N + 1):
        L = col(i)
        ws[f"{L}{C['pag']}"] = f"='DRE Mensal'!{L}${D['pag']}"
        ws[f"{L}{C['free']}"] = f"='DRE Mensal'!{L}${D['free']}"
