# -*- coding: utf-8 -*-
"""Aba Resumo — comeca com a pergunta que o usuario faz primeiro.

Construida em duas passadas: `controles` escreve o cenario e a tabela de
multiplicadores (que a DRE le), e `painel` escreve o resto (que le a DRE).
"""
from openpyxl.styles import Alignment
from base import *

ABA = "Resumo"
CEN = ["CONSERVADOR", "PROVÁVEL", "AGRESSIVO"]
CORES = ["C00000", "1F3864", "1F7040"]
MULT = [("captação de usuários", 0.60, 1.00, 1.50),
        ("conversão em pagante", 0.60, 1.00, 1.50),
        ("churn",                1.30, 1.00, 0.80)]


def controles(wb):
    ws = wb.create_sheet(ABA)
    for c, w in zip("ABCDEFGH", (3, 52, 19, 19, 19, 19, 3, 66)):
        ws.column_dimensions[c].width = w
    titulo(ws, "OPPA — VIABILIDADE FINANCEIRA")
    ws["A2"] = ("Projeção mensal de jan/2026 a dez/2030. Tudo vem de fórmula — nenhum número "
                "desta aba é digitado à mão.")
    ws["A2"].font = f(9, False, CINZA, italic=True)
    ws["A2"].alignment = Alignment(vertical="center", indent=1)
    C = {}
    r = 4
    secao(ws, r, "CENÁRIO — troque o número e o modelo inteiro responde", 6); r += 1
    ws.cell(r, 2, "Cenário ativo   1 = Conservador · 2 = Provável · 3 = Agressivo").font = \
        f(11, True, PRETO)
    ws.cell(r, 2).alignment = Alignment(indent=1, vertical="center")
    entrada(ws.cell(r, 3), 2, "0", chave=True)
    ws.cell(r, 3).font = f(14, True, AZUL)
    cn = ws.cell(r, 4, f'=CHOOSE($C${r},"{CEN[0]}","{CEN[1]}","{CEN[2]}")')
    cn.font = f(12, True, PRETO)
    cn.alignment = Alignment(horizontal="left", indent=1)
    ws.row_dimensions[r].height = 24
    C["cenario"] = r
    r += 2

    ws.cell(r, 3, "Multiplicador aplicado ao driver").font = f(9, True, BRANCO)
    ws.cell(r, 3).fill = F_HDR
    ws.cell(r, 3).alignment = Alignment(indent=1)
    for k in range(3):
        c = ws.cell(r, 4 + k, CEN[k])
        c.font = f(9, True, BRANCO)
        c.fill = PatternFill("solid", fgColor=CORES[k])
        c.alignment = Alignment(horizontal="center")
    ws.cell(r, 2, "Como cada cenário mexe nos drivers").font = f(10, True, PRETO)
    ws.cell(r, 2).alignment = Alignment(indent=1)
    r += 1
    C["mult_ini"] = r
    for nome, a, b, c3 in MULT:
        ws.cell(r, 3, nome).font = f(10, False, PRETO)
        ws.cell(r, 3).alignment = Alignment(indent=1)
        for k, v in enumerate((a, b, c3)):
            entrada(ws.cell(r, 4 + k), v, '0.00"x"')
        r += 1
    ws.cell(r, 2, "São multiplicadores, não valores: 0,60x significa captar 60% do que está na "
                  "aba DRE Mensal. Edite à vontade — é só isto que separa os três cenários.")
    ws.cell(r, 2).font = f(9, False, CINZA, italic=True)
    ws.cell(r, 2).alignment = Alignment(wrap_text=True, vertical="top", indent=1)
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6)
    ws.row_dimensions[r].height = 28
    r += 2
    C["prox"] = r
    return ws, C


def painel(ws, C, D, A, NU):
    r = C["prox"]
    R = {}

    def rng(key):
        return f"'DRE Mensal'!${L_INI}${D[key]}:${L_FIM}${D[key]}"

    def fim(key):
        return f"'DRE Mensal'!${L_FIM}${D[key]}"

    def mes_de(expr):
        """Converte um numero de mes (1..60) em texto mmm/aaaa."""
        return (f'CHOOSE(MOD({expr}-1,12)+1,"jan","fev","mar","abr","mai","jun","jul","ago",'
                f'"set","out","nov","dez")&"/"&(2026+INT(({expr}-1)/12))')

    def bloco(txt, cor="2E5C8A"):
        nonlocal r
        ws.cell(r, 2, txt).font = f(11, True, BRANCO)
        for c in range(2, 7):
            ws.cell(r, c).fill = PatternFill("solid", fgColor=cor)
        ws.row_dimensions[r].height = 20
        r += 1

    def L(rot, fx, key=None, fmt=BRL, bold=False, obs=None, grande=False, cor=PRETO,
          nivel=1):
        nonlocal r
        c = ws.cell(r, 2, rot)
        c.font = f(13 if grande else 10, bold or grande, cor)
        c.alignment = Alignment(indent=nivel, vertical="center")
        cc = ws.cell(r, 3, fx)
        cc.number_format = fmt
        cc.font = f(16 if grande else 10, True, BRANCO if grande else PRETO)
        cc.alignment = Alignment(horizontal="center", vertical="center")
        cc.border = BORDA
        if grande:
            cc.fill = F_DESTAQ
            ws.row_dimensions[r].height = 30
        elif bold:
            cc.fill = F_TOT
        if obs:
            o = ws.cell(r, 8, obs)
            o.font = f(8, False, CINZA, italic=True)
            o.alignment = Alignment(wrap_text=True, vertical="center", indent=1)
            ws.row_dimensions[r].height = max(ws.row_dimensions[r].height or 15,
                                              15, 13 + 11 * (len(obs) // 80))
        if key:
            R[key] = r
        r += 1
        return r - 1

    # ============================================= 1. QUANTO PRECISA
    bloco("1.  QUANTO A EMPRESA PRECISA DE DINHEIRO", "C00000")
    L("QUANTO AINDA FALTA CAPTAR", f"={fim('falta_pior')}", "falta", grande=True,
      obs="A resposta curta. É o pior mês de toda a projeção: quanto teria de estar em caixa e "
          "não está, já contando a reserva mínima e já descontando os aportes que a aba "
          "Aportes considera. Detalhe mês a mês na DRE Mensal, seção 6.")
    _mk_neg = rng("mk_neg")
    L("Até quando o dinheiro precisa estar na conta",
      '=IF(MIN(' + _mk_neg + ')=9999,"o caixa nunca fica negativo",'
      + mes_de("MIN(" + _mk_neg + ")") + ')',
      "prazo", fmt=TXT, bold=True,
      obs="Primeiro mês em que o caixa fica negativo. É o prazo que não se negocia.")
    L("O que a operação consome antes de se sustentar",
      f'=-MIN(0,MIN({rng("fcl_ac")}))', "consumo",
      obs="Fundo do poço do fluxo de caixa livre acumulado. Independe de quando o dinheiro entra.")
    L("(+) Reserva de caixa no mês mais apertado",
      f'=INDEX({rng("reserva")},MATCH(MIN({rng("caixa")}),{rng("caixa")},0))', "reserva")
    L("(=) Necessidade total de capital",
      f'=$C${R["consumo"]}+$C${R["reserva"]}', "necessidade", bold=True)
    L("(−) Capital que a aba Aportes considera", f"=-Aportes!$B${A['res_total']}", "cap")
    L("   do qual já entrou de verdade (Realizado)", f"=Aportes!$B${A['res_real']}", "cap_real",
      cor=CINZA, nivel=2,
      obs="O resto depende de assinatura. Desligue os interruptores na aba Aportes para ver a "
          "empresa só com o dinheiro que já está na conta.")
    L("(=) Falta captar, ignorando as datas de entrada",
      f'=MAX(0,$C${R["necessidade"]}+$C${R["cap"]})', "falta_simples",
      obs="Serve para dimensionar a rodada. A linha grande lá em cima é maior quando parte do "
          "capital chega depois do mês em que o caixa aperta.")
    L("Caixa mínimo projetado", f'=MIN({rng("caixa")})', "caixa_min", bold=True)
    _cx = rng("caixa")
    L("Mês do caixa mínimo",
      "=" + mes_de("MATCH(MIN(" + _cx + ")," + _cx + ",0)"),
      "caixa_min_mes", fmt=TXT)
    r += 1

    # ============================================= 2. O NEGOCIO
    bloco("2.  O NEGÓCIO")
    L("Base de usuários em dez/2030", f"={fim('base_fim')}", "base", INT)
    L("Assinantes pagantes em dez/2030", f"={fim('pag')}", "pag", INT, bold=True)
    L("% da base que paga em dez/2030", f"={fim('conv_real')}", "conv", PCT)
    L("Faturamento bruto de 2027",
      f'=SUMIFS({rng("bruto")},\'DRE Mensal\'!${L_INI}${R_ANO}:${L_FIM}${R_ANO},2027)', "fat27")
    L("Faturamento bruto de 2030",
      f'=SUMIFS({rng("bruto")},\'DRE Mensal\'!${L_INI}${R_ANO}:${L_FIM}${R_ANO},2030)', "fat30")
    L("Faturamento bruto acumulado em 5 anos", f'=SUM({rng("bruto")})', "fat", bold=True)
    L("EBITDA acumulado em 5 anos", f'=SUM({rng("ebitda")})', "ebitda", bold=True)
    L("Margem EBITDA em dez/2030", f"={fim('mg_ebitda')}", "mg", PCT)
    _mk_eb = rng("mk_eb")
    L("Primeiro mês com EBITDA positivo",
      '=IF(MIN(' + _mk_eb + ')=9999,"não atinge",' + mes_de("MIN(" + _mk_eb + ")") + ')',
      "be", fmt=TXT, bold=True,
      obs="Quando a operação passa a se pagar. Antes disso, cada mês é financiado por capital.")
    r += 1

    # ============================================= 3. PRECO E CUSTO
    bloco("3.  PREÇO E CUSTO POR ASSINANTE — em dez/2030")
    L("Preço médio por assinante", f"={fim('preco')}", "preco", BRL2)
    L("Fica para a empresa, de cada R$ 1 faturado", f"={fim('retencao')}", "ret", PCT,
      obs="O resto é imposto, comissão de loja e meio de pagamento.")
    L("Custo de nuvem por assinante pagante", f"={fim('nuvem_unit')}", "nuvem_u", BRL2,
      obs="A conta inteira de nuvem dividida pelos pagantes — inclui o que o plano gratuito "
          "consome. Cai com a escala: no lançamento é muito maior.")
    L("   do qual é subsídio do plano gratuito", f"={fim('subsidio')}", "subsidio", BRL2,
      nivel=2, cor=VERM,
      obs="O que cada assinante paga, sem saber, para sustentar quem não paga. No mês do "
          "lançamento este número passa do próprio preço da assinatura.")
    L("Custo variável por assinante (nuvem e suporte)", f"={fim('cv_unit')}", "cv", BRL2)
    L("Sobra por assinante antes do custo fixo", f"={fim('mc_unit')}", "mc", BRL2, bold=True,
      obs="Se for negativa, cada assinante novo dá prejuízo e nenhum volume resolve.")
    L("CAC por assinante pagante conquistado", f"={fim('cac_pag')}", "cac", BRL2)
    L("Assinantes necessários para o mês fechar no zero", f"={fim('n_eq')}", "n_eq", INT,
      bold=True)
    L("Assinantes projetados", f"={fim('pag')}", "n_proj", INT, cor=CINZA)
    r += 1

    # ============================================= 4. VIABILIDADE
    bloco("4.  VIABILIDADE")
    ws.cell(r, 2, "TMA — retorno mínimo exigido ao ano").font = f(10, True, PRETO)
    ws.cell(r, 2).alignment = Alignment(indent=1)
    entrada(ws.cell(r, 3), 0.25, PCT, chave=True)
    R["tma"] = r
    o = ws.cell(r, 8, "Custo de oportunidade do capital. 25% ao ano é referência para risco de "
                      "early stage no Brasil.")
    o.font = f(8, False, CINZA, italic=True)
    o.alignment = Alignment(wrap_text=True, vertical="center", indent=1)
    r += 1
    L("TMA equivalente ao mês", f'=(1+$C${R["tma"]})^(1/12)-1', "tma_m", PCT2, cor=CINZA)
    L("VPL — Valor Presente Líquido",
      f'=NPV($C${R["tma_m"]},{rng("fcl")})', "vpl", bold=True,
      obs="Sobre o fluxo de caixa livre, sem aportes. Positivo significa que o projeto cria "
          "valor acima da TMA.")
    L("TIR ao ano", f'=IFERROR((1+IRR({rng("fcl")}))^12-1,"n/d")', "tir", PCT)
    L("Payback (meses até o acumulado virar)",
      f'=IF({fim("fcl_ac")}<0,"> 60",MAX({rng("mk_pb")})+1)', "pb", INT)
    L("Caixa em dez/2030", f"={fim('caixa')}", "caixa_fim")
    r += 1

    # ============================================= 5. CONFERENCIAS
    bloco("5.  CONFERÊNCIAS — têm de ficar zeradas", "808080")
    # MAX(ABS(faixa)) sem formula matricial: SUMPRODUCT forca o calculo em array
    L("Faturamento líquido − custo − EBITDA",
      f'=SUMPRODUCT(MAX(ABS({rng("chk1")})))', "k1", BRL2, cor=CINZA)
    L("Base de usuários − pagantes − gratuitos",
      f'=SUMPRODUCT(MAX(ABS({rng("chk2")})))', "k2", INT, cor=CINZA)
    L("A conta de nuvem nunca cai quando a base sobe",
      f"='{NU['aba']}'!${L_TOT}${NU['chk']}", "k3", fmt=TXT, cor=CINZA)
    r += 1

    # ============================================= 6. O QUE OLHAR
    bloco("6.  O QUE MEXER PARA MELHORAR O RESULTADO", "ED7D31")
    for txt in [
        "Nuvem — aba Nuvem e Escala, coluna B. É o custo que mais pesa por assinante. Peça "
        "cotação ao provedor em dois ou três tamanhos de base e troque as âncoras.",
        "Conversão — DRE Mensal, seção 2. Mover a conversão move a receita sem mexer no preço "
        "nem no custo de aquisição. É a alavanca mais barata que existe.",
        "Comissão das lojas — DRE Mensal, seção 4. Vender a assinatura fora do app reduz os "
        "12,8% para a taxa de cartão. Em compra única é permitido pelas regras das lojas; "
        "em assinatura recorrente há restrição, vale conferir antes.",
        "Data dos aportes — aba Aportes. Dinheiro certo na data errada não paga conta. "
        "Antecipar um aporte já previsto pode resolver parte da lacuna sem captar mais nada.",
        "Equipe — aba Equipe. Depois da nuvem, é o maior custo fixo. Mexer no mês de início de "
        "cada contratação muda o caixa de 2027 inteiro.",
    ]:
        c = ws.cell(r, 2, "•  " + txt)
        c.font = f(9, False, "262626")
        c.alignment = Alignment(wrap_text=True, vertical="top", indent=1)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6)
        ws.row_dimensions[r].height = 13 + 11 * (len(txt) // 95 + 1)
        r += 1

    ws.sheet_properties.tabColor = "1F3864"
    ws.freeze_panes = "A4"
    return R
