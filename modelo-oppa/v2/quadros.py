# -*- coding: utf-8 -*-
"""Abas Equipe e Aportes — as duas tabelas que o usuario ja mantem.

Os dados vieram do arquivo dele. A estrutura e a mesma de antes porque funciona:
uma linha por item, data de inicio e de fim, e um espalhamento mensal embaixo.
"""
import datetime as dt
from openpyxl.styles import Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from base import *

# ------------------------------------------------------------------ EQUIPE
# funcao, tipo, status, qtd, custo unit., encargos, inicio, fim, classificacao
TIME = [
 ("Equipe PJ — desenvolvimento do MVP",                 "PJ","Contratado",1,10500,0,(2026,9),(2026,10),"Investimento"),
 ("Desenvolvimento — time base",                        "PJ","Plano",2, 9000,0,(2027,1),(2030,12),"Despesa"),
 ("Desenvolvimento — ampliação 2028 (adicional)",       "PJ","Plano",3, 9500,0,(2028,1),(2030,12),"Despesa"),
 ("Desenvolvimento — ampliação 2029 (adicional)",       "PJ","Plano",5,10000,0,(2029,1),(2030,12),"Despesa"),
 ("Desenvolvimento — ampliação 2030 (adicional)",       "PJ","Plano",6,10500,0,(2030,1),(2030,12),"Despesa"),
 ("Suporte ao cliente — time base",                     "PJ","Plano",1, 3500,0,(2027,1),(2030,12),"Despesa"),
 ("Suporte — ampliação 2028 (adicional)",               "PJ","Plano",3, 3500,0,(2028,1),(2030,12),"Despesa"),
 ("Suporte — ampliação 2029 (adicional)",               "PJ","Plano",6, 3800,0,(2029,1),(2030,12),"Despesa"),
 ("Suporte — ampliação 2030 (adicional)",               "PJ","Plano",8, 4000,0,(2030,1),(2030,12),"Despesa"),
 ("Growth e marketing",                                 "PJ","Plano",1, 6000,0,(2027,1),(2030,12),"Despesa"),
 ("Growth — ampliação 2028 (adicional)",                "PJ","Plano",2, 7000,0,(2028,1),(2030,12),"Despesa"),
 ("Growth — ampliação 2029 (adicional)",                "PJ","Plano",3, 8000,0,(2029,1),(2030,12),"Despesa"),
 ("Dados e infraestrutura",                             "PJ","Plano",1,10000,0,(2028,1),(2030,12),"Despesa"),
 ("Dados e infraestrutura — ampliação 2029 (adicional)","PJ","Plano",2,11000,0,(2029,1),(2030,12),"Despesa"),
 ("Produto e gestão",                                   "PJ","Plano",1,12000,0,(2028,1),(2030,12),"Despesa"),
 ("Produto e gestão — ampliação 2030 (adicional)",      "PJ","Plano",2,13000,0,(2030,1),(2030,12),"Despesa"),
 ("Pró-labore dos sócios (2027–2028)",                  "Sócio","Plano",3,5000,0,(2027,1),(2028,12),"Despesa"),
 ("Pró-labore dos sócios (2029–2030)",                  "Sócio","Plano",4,8000,0,(2029,1),(2030,12),"Despesa"),
 ("(exemplo) Contratação CLT",                          "CLT","Plano",0,8000,0.7,(2028,1),(2030,12),"Despesa"),
]
VAZIAS_TIME = 8

# ------------------------------------------------------------------ APORTES
# data, nome, valor, status, tipo
LANC = [
 ((2026,4,16),  "Paulo Kapp",        10000, "Realizado",     "Fundador"),
 ((2026,7,4),   "Lucas Medina",       5000, "Realizado",     "Fundador"),
 ((2026,8,20),  "Paulo Kapp",        10000, "Realizado",     "Fundador"),
 ((2026,8,20),  "Felipe Toescher",    3000, "Realizado",     "Fundador"),
 ((2026,8,20),  "Lucas Iung",         2000, "Realizado",     "Fundador"),
 ((2026,8,29),  "Lucas Iung",         3000, "Realizado",     "Fundador"),
 ((2026,9,4),   "Felipe Toescher",    2000, "Realizado",     "Fundador"),
 ((2026,9,16),  "Lucas Iung",         1000, "Realizado",     "Fundador"),
 ((2026,9,16),  "Lucas Medina",       1000, "Realizado",     "Fundador"),
 ((2026,10,15), "C-ioT",             40000, "Em negociação", "Novo sócio"),
 ((2026,10,15), "Felipe Martinelli", 40000, "Em negociação", "Novo sócio"),
]
VAZIAS_LANC = 14
# investidor, status, aporte, mes do aporte, receita B2B/mes, inicio, fim
ESTRAT = [("Boston Scientific", "Em negociação", 200000, (2027,2), 5000, (2027,3), (2030,12))]
VAZIAS_ESTRAT = 6


def _dv(ws, col, opcoes, r0, r1):
    dv = DataValidation(type="list", formula1=f'"{opcoes}"', allow_blank=True, showDropDown=False)
    ws.add_data_validation(dv)
    dv.add(f"{col}{r0}:{col}{r1}")


# =====================================================================
def equipe(wb):
    ws = wb.create_sheet("Equipe")
    larguras(ws, a=46)
    titulo(ws, "EQUIPE — UMA LINHA POR PAPEL",
           "Cada linha vale do mês de início ao mês de fim. As linhas de ampliação são "
           "ADICIONAIS: somam-se às anteriores. O interruptor de plano desliga tudo o que ainda "
           "não foi contratado.")
    tempo(ws)
    E = {}
    r = 8
    c = ws.cell(r, 1, "Considerar o plano de contratações?   1 = Sim · 0 = só o que já está contratado")
    c.font = f(10, True, PRETO)
    entrada(ws.cell(r, 2), 1, "0", chave=True)
    nota(ws, r, "Com 0, o modelo mostra a Oppa apenas com a equipe do MVP. Serve para ver o piso "
                "de custo.", col_ini=4, ate=12)
    E["plano_on"] = r
    r += 2

    secao(ws, r, "QUADRO — edite, acrescente ou apague linhas")
    r += 1
    hdr = ["Função / papel", "Tipo", "Status", "Qtd", "Custo mensal\npor pessoa (R$)",
           "Encargos\n(% sobre o custo)", "Custo total\nmensal (R$)", "Início", "Fim",
           "Classificação"]
    for j, t in enumerate(hdr):
        cc = ws.cell(r, 1 + j, t)
        cc.font = f(9, True, BRANCO)
        cc.fill = F_HDR
        cc.alignment = Alignment(horizontal="center", wrap_text=True, vertical="center")
    ws.row_dimensions[r].height = 30
    r += 1
    E["ini"] = r
    linhas = list(TIME) + [("", "PJ", "Plano", 0, 0, 0, (2027, 1), (2030, 12), "Despesa")] * VAZIAS_TIME
    for (fn, tipo, st, qtd, custo, enc, ini, fim, cls) in linhas:
        ws.cell(r, 1, fn).font = f(9, bool(fn), PRETO if fn else CINZA)
        ws.cell(r, 1).alignment = Alignment(indent=1)
        for cidx, val, fmt in ((2, tipo, TXT), (3, st, TXT), (4, qtd, INT),
                               (5, custo, BRL), (6, enc, PCT)):
            entrada(ws.cell(r, cidx), val, fmt)
        ct = ws.cell(r, 7, f"=$D{r}*$E{r}*(1+$F{r})")
        ct.number_format = BRL
        ct.font = f(9, True, PRETO)
        ct.border = BORDA
        ct.alignment = Alignment(horizontal="center")
        entrada(ws.cell(r, 8), dt.datetime(ini[0], ini[1], 1), MES)
        entrada(ws.cell(r, 9), dt.datetime(fim[0], fim[1], 1), MES)
        entrada(ws.cell(r, 10), cls, TXT)
        r += 1
    E["fim"] = r - 1
    i, fi = E["ini"], E["fim"]
    _dv(ws, "B", "PJ,CLT,Sócio", i, fi)
    _dv(ws, "C", "Contratado,Plano", i, fi)
    _dv(ws, "J", "Despesa,Investimento", i, fi)
    ws.cell(r, 1, "TOTAL DO QUADRO (soma de tudo, sem olhar data)").font = f(10, True)
    ct = ws.cell(r, 7, f"=SUM(G{i}:G{fi})")
    ct.number_format = BRL
    ct.font = f(10, True)
    ct.fill = F_TOT
    r += 2

    secao(ws, r, "O CUSTO MÊS A MÊS")
    r += 1

    def mensal(rot, filtro, key, bold=False, cor=PRETO, obs=None, fmt=BRL):
        nonlocal r
        cl = ws.cell(r, 1, rot)
        cl.font = f(10, bold, cor)
        cl.alignment = Alignment(indent=0 if bold else 1, vertical="center")
        for k in range(1, N + 1):
            L = col(k)
            cc = ws[f"{L}{r}"]
            cc.value = (f'=SUMPRODUCT({filtro}'
                        f'(($C${i}:$C${fi}="Contratado")+($C${i}:$C${fi}="Plano")*$B${E["plano_on"]})'
                        f'*($H${i}:$H${fi}<={L}${R_DATA})*($I${i}:$I${fi}>={L}${R_DATA})'
                        f'*{"$D" if fmt == INT else "$G"}${i}:{"$D" if fmt == INT else "$G"}${fi})')
            cc.number_format = fmt
            cc.font = f(9, bold, cor)
        ct = ws[f"{L_TOT}{r}"]
        ct.value = f"=SUM({L_INI}{r}:{L_FIM}{r})" if fmt == BRL else f"={L_FIM}{r}"
        ct.number_format = fmt
        ct.font = f(9, True, PRETO)
        ct.fill = F_TOT
        if obs:
            nota(ws, r, obs)
        E[key] = r
        r += 1
        return E[key]

    mensal("CUSTO DE EQUIPE NO MÊS (despesa)", f'($J${i}:$J${fi}="Despesa")*', "despesa",
           bold=True, obs="Entra na DRE como despesa de pessoal.")
    mensal("Equipe PJ", f'($B${i}:$B${fi}="PJ")*($J${i}:$J${fi}="Despesa")*', "pj")
    mensal("Equipe CLT (com encargos)", f'($B${i}:$B${fi}="CLT")*($J${i}:$J${fi}="Despesa")*', "clt")
    mensal("Pró-labore dos sócios", f'($B${i}:$B${fi}="Sócio")*($J${i}:$J${fi}="Despesa")*', "socio")
    mensal("Desenvolvimento capitalizado (investimento, não despesa)",
           f'($J${i}:$J${fi}="Investimento")*', "invest", cor=VERDE,
           obs="Vai para a linha de investimento da DRE, não para a de pessoal.")
    mensal("Pessoas na operação (headcount)", "", "headcount", fmt=INT, cor=CINZA)

    ws.sheet_properties.tabColor = "ED7D31"
    return ws, E


# =====================================================================
def aportes(wb):
    ws = wb.create_sheet("Aportes")
    larguras(ws, a=34)
    titulo(ws, "APORTES E INVESTIDORES — TODO O CAPITAL QUE ENTRA",
           "O STATUS de cada linha decide se ela entra no modelo. Os dois interruptores abaixo "
           "ligam e desligam o que ainda não está assinado.")
    tempo(ws)
    A = {}
    r = 8
    for rot, val, obs in [
        ("Considerar aportes PREVISTOS?   1 = Sim · 0 = Não", 1,
         "Negociação adiantada, sem assinatura."),
        ("Considerar acordos EM NEGOCIAÇÃO?   1 = Sim · 0 = Não", 1,
         "Sem compromisso nenhum. Hoje são os R$ 80 mil de C-ioT e Martinelli e os R$ 200 mil "
         "da Boston. Desligue para ver a Oppa só com o dinheiro que já entrou."),
    ]:
        c = ws.cell(r, 1, rot)
        c.font = f(10, True, PRETO)
        entrada(ws.cell(r, 2), val, "0", chave=True)
        nota(ws, r, obs, col_ini=4, ate=12)
        A["inc_previsto" if "PREVISTOS" in rot else "inc_negoc"] = r
        r += 1
    r += 1

    secao(ws, r, "1. LANÇAMENTOS DE APORTE")
    r += 1
    for j, t in enumerate(["Data", "Sócio / Investidor", "Valor (R$)", "Status", "Tipo"]):
        c = ws.cell(r, 1 + j, t)
        c.font = f(9, True, BRANCO)
        c.fill = F_HDR
        c.alignment = Alignment(horizontal="center", indent=1)
    r += 1
    A["ini"] = r
    for (d, nome, val, st, tipo) in LANC + [((2027, 1, 15), "", 0, "Previsto", "Novo sócio")] * VAZIAS_LANC:
        entrada(ws.cell(r, 1), dt.datetime(*d), 'dd/mm/yyyy')
        entrada(ws.cell(r, 2), nome, TXT)
        entrada(ws.cell(r, 3), val, BRL)
        entrada(ws.cell(r, 4), st, TXT)
        entrada(ws.cell(r, 5), tipo, TXT)
        r += 1
    A["fim"] = r - 1
    i, fi = A["ini"], A["fim"]
    _dv(ws, "D", "Realizado,Compromisso,Previsto,Em negociação", i, fi)
    _dv(ws, "E", "Fundador,Novo sócio,Investidor estratégico", i, fi)
    gate = (f'(($D${i}:$D${fi}="Realizado")+($D${i}:$D${fi}="Compromisso")'
            f'+($D${i}:$D${fi}="Previsto")*$B${A["inc_previsto"]}'
            f'+($D${i}:$D${fi}="Em negociação")*$B${A["inc_negoc"]})')
    A["gate"] = gate
    ws.cell(r, 2, "TOTAL que o modelo considera").font = f(10, True)
    ct = ws.cell(r, 3, f"=SUMPRODUCT({gate}*$C${i}:$C${fi})")
    ct.number_format = BRL
    ct.font = f(10, True)
    ct.fill = F_TOT
    A["total_lanc"] = r
    r += 2

    secao(ws, r, "2. INVESTIDOR ESTRATÉGICO — aporta e traz receita recorrente")
    r += 1
    for j, t in enumerate(["Investidor", "Status", "Aporte (R$)", "Mês do aporte",
                           "Receita B2B (R$/mês)", "Início da receita", "Fim da receita"]):
        c = ws.cell(r, 1 + j, t)
        c.font = f(9, True, BRANCO)
        c.fill = F_HDR
        c.alignment = Alignment(horizontal="center", wrap_text=True, indent=1)
    ws.row_dimensions[r].height = 26
    r += 1
    A["est_ini"] = r
    for (nome, st, ap, mes, rec, ri, rf) in ESTRAT + [("", "Em negociação", 0, (2028,1), 0, (2028,1), (2030,12))] * VAZIAS_ESTRAT:
        entrada(ws.cell(r, 1), nome, TXT)
        entrada(ws.cell(r, 2), st, TXT)
        entrada(ws.cell(r, 3), ap, BRL)
        entrada(ws.cell(r, 4), dt.datetime(*mes, 1), MES)
        entrada(ws.cell(r, 5), rec, BRL)
        entrada(ws.cell(r, 6), dt.datetime(*ri, 1), MES)
        entrada(ws.cell(r, 7), dt.datetime(*rf, 1), MES)
        r += 1
    A["est_fim"] = r - 1
    ei, ef = A["est_ini"], A["est_fim"]
    _dv(ws, "B", "Confirmado,Em negociação", ei, ef)
    gate_e = (f'(($B${ei}:$B${ef}="Confirmado")'
              f'+($B${ei}:$B${ef}="Em negociação")*$B${A["inc_negoc"]})')
    A["gate_est"] = gate_e
    ws.cell(r, 1, "TOTAL").font = f(10, True)
    for cidx, rng in ((3, f"$C${ei}:$C${ef}"), (5, f"$E${ei}:$E${ef}")):
        ct = ws.cell(r, cidx, f"=SUMPRODUCT({gate_e}*{rng})")
        ct.number_format = BRL
        ct.font = f(10, True)
        ct.fill = F_TOT
    r += 2

    secao(ws, r, "3. O CAPITAL MÊS A MÊS")
    r += 1

    def mensal(rot, gen, key, bold=False, cor=PRETO, obs=None):
        nonlocal r
        cl = ws.cell(r, 1, rot)
        cl.font = f(10, bold, cor)
        cl.alignment = Alignment(indent=0 if bold else 1, vertical="center")
        for k in range(1, N + 1):
            L = col(k)
            cc = ws[f"{L}{r}"]
            cc.value = gen(k, L)
            cc.number_format = BRL
            cc.font = f(9, bold, cor)
        ct = ws[f"{L_TOT}{r}"]
        ct.value = f"=SUM({L_INI}{r}:{L_FIM}{r})"
        ct.number_format = BRL
        ct.font = f(9, True, PRETO)
        ct.fill = F_TOT
        if obs:
            nota(ws, r, obs)
        A[key] = r
        r += 1
        return A[key]

    mes_lanc = (f'($A${i}:$A${fi}>={{L}}${R_DATA})'
                f'*($A${i}:$A${fi}<=EOMONTH({{L}}${R_DATA},0))')
    mensal("Aportes de sócios e novos sócios",
           lambda k, L: f"=SUMPRODUCT({gate}*{mes_lanc.format(L=L)}*$C${i}:$C${fi})", "lanc")
    mensal("Aporte do investidor estratégico",
           lambda k, L: (f"=SUMPRODUCT({gate_e}*($D${ei}:$D${ef}={L}${R_DATA})"
                         f"*$C${ei}:$C${ef})"), "estr")
    mensal("TOTAL DE APORTES NO MÊS",
           lambda k, L: f'=SUM({L}{A["lanc"]}:{L}{A["estr"]})', "total", bold=True)
    mensal("Receita B2B do investidor estratégico",
           lambda k, L: (f"=SUMPRODUCT({gate_e}*($F${ei}:$F${ef}<={L}${R_DATA})"
                         f"*($G${ei}:$G${ef}>={L}${R_DATA})*$E${ei}:$E${ef})"), "b2b",
           cor=VERDE, obs="Entra na DRE como outra receita, não como aporte.")
    r += 1

    secao(ws, r, "4. DE ONDE VEM O CAPITAL")
    r += 1
    for rot, fx, key in [
        ("Realizado — dinheiro que já entrou", f'=SUMIFS($C${i}:$C${fi},$D${i}:$D${fi},"Realizado")', "res_real"),
        ("Compromisso assinado, ainda não integralizado", f'=SUMIFS($C${i}:$C${fi},$D${i}:$D${fi},"Compromisso")', "res_comp"),
        ("Previsto — negociação adiantada", f'=SUMIFS($C${i}:$C${fi},$D${i}:$D${fi},"Previsto")*$B${A["inc_previsto"]}', "res_prev"),
        ("Em negociação — sem compromisso", f'=SUMIFS($C${i}:$C${fi},$D${i}:$D${fi},"Em negociação")*$B${A["inc_negoc"]}', "res_negoc"),
        ("Investidor estratégico", f"=SUMPRODUCT({gate_e}*$C${ei}:$C${ef})", "res_estr"),
        ("CAPITAL TOTAL CONSIDERADO", None, "res_total"),
    ]:
        ws.cell(r, 1, rot).font = f(10, key == "res_total", PRETO)
        ws.cell(r, 1).alignment = Alignment(indent=1)
        ct = ws.cell(r, 2, fx if fx else
                     f'=SUM(B{A["res_real"]}:B{A["res_estr"]})')
        ct.number_format = BRL
        ct.font = f(10, True, PRETO)
        if key == "res_total":
            ct.fill = F_TOT
        A[key] = r
        r += 1

    ws.sheet_properties.tabColor = "ED7D31"
    return ws, A
