# -*- coding: utf-8 -*-
"""Aba 'DRE Mensal' — a espinha do modelo. Tudo em uma aba, mes a mes.

Como funciona uma linha de driver:

    B  valor inicial          o valor no primeiro mes em que o item existe
    C  evolucao por mes       quanto ele cresce (ou cai) a cada mes
    D  mes de inicio          antes disso a linha vale zero
    E..BL                     calculado: mes anterior x (1 + evolucao)

Digitar POR CIMA de qualquer mes reancora a curva dali para frente, porque cada
mes le o mes anterior. E a resposta para "a premissa precisa ser mensal": e
mensal, e mexer em um mes nao obriga a mexer nos outros 59.
"""
from openpyxl.styles import Alignment
from base import *

ABA = "DRE Mensal"

# faixas do Anexo V do Simples Nacional (LC 123/2006)
FAIXAS = [(0, 0.155, 0), (180_000.01, 0.18, 4_500), (360_000.01, 0.195, 9_900),
          (720_000.01, 0.205, 17_100), (1_800_000.01, 0.23, 62_100),
          (3_600_000.01, 0.305, 540_000)]
ALIQ_LP = 0.1633     # acima do teto do Simples


def build(wb, CTRL, E, A, NU):
    """CTRL: linhas de controle no Resumo · E: Equipe · A: Aportes · NU: Nuvem."""
    ws = wb.create_sheet(ABA)
    larguras(ws, a=52)
    titulo(ws, "DRE MENSAL — DOS DRIVERS AO RESULTADO, EM UMA ABA",
           "Uma coluna por mês, de jan/2026 a dez/2030. Você digita nas colunas B, C e D — e "
           "pode digitar por cima de qualquer mês. Nada aqui é anual.")
    tempo(ws)
    D = {}
    r = 8

    # -------------------------------------------------- helpers
    def rot(r_, texto, bold=False, nivel=1, cor=PRETO, italic=False):
        c = ws.cell(r_, 1, texto)
        c.font = f(10, bold, cor, italic)
        c.alignment = Alignment(indent=0 if bold else nivel, vertical="center")
        return c

    def calc(texto, gen, key, fmt=BRL, bold=False, cor=PRETO, nivel=1,
             destaque=False, obs=None, total="soma", italic=False):
        """Linha calculada: sem colunas B/C/D."""
        nonlocal r
        rot(r, texto, bold, nivel, cor, italic)
        for i in range(1, N + 1):
            L = col(i)
            cc = ws[f"{L}{r}"]
            cc.value = gen(i, L, ant(i) if i > 1 else None)
            cc.number_format = fmt
            cc.font = f(9, bold, cor)
            if destaque:
                cc.fill = F_TOT
        ct = ws[f"{L_TOT}{r}"]
        ct.value = (f"=SUM({L_INI}{r}:{L_FIM}{r})" if total == "soma" else
                    f"={L_FIM}{r}" if total == "ultimo" else None)
        if ct.value:
            ct.number_format = fmt
            ct.font = f(9, True, PRETO)
            ct.fill = F_TOT
        if obs:
            nota(ws, r, obs)
        D[key] = r
        r += 1
        return D[key]

    def driver(texto, inicial, evol, mes_ini, key, fmt=BRL, chave=False, obs=None,
               teto=None, bold=False):
        """Linha de premissa mensal: inicial + evolucao, reancoravel."""
        nonlocal r
        rot(r, texto, bold)
        entrada(ws.cell(r, 2), inicial, fmt, chave=chave)
        entrada(ws.cell(r, 3), evol, PCT)
        entrada(ws.cell(r, 4), mes_ini, INT)
        for i in range(1, N + 1):
            L, P = col(i), ant(i) if i > 1 else None
            corpo = (f"$B{r}" if i == 1 else
                     f"IF({P}{r}=0,$B{r},{P}{r}*(1+$C{r}))")
            if teto:
                corpo = f"MIN({teto},{corpo})"
            cc = ws[f"{L}{r}"]
            cc.value = f"=IF({L}${R_IDX}<$D{r},0,{corpo})"
            cc.number_format = fmt
            cc.font = f(9, False, PRETO)
            cc.fill = F_SUB
        ct = ws[f"{L_TOT}{r}"]
        ct.value = f"={L_FIM}{r}"
        ct.number_format = fmt
        ct.font = f(9, True, PRETO)
        ct.fill = F_TOT
        if obs:
            nota(ws, r, obs)
        D[key] = r
        r += 1
        return D[key]

    def avulso(texto, key, fmt=BRL, obs=None):
        """Linha de valor pontual: cada mes e uma celula em branco para digitar."""
        nonlocal r
        rot(r, texto)
        for cidx in (2, 3, 4):
            c = ws.cell(r, cidx, "—")
            c.font = f(9, False, CINZA)
            c.alignment = Alignment(horizontal="center")
        for i in range(1, N + 1):
            cc = ws[f"{col(i)}{r}"]
            cc.number_format = fmt
            cc.font = f(9, False, AZUL)
            cc.fill = F_IN
            cc.border = BORDA
        ct = ws[f"{L_TOT}{r}"]
        ct.value = f"=SUM({L_INI}{r}:{L_FIM}{r})"
        ct.number_format = fmt
        ct.font = f(9, True, PRETO)
        ct.fill = F_TOT
        if obs:
            nota(ws, r, obs)
        D[key] = r
        r += 1
        return D[key]

    # ================================================== 1. CENARIO
    secao(ws, r, "1. CENÁRIO ATIVO  —  troque no Resumo e esta aba inteira responde"); r += 1
    rot(r, "Cenário em uso", bold=True)
    c = ws.cell(r, 2, f"=Resumo!$C${CTRL['cenario']}")
    c.number_format = "0"; c.font = f(10, True, VERDE); c.fill = F_TOT
    c.alignment = Alignment(horizontal="center")
    cn = ws.cell(r, 3, f'=CHOOSE($B{r},"CONSERVADOR","PROVÁVEL","AGRESSIVO")')
    cn.font = f(10, True, VERDE)
    nota(ws, r, "O número e os três multiplicadores abaixo estão na aba Resumo, no alto. "
                "Trocar um número só muda o modelo inteiro.", col_ini=6, ate=14)
    D["cenario"] = r; r += 1
    for nome, key in (("captação de usuários", "m_cap"), ("conversão em pagante", "m_conv"),
                      ("churn", "m_churn")):
        rot(r, f"Multiplicador do cenário — {nome}", nivel=2, cor=CINZA)
        c = ws.cell(r, 2, f"=INDEX(Resumo!$D${CTRL['mult_ini']}:$F${CTRL['mult_ini']+2},"
                          f"MATCH(\"{nome}\",Resumo!$C${CTRL['mult_ini']}:$C${CTRL['mult_ini']+2},0),"
                          f"$B${D['cenario']})")
        c.number_format = MULT; c.font = f(10, True, VERDE); c.alignment = Alignment(horizontal="center")
        D[key] = r; r += 1
    r += 1

    # ================================================== 2. DRIVERS
    secao(ws, r, "2. DRIVERS DO NEGÓCIO  —  digite em B (valor inicial), C (evolução por mês) "
                 "e D (mês de início)"); r += 1
    driver("Novos usuários captados no mês", 18100, 0.034, 12, "novos", INT, chave=True,
           obs="Mês 12 = dez/2026, o lançamento. 18.100 no primeiro mês crescendo 3,4% ao mês "
               "chega a ~89 mil/mês em dez/2030 — a mesma trajetória que você vinha usando. "
               "Digite por cima de qualquer mês para ancorar a curva ali.")
    driver("Churn mensal da base de usuários", 0.08, -0.005, 12, "churn", PCT, chave=True,
           obs="8% no lançamento caindo 0,5% ao mês (relativo) chega a 6,3% em 2030. "
               "Health/fitness roda entre 7% e 10% ao mês no primeiro ano.")
    driver("Conversão da base em assinante pagante", 0.005, 0.04, 12, "conv", PCT, chave=True,
           teto=0.5,
           obs="0,5% no lançamento crescendo 4% ao mês: 0,8% em um ano, 3,3% em dez/2030. "
               "Isto substitui a rampa de maturação da versão anterior — aqui a rampa é a "
               "própria curva. Mediana de freemium é 2% a 3%; o quartil de cima chega a 8%.")
    driver("Preço médio por assinante no mês (R$)", 39.90, 0.0, 12, "preco", BRL2, chave=True,
           obs="A média ponderada do que cada pagante paga. Com um plano só, é o preço dele. "
               "Ponha a evolução em 0,4%/mês para embutir reajuste anual de ~5%.")
    r += 1

    # ================================================== 3. USUARIOS
    secao(ws, r, "3. USUÁRIOS"); r += 1
    # Referencia para a frente: a base do inicio do mes le a base do fim do mes
    # ANTERIOR, que e escrita tres linhas abaixo. O numero precisa ser conhecido
    # antes, entao e reservado aqui.
    r_bfim = r + 3
    calc("Base no início do mês", lambda i, L, P: "=0" if i == 1 else f"={P}{r_bfim}",
         "base_ini", INT, total="ultimo")
    calc("(+) Novos usuários captados",
         lambda i, L, P: f"=ROUND({L}{D['novos']}*$B${D['m_cap']},0)", "novos_cen", INT)
    calc("(−) Usuários perdidos (churn)",
         lambda i, L, P: f"=-ROUND({L}{D['base_ini']}*{L}{D['churn']}*$B${D['m_churn']},0)",
         "perdidos", INT)
    calc("(=) BASE NO FIM DO MÊS",
         lambda i, L, P: f"={L}{D['base_ini']}+{L}{D['novos_cen']}+{L}{D['perdidos']}",
         "base_fim", INT, bold=True, destaque=True, total="ultimo",
         obs="Pagantes + gratuitos. É esta linha que a aba Nuvem e Escala usa para achar a faixa.")
    calc("Assinantes pagantes",
         lambda i, L, P: f"=ROUND({L}{D['base_fim']}*{L}{D['conv']}*$B${D['m_conv']},0)",
         "pag", INT, bold=True, destaque=True, total="ultimo")
    calc("Usuários no plano gratuito",
         lambda i, L, P: f"={L}{D['base_fim']}-{L}{D['pag']}", "free", INT, cor=CINZA,
         total="ultimo", obs="Não pagam e consomem nuvem. Quem assina sustenta quem não assina.")
    calc("% da base que paga", lambda i, L, P: f"=IFERROR({L}{D['pag']}/{L}{D['base_fim']},0)",
         "conv_real", PCT, cor=CINZA, total="ultimo")
    r += 1

    # ================================================== 4. FATURAMENTO
    secao(ws, r, "4. FATURAMENTO"); r += 1
    calc("Receita de assinaturas",
         lambda i, L, P: f"={L}{D['pag']}*{L}{D['preco']}", "rec_ass", BRL, bold=True)
    calc("Receita B2B do investidor estratégico",
         lambda i, L, P: f"=Aportes!{L}{A['b2b']}", "rec_b2b", BRL, cor=VERDE,
         obs="Vem da aba Aportes, tabela 2. Entra só se o status permitir.")
    avulso("Outras receitas (digite no mês)", "rec_out",
           obs="Venda de dados, marketplace, parceria pontual. Cada mês é uma célula livre.")
    calc("(=) FATURAMENTO BRUTO",
         lambda i, L, P: f"=SUM({L}{D['rec_ass']}:{L}{D['rec_out']})", "bruto", BRL,
         bold=True, destaque=True)
    r_rbt = r
    calc("Faturamento dos últimos 12 meses (RBT12)",
         lambda i, L, P: (f"={L}{D['bruto']}*12" if i == 1 else
                          (f"=SUM({L_INI}{D['bruto']}:{P}{D['bruto']})/{i-1}*12" if i <= 12 else
                           f"=SUM({col(i-12)}{D['bruto']}:{P}{D['bruto']})")),
         "rbt12", BRL, cor=CINZA, total="ultimo",
         obs="Base de cálculo da faixa do Simples. Nos 12 primeiros meses é a média anualizada.")
    driver("(−) Imposto sobre a receita", 0.155, 0.0, 1, "p_imp", PCT2, chave=True,
           obs="Você digita a alíquota. A linha cinza logo abaixo sugere a alíquota efetiva pela "
               "tabela do Anexo V, calculada a partir do RBT12 do mês.")
    FX_INI = None   # preenchido no fim; placeholder resolvido abaixo
    D["_sug_imp"] = r      # reservada: a formula entra depois da tabela de faixas
    rot(r, "sugestão pela tabela do Simples (Anexo V)", nivel=2, cor=CINZA, italic=True)
    r += 1
    calc("(−) Impostos sobre a receita",
         lambda i, L, P: f"=-{L}{D['bruto']}*{L}{D['p_imp']}", "imp", BRL)
    driver("(−) Comissão das lojas de aplicativos (% do bruto)", 0.128, -0.003, 1, "p_loja",
           PCT2, obs="15% da Apple/Google aplicados sobre a fatia cobrada dentro do app. "
                     "Com 85% dentro do app: 15% × 85% = 12,8%. Vender fora do app reduz isto.")
    calc("(−) Comissão das lojas de aplicativos",
         lambda i, L, P: f"=-{L}{D['rec_ass']}*{L}{D['p_loja']}", "loja", BRL)
    driver("(−) Taxa de meio de pagamento (% do bruto)", 0.0068, 0.0, 1, "p_pgto", PCT2,
           obs="4,5% de cartão/PIX sobre os 15% cobrados fora do app = 0,68%.")
    calc("(−) Taxas de meio de pagamento",
         lambda i, L, P: f"=-{L}{D['bruto']}*{L}{D['p_pgto']}", "pgto", BRL)
    calc("(=) FATURAMENTO LÍQUIDO",
         lambda i, L, P: f"={L}{D['bruto']}+{L}{D['imp']}+{L}{D['loja']}+{L}{D['pgto']}",
         "liquido", BRL, bold=True, destaque=True)
    calc("Fica para a empresa, de cada R$ 1 faturado",
         lambda i, L, P: f"=IFERROR({L}{D['liquido']}/{L}{D['bruto']},0)", "retencao", PCT,
         cor=CINZA, total="ultimo",
         obs="É o divisor de qualquer conta de preço: cobrar R$ 1 não é receber R$ 1.")
    r += 1

    # ================================================== 5. CUSTOS
    secao(ws, r, "5. CUSTOS E DESPESAS"); r += 1
    calc("(−) Nuvem e infraestrutura",
         lambda i, L, P: f"=-'{NU['aba']}'!{L}{NU['conta']}", "nuvem", BRL, cor=VERDE,
         obs="Vem da aba Nuvem e Escala: a conta interpolada pela base deste mês.")
    calc("do qual vem do plano gratuito",
         lambda i, L, P: f"=-'{NU['aba']}'!{L}{NU['parte_free']}", "nuvem_free", BRL, cor=CINZA,
         nivel=2, obs="O plano gratuito não paga e consome nuvem. Esta linha é o tamanho do "
                      "subsídio.")
    calc("Custo de nuvem por assinante pagante",
         lambda i, L, P: f"='{NU['aba']}'!{L}{NU['por_pag']}", "nuvem_unit", BRL2, cor=CINZA,
         total="ultimo",
         obs="A conta inteira dividida pelos pagantes — é com este número que o preço se "
             "compara. Cai com a escala: compare o lançamento com dez/2030.")
    calc("Subsídio do plano gratuito por assinante pagante",
         lambda i, L, P: f"='{NU['aba']}'!{L}{NU['subsidio']}", "subsidio", BRL2, cor=CINZA,
         total="ultimo")
    driver("(−) Suporte ao cliente (R$ por assinante pagante)", 0.35, -0.005, 12, "p_sup",
           BRL2, obs="Ferramentas de atendimento e tempo de equipe que escalam com o número de "
                     "pagantes. O time de suporte em si está na aba Equipe.")
    calc("(−) Suporte ao cliente",
         lambda i, L, P: f"=-{L}{D['pag']}*{L}{D['p_sup']}", "sup", BRL)
    calc("(−) Equipe (PJ, CLT e pró-labore)",
         lambda i, L, P: f"=-Equipe!{L}{E['despesa']}", "equipe", BRL, cor=VERDE,
         obs="Vem da aba Equipe, linha a linha.")
    driver("(−) Marketing de aquisição (R$ por novo usuário captado)", 2.50, 0.009, 12,
           "p_cac", BRL2, chave=True,
           obs="CAC por usuário captado, não por assinante. O CAC por ASSINANTE é este número "
               "dividido pela conversão — veja a seção 7.")
    calc("(−) Marketing de aquisição",
         lambda i, L, P: f"=-{L}{D['novos_cen']}*{L}{D['p_cac']}", "mkt_aq", BRL)
    driver("(−) Marketing recorrente (R$ no mês)", 500, 0.07, 11, "mkt_rec", BRL,
           obs="Conteúdo, ferramentas, influenciador. Começa um mês antes do lançamento.")
    driver("(−) Administrativo: contabilidade, jurídico, ferramentas (R$ no mês)", 600, 0.03,
           4, "adm", BRL,
           obs="Só o contador de uma empresa no Simples costuma custar de R$ 300 a R$ 800 por "
               "mês. Começa no mês 4 (abr/2026), com a constituição.")
    avulso("(−) Outras despesas (digite no mês)", "desp_out")
    calc("(=) CUSTO TOTAL DO MÊS",
         lambda i, L, P: (f"={L}{D['nuvem']}+{L}{D['sup']}+{L}{D['equipe']}+{L}{D['mkt_aq']}"
                          f"-{L}{D['mkt_rec']}-{L}{D['adm']}-{L}{D['desp_out']}"),
         "custo", BRL, bold=True, destaque=True,
         obs="Soma de tudo acima, em negativo. Não inclui imposto, loja e meio de pagamento, "
             "que já saíram do faturamento.")
    calc("(=) EBITDA",
         lambda i, L, P: f"={L}{D['liquido']}+{L}{D['custo']}", "ebitda", BRL, bold=True,
         destaque=True)
    calc("Margem EBITDA (% do faturamento líquido)",
         lambda i, L, P: f"=IFERROR({L}{D['ebitda']}/{L}{D['liquido']},0)", "mg_ebitda", PCT,
         cor=CINZA, total="ultimo")
    calc("EBITDA acumulado",
         lambda i, L, P: f"={L}{D['ebitda']}" if i == 1 else f"={P}{r}+{L}{D['ebitda']}",
         "ebitda_ac", BRL, cor=CINZA, total="ultimo")
    r += 1

    # ================================================== 6. CAIXA E CAPITAL
    secao(ws, r, "6. CAIXA E QUANTO FALTA DE DINHEIRO"); r += 1
    avulso("(−) Investimentos: equipamentos, marca, certificações (digite no mês)", "capex",
           obs="Sai do caixa no mês em que você digitar. Não é despesa do mês — é investimento.")
    calc("(−) Desenvolvimento capitalizado (da aba Equipe)",
         lambda i, L, P: f"=-Equipe!{L}{E['invest']}", "capex_dev", BRL, cor=VERDE)
    calc("(=) FLUXO DE CAIXA LIVRE (FCL)",
         lambda i, L, P: f"={L}{D['ebitda']}-{L}{D['capex']}+{L}{D['capex_dev']}", "fcl",
         BRL, bold=True, destaque=True,
         obs="O que a operação gera ou consome. Aporte NÃO entra aqui: aporte é financiamento.")
    calc("FCL acumulado",
         lambda i, L, P: f"={L}{D['fcl']}" if i == 1 else f"={P}{r}+{L}{D['fcl']}", "fcl_ac",
         BRL, cor=CINZA, total="ultimo",
         obs="O fundo do poço desta linha é quanto a operação consome antes de virar.")
    calc("(+) Aportes de sócios e investidores",
         lambda i, L, P: f"=Aportes!{L}{A['total']}", "aportes", BRL, cor=VERDE)
    calc("(=) Fluxo do mês (FCL + aportes)",
         lambda i, L, P: f"={L}{D['fcl']}+{L}{D['aportes']}", "fluxo", BRL)
    calc("(=) CAIXA ACUMULADO",
         lambda i, L, P: f"={L}{D['fluxo']}" if i == 1 else f"={P}{r}+{L}{D['fluxo']}",
         "caixa", BRL, bold=True, destaque=True, total="ultimo",
         obs="Se esta linha fica negativa, a conta da empresa não tem fundo naquele mês.")
    calc("Custo fixo do mês (referência da reserva)",
         lambda i, L, P: (f"=-({L}{D['nuvem']}+{L}{D['equipe']})"
                          f"+{L}{D['mkt_rec']}+{L}{D['adm']}"), "cf", BRL, cor=CINZA,
         obs="Nuvem, equipe, marketing recorrente e administrativo: o que se paga mesmo sem "
             "nenhum assinante novo.")
    rot(r, "Reserva mínima de caixa desejada (meses de custo fixo)", bold=True)
    entrada(ws.cell(r, 2), 3, "0", chave=True)
    nota(ws, r, "Caixa que precisa sobrar depois de pagar tudo. Três meses é o piso usual em "
                "estágio inicial. Zere para ver o mínimo matemático.", col_ini=6, ate=14)
    D["res_meses"] = r; r += 1
    calc("Reserva mínima em reais",
         lambda i, L, P: f"={L}{D['cf']}*$B${D['res_meses']}", "reserva", BRL, cor=CINZA, total="ultimo")
    calc("FALTA CAPTAR neste mês",
         lambda i, L, P: f"=MAX(0,{L}{D['reserva']}-{L}{D['caixa']})", "falta", BRL,
         bold=True, destaque=True, total="ultimo",
         obs="Quanto teria de estar em caixa neste mês e não está, já contando a reserva.")
    calc("FALTA CAPTAR — pior mês até aqui",
         lambda i, L, P: f"=MAX({L_INI}{D['falta']}:{L}{D['falta']})", "falta_pior", BRL,
         bold=True, total="ultimo",
         obs="Só sobe. O último valor é a resposta: quanto a empresa precisa de investimento.")
    calc("Marcador — mês com caixa negativo",
         lambda i, L, P: f"=IF({L}{D['caixa']}<0,{L}${R_IDX},9999)", "mk_neg", INT, cor=CINZA,
         total=None, obs="Auxiliar: o menor valor desta linha é o primeiro mês sem fundo.")
    calc("Marcador — mês com EBITDA positivo",
         lambda i, L, P: f"=IF({L}{D['ebitda']}>0,{L}${R_IDX},9999)", "mk_eb", INT, cor=CINZA,
         total=None)
    calc("Marcador — mês com FCL acumulado ainda negativo",
         lambda i, L, P: f"=IF({L}{D['fcl_ac']}<0,{L}${R_IDX},0)", "mk_pb", INT, cor=CINZA,
         total=None, obs="Auxiliar do payback: o maior valor, mais um, é o mês em que o "
                         "acumulado vira.")
    r += 1

    # ================================================== 7. CONFERENCIA
    secao(ws, r, "7. CONTAS DE APOIO E CONFERÊNCIA"); r += 1
    calc("CAC por assinante pagante conquistado",
         lambda i, L, P: f"=IFERROR({L}{D['p_cac']}/({L}{D['conv']}*$B${D['m_conv']}),0)",
         "cac_pag", BRL2, cor=CINZA, total="ultimo",
         obs="Quanto custa trazer um PAGANTE: o CAC por usuário dividido pela conversão.")
    calc("Receita por assinante, líquida de imposto, loja e pagamento",
         lambda i, L, P: f"={L}{D['preco']}*{L}{D['retencao']}", "arpu_liq", BRL2, cor=CINZA, total="ultimo")
    calc("Custo variável por assinante (nuvem, suporte, gratuitos rateados)",
         lambda i, L, P: (f"=IFERROR((-{L}{D['nuvem']}-{L}{D['sup']})/{L}{D['pag']},0)"),
         "cv_unit", BRL2, cor=CINZA, total="ultimo")
    calc("Sobra por assinante antes do custo fixo",
         lambda i, L, P: f"={L}{D['arpu_liq']}-{L}{D['cv_unit']}", "mc_unit", BRL2, bold=True, total="ultimo",
         obs="Se esta linha for negativa, cada assinante novo dá prejuízo e nenhum volume "
             "resolve. É a primeira coisa a conferir num preço.")
    calc("Assinantes necessários para o mês fechar no zero",
         lambda i, L, P: (f'=IF({L}{D["mc_unit"]}<=0,"",'
                          f'ROUNDUP(({L}{D["cf"]}+{L}{D["mkt_aq"]}*-1)/{L}{D["mc_unit"]},0))'),
         "n_eq", INT, bold=True, total="ultimo",
         obs="Custo fixo mais aquisição, dividido pela sobra por assinante. Vazio = nenhum "
             "volume fecha a esse preço.")
    calc("✔ Conferência: faturamento líquido − custo − EBITDA (deve ser zero)",
         lambda i, L, P: f"={L}{D['liquido']}+{L}{D['custo']}-{L}{D['ebitda']}", "chk1", BRL2,
         cor=CINZA)
    calc("✔ Conferência: base = pagantes + gratuitos (deve ser zero)",
         lambda i, L, P: f"={L}{D['base_fim']}-{L}{D['pag']}-{L}{D['free']}", "chk2", INT,
         cor=CINZA)
    r += 2

    # ---- tabela de faixas do Simples, usada pela linha de sugestao ----
    secao(ws, r, "TABELA DO SIMPLES NACIONAL — ANEXO V  (referência para a alíquota acima)", 8)
    r += 1
    for j, t in enumerate(["RBT12 a partir de", "Alíquota nominal", "Dedução (R$)"]):
        c = ws.cell(r, 1 + j, t)
        c.font = f(9, True, BRANCO); c.fill = F_HDR
        c.alignment = Alignment(horizontal="center", indent=1)
    r += 1
    FX_INI = r
    for de, nom, ded in FAIXAS:
        ws.cell(r, 1, de).number_format = BRL
        ws.cell(r, 2, nom).number_format = PCT
        ws.cell(r, 3, ded).number_format = BRL
        for cidx in (1, 2, 3):
            ws.cell(r, cidx).font = f(9, False, PRETO)
            ws.cell(r, cidx).border = BORDA
            ws.cell(r, cidx).alignment = Alignment(horizontal="center")
        r += 1
    FX_FIM = r - 1
    ws.cell(r, 1, f"Acima de R$ 4.800.000 sai do Simples: alíquota de Lucro Presumido").font = \
        f(9, False, CINZA, italic=True)
    ws.cell(r, 2, ALIQ_LP).number_format = PCT
    ws.cell(r, 2).font = f(9, True, PRETO)
    r_lp = r
    r += 2

    # agora que a tabela existe, preenche a linha de sugestao reservada
    rs = D["_sug_imp"]
    for i in range(1, N + 1):
        L = col(i)
        cc = ws[f"{L}{rs}"]
        cc.value = (f'=IF({L}{D["rbt12"]}<=0,$B${FX_INI},'
                    f'IF({L}{D["rbt12"]}>4800000,$B${r_lp},'
                    f'MAX($B${FX_INI},(INDEX($B${FX_INI}:$B${FX_FIM},'
                    f'MATCH({L}{D["rbt12"]},$A${FX_INI}:$A${FX_FIM},1))*{L}{D["rbt12"]}'
                    f'-INDEX($C${FX_INI}:$C${FX_FIM},'
                    f'MATCH({L}{D["rbt12"]},$A${FX_INI}:$A${FX_FIM},1)))/{L}{D["rbt12"]})))')
        cc.number_format = PCT2
        cc.font = f(9, False, CINZA)
    ct = ws[f"{L_TOT}{rs}"]
    ct.value = f"={L_FIM}{rs}"
    ct.number_format = PCT2
    ct.font = f(9, True, CINZA)
    ct.fill = F_TOT
    nota(ws, rs, "Alíquota efetiva do Anexo V pelo RBT12 do mês. Se o número que você digitou "
                 "acima estiver muito abaixo desta linha, o modelo está otimista no imposto.")
    D["sug_imp"] = rs

    ws.sheet_properties.tabColor = "1F7040"
    return ws, D
