# -*- coding: utf-8 -*-
"""Aba Pessoas: quadro de pessoal editavel linha a linha."""
import datetime as dt
from openpyxl.styles import Alignment
from openpyxl.comments import Comment
from openpyxl.worksheet.datavalidation import DataValidation
from common import *

D = lambda a, m: dt.datetime(a, m, 1)
FIM = D(2030, 12)

# Funcao | Tipo | Status | Qtd | Custo mensal | Encargos | Inicio | Fim | Classificacao | Nota
QUADRO = [
    ("Equipe PJ — desenvolvimento do MVP", "PJ", "Contratado", 1, 10500, 0.0,
     D(2026, 9), D(2026, 10), "Investimento",
     "Único custo de pessoal efetivamente comprometido. Brief: equipe PJ contratada em setembro "
     "para entregar o app até outubro, orçamento total de R$ 21.000 (2 meses × R$ 10.500). "
     "Classificado como Investimento porque é o desenvolvimento do ativo — amortizado em 60 meses na DRE."),

    ("Desenvolvimento — time base", "PJ", "Plano", 2, 9000, 0.0,
     D(2027, 1), FIM, "Despesa",
     "Manutenção e evolução do app após o lançamento. Sugestão a validar."),
    ("Desenvolvimento — ampliação 2028 (adicional)", "PJ", "Plano", 3, 9500, 0.0,
     D(2028, 1), FIM, "Despesa", "Soma-se ao time base — não o substitui."),
    ("Desenvolvimento — ampliação 2029 (adicional)", "PJ", "Plano", 5, 10000, 0.0,
     D(2029, 1), FIM, "Despesa", "Soma-se aos anteriores."),
    ("Desenvolvimento — ampliação 2030 (adicional)", "PJ", "Plano", 6, 10500, 0.0,
     D(2030, 1), FIM, "Despesa", "Soma-se aos anteriores."),

    ("Suporte ao cliente — time base", "PJ", "Plano", 1, 3500, 0.0,
     D(2027, 1), FIM, "Despesa",
     "Atendimento humano. As ferramentas de suporte (chatbot, help desk) estão em Premissas, "
     "no custo variável por pagante."),
    ("Suporte — ampliação 2028 (adicional)", "PJ", "Plano", 3, 3500, 0.0,
     D(2028, 1), FIM, "Despesa", ""),
    ("Suporte — ampliação 2029 (adicional)", "PJ", "Plano", 6, 3800, 0.0,
     D(2029, 1), FIM, "Despesa", ""),
    ("Suporte — ampliação 2030 (adicional)", "PJ", "Plano", 8, 4000, 0.0,
     D(2030, 1), FIM, "Despesa", ""),

    ("Growth e marketing", "PJ", "Plano", 1, 6000, 0.0,
     D(2027, 1), FIM, "Despesa",
     "Pessoa dedicada a aquisição e retenção. A mídia paga está em Premissas (CAC e marketing "
     "recorrente), não aqui."),
    ("Growth — ampliação 2028 (adicional)", "PJ", "Plano", 2, 7000, 0.0,
     D(2028, 1), FIM, "Despesa", ""),
    ("Growth — ampliação 2029 (adicional)", "PJ", "Plano", 3, 8000, 0.0,
     D(2029, 1), FIM, "Despesa", ""),

    ("Dados e infraestrutura", "PJ", "Plano", 1, 10000, 0.0,
     D(2028, 1), FIM, "Despesa",
     "Confiabilidade da ingestão de sinais vitais e custo de nuvem sob controle."),
    ("Dados e infraestrutura — ampliação 2029 (adicional)", "PJ", "Plano", 2, 11000, 0.0,
     D(2029, 1), FIM, "Despesa", ""),

    ("Produto e gestão", "PJ", "Plano", 1, 12000, 0.0,
     D(2028, 1), FIM, "Despesa", ""),
    ("Produto e gestão — ampliação 2030 (adicional)", "PJ", "Plano", 2, 13000, 0.0,
     D(2030, 1), FIM, "Despesa", ""),

    ("Pró-labore dos sócios (2027–2028)", "Sócio", "Plano", 3, 5000, 0.0,
     D(2027, 1), D(2028, 12), "Despesa",
     "Retirada dos sócios ativos. Zero em 2026. Conta para o Fator R — ver quadro no fim desta aba."),
    ("Pró-labore dos sócios (2029–2030)", "Sócio", "Plano", 4, 8000, 0.0,
     D(2029, 1), FIM, "Despesa", "Substitui a linha anterior, que se encerra em dez/2028."),

    ("(exemplo) Contratação CLT", "CLT", "Plano", 0, 8000, 0.70,
     D(2028, 1), FIM, "Despesa",
     "Linha de exemplo, com quantidade zero. Encargos de 70% cobrem INSS patronal, FGTS, férias, "
     "13º e provisões. Contratar em CLT eleva o Fator R e pode levar a empresa do Anexo V ao "
     "Anexo III — compare o custo da folha com a economia tributária no quadro abaixo."),
]
VAZIAS = 8
PE = {}


def build(wb, P):
    ws = wb.create_sheet("Pessoas")
    titulo(ws, "PESSOAS — QUADRO DE PESSOAL",
           "Uma linha por papel. Edite quantidade, custo, início e fim livremente. "
           "Status 'Contratado' entra sempre; 'Plano' só entra se o controle 1 de Premissas estiver em 1. "
           "Todo o custo de pessoal do modelo sai daqui.")
    largura_padrao(ws, col_a=44, col_mes_w=12)
    for col, w in zip("BCDEFGHIJ", (11, 12, 8, 14, 11, 15, 11, 11, 14)):
        ws.column_dimensions[col].width = w

    r = 4
    secao(ws, r, "QUADRO — edite livremente, inclusive as linhas em branco no fim", 10); r += 1
    hdr = ["Função / papel", "Tipo", "Status", "Qtd", "Custo mensal\npor pessoa",
           "Encargos\n(% sobre CLT)", "Custo total\nmensal", "Início", "Fim", "Classificação"]
    for j, h in enumerate(hdr):
        c = ws.cell(r, 1 + j, h)
        c.font = f(9, True, BRANCO); c.fill = FILL_HEADER
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[r].height = 30
    r += 1
    PE["ini"] = r

    linhas = list(QUADRO) + [("", "PJ", "Plano", 0, 0, 0.0, D(2027, 1), FIM, "Despesa", "")
                             for _ in range(VAZIAS)]
    for (fn, tipo, st, qtd, custo, enc, ini, fim, cls, obs) in linhas:
        contratado = (st == "Contratado")
        ws.cell(r, 1, fn)
        ws.cell(r, 2, tipo)
        ws.cell(r, 3, st)
        ws.cell(r, 4, qtd).number_format = INT
        ws.cell(r, 5, custo).number_format = BRL
        ws.cell(r, 6, enc).number_format = PCT
        ws.cell(r, 7, f'=$D{r}*$E{r}*(1+$F{r})').number_format = BRL
        ws.cell(r, 8, ini).number_format = MES
        ws.cell(r, 9, fim).number_format = MES
        ws.cell(r, 10, cls)
        for j in range(1, 11):
            c = ws.cell(r, j)
            c.border = BORDA_FINA
            c.alignment = Alignment(horizontal="left" if j == 1 else "center",
                                    indent=1 if j == 1 else 0)
            if j == 7:
                c.font = f(9, True, PRETO)
            else:
                c.font = f(9, contratado, AZUL_INPUT)
            if contratado:
                c.fill = FILL_OK
            elif st == "Plano":
                c.fill = FILL_SUB if fn else PatternFill()
        if obs:
            ws.cell(r, 1).comment = Comment(obs, "Modelo Oppa", height=140, width=380)
        r += 1
    PE["fim"] = r - 1
    i, fi = PE["ini"], PE["fim"]

    ws.cell(r, 1, "TOTAL DO QUADRO (custo mensal se todas as linhas estivessem ativas)").font = f(10, True)
    c = ws.cell(r, 7, f'=SUM(G{i}:G{fi})'); c.number_format = BRL
    c.font = f(10, True); c.fill = FILL_TOTAL
    r += 1
    nota(ws, f"A{r}", "Azul-escuro/verde = já contratado · azul-claro = plano sugerido, a validar. "
                      "Coluna 'Custo total mensal' é calculada.")
    r += 2

    # --- validacoes ------------------------------------------------------
    for col, opcoes in (("B", '"PJ,CLT,Sócio"'),
                        ("C", '"Contratado,Plano"'),
                        ("J", '"Despesa,Investimento"')):
        dv = DataValidation(type="list", formula1=opcoes, allow_blank=True, showDropDown=False)
        ws.add_data_validation(dv)
        dv.add(f"{col}{i}:{col}{fi}")

    # --- bloco mensal ----------------------------------------------------
    cabecalho_tempo(ws, linha_data=r, congelar=f"B{r+5}", anos_row=P["anos"])
    PE["tempo"] = r
    t = r
    r += 4
    secao(ws, r, "CUSTO MENSAL DE PESSOAL", COL_PCT); r += 1

    def ativo(L):
        return (f'($H${i}:$H${fi}<={L}${t})*($I${i}:$I${fi}>={L}${t})'
                f'*(($C${i}:$C${fi}="Contratado")'
                f'+($C${i}:$C${fi}="Plano")*Premissas!$B${P["plano"]})')

    def bloco(lbl, cond, key, fmt=BRL, bold=False, cor=PRETO, col_soma="$G"):
        nonlocal r
        rotulo(ws, r, lbl, 0 if bold else 1, bold=bold, color=cor)
        preencher_linha(ws, r, lambda idx, L: (
            f'=SUMPRODUCT({cond(L)}{ativo(L)}*{col_soma}${i}:{col_soma}${fi})'),
            fmt, bold=bold, color=cor)
        PE[key] = r
        r += 1

    bloco("Pessoas ativas no mês (headcount)",
          lambda L: '', "headcount", INT, col_soma="$D")
    ws[f"{L_TOT}{PE['headcount']}"] = f'={L_FIM}{PE["headcount"]}'
    ws[f"{L_TOT}{PE['headcount']}"].number_format = INT
    nota(ws, f"{L_PCT}{PE['headcount']}", "Quadro em dez/2030, não soma.")

    bloco("Equipe PJ (despesa)",
          lambda L: f'($B${i}:$B${fi}="PJ")*($J${i}:$J${fi}="Despesa")*', "pj")
    bloco("Equipe CLT — salários e encargos (despesa)",
          lambda L: f'($B${i}:$B${fi}="CLT")*($J${i}:$J${fi}="Despesa")*', "clt")
    bloco("Pró-labore dos sócios (despesa)",
          lambda L: f'($B${i}:$B${fi}="Sócio")*($J${i}:$J${fi}="Despesa")*', "socio")
    bloco("TOTAL DE PESSOAL — DESPESA",
          lambda L: f'($J${i}:$J${fi}="Despesa")*', "despesa", bold=True)
    for k in range(1, N_MESES + 1):
        ws[f"{col_mes(k)}{PE['despesa']}"].fill = FILL_TOTAL
    bloco("Pessoal capitalizado — INVESTIMENTO",
          lambda L: f'($J${i}:$J${fi}="Investimento")*', "invest", cor=VERDE_LINK)
    nota(ws, f"{L_PCT}{PE['invest']}", "Vai para a aba Investimentos e é amortizado em 60 meses.")
    bloco("Folha que conta para o Fator R (CLT + pró-labore)",
          lambda L: f'(($B${i}:$B${fi}="CLT")+($B${i}:$B${fi}="Sócio"))*', "folha_fr", cor=CINZA)
    nota(ws, f"{L_PCT}{PE['folha_fr']}", "Pagamento a PJ não entra no Fator R.")
    r += 1

    PE["prox"] = r
    return ws, PE


def add_fator_r(ws, P, PE, FAT):
    """Bloco anual com Fator R. Chamado depois que a aba Faturamento existe."""
    r = PE["prox"]
    t = PE["tempo"]
    i, fi = PE["ini"], PE["fim"]
    ANO_FAT = FAT

    # --- resumo anual + Fator R -------------------------------------------
    secao(ws, r, "RESUMO ANUAL E FATOR R", 10); r += 1
    ahdr = r
    ws.cell(r, 1, "").font = f(9)
    for j, a in enumerate([2026, 2027, 2028, 2029, 2030]):
        c = ws.cell(r, 2 + j, a)
        c.font = f(10, True, BRANCO); c.fill = FILL_HEADER
        c.alignment = Alignment(horizontal="center"); c.number_format = '0'
    r += 1
    PE["anual_hdr"] = ahdr

    def anual(lbl, src, fmt=BRL, modo="soma", bold=False):
        nonlocal r
        rotulo(ws, r, lbl, 0 if bold else 1, bold=bold)
        for j in range(5):
            L = gcl(2 + j)
            if modo == "soma":
                fx = (f'=SUMIFS($B${src}:${L_FIM}${src},'
                      f'$B${t+2}:${L_FIM}${t+2},{L}${ahdr})')
            else:  # maximo do ano
                fx = (f'=MAX(IF($B${t+2}:${L_FIM}${t+2}={L}${ahdr},$B${src}:${L_FIM}${src}))')
            c = ws.cell(r, 2 + j, fx)
            c.number_format = fmt; c.font = f(9, bold)
            c.border = BORDA_FINA
            if bold:
                c.fill = FILL_SUB
        PE[f"a_{src}"] = r
        r += 1
        return r - 1

    r_desp = anual("Custo total de pessoal (despesa)", PE["despesa"], BRL, "soma", True)
    r_inv = anual("Pessoal capitalizado (investimento)", PE["invest"])
    r_folha = anual("Folha do Fator R (CLT + pró-labore)", PE["folha_fr"])

    rotulo(ws, r, "Pessoas ativas em dezembro", 1)
    for j in range(5):
        L = gcl(2 + j)
        col_dez = col_am(2026 + j, 12)
        c = ws.cell(r, 2 + j, f'={col_dez}{PE["headcount"]}')
        c.number_format = INT; c.font = f(9); c.border = BORDA_FINA
    r_hc = r; r += 1

    rotulo(ws, r, "Receita bruta do ano (cenário ativo)", 1)
    for j in range(5):
        L = gcl(2 + j)
        c = ws.cell(r, 2 + j,
            f"=SUMIFS(Faturamento!$B${ANO_FAT['bruta']}:${L_FIM}${ANO_FAT['bruta']},"
            f"Faturamento!$B$6:${L_FIM}$6,{L}${ahdr})")
        c.number_format = BRL; c.font = f(9, False, VERDE_LINK); c.border = BORDA_FINA
    r_rec = r; r += 1

    rotulo(ws, r, "FATOR R  (folha ÷ receita bruta)", 0, bold=True)
    for j in range(5):
        L = gcl(2 + j)
        c = ws.cell(r, 2 + j, f'=IFERROR({L}{r_folha}/{L}{r_rec},0)')
        c.number_format = PCT; c.font = f(10, True); c.border = BORDA_FINA
        c.fill = FILL_TOTAL
    r_fr = r; r += 1

    rotulo(ws, r, "Anexo aplicável", 1)
    for j in range(5):
        L = gcl(2 + j)
        c = ws.cell(r, 2 + j, f'=IF({L}{r_fr}>=0.28,"Anexo III","Anexo V")')
        c.font = f(10, True, "C00000"); c.border = BORDA_FINA
        c.alignment = Alignment(horizontal="center")
    r += 1
    rotulo(ws, r, "Regime escolhido em Premissas", 1, italic=True)
    for j in range(5):
        L = gcl(2 + j)
        c = ws.cell(r, 2 + j,
            f'=CHOOSE(Premissas!$B${P["regime"]},"Anexo III","Anexo V","Lucro Presumido")')
        c.font = f(9, False, VERDE_LINK); c.border = BORDA_FINA
        c.alignment = Alignment(horizontal="center")
    r_esc = r; r += 1
    rotulo(ws, r, "Escolha compatível com o Fator R?", 0, bold=True)
    for j in range(5):
        L = gcl(2 + j)
        c = ws.cell(r, 2 + j,
            f'=IF(Premissas!$B${P["regime"]}=3,"n/a",'
            f'IF({L}{r_esc}={L}{r_fr+1},"✔ SIM","⚠ NÃO"))')
        c.font = f(10, True); c.border = BORDA_FINA
        c.alignment = Alignment(horizontal="center")
    r += 1
    rotulo(ws, r, "Custo de pessoal ÷ receita bruta", 1, italic=True)
    for j in range(5):
        L = gcl(2 + j)
        c = ws.cell(r, 2 + j, f'=IFERROR({L}{r_desp}/{L}{r_rec},0)')
        c.number_format = PCT; c.font = f(9, False, CINZA); c.border = BORDA_FINA
    r += 2

    for t_, txt in [
        ("Como o Fator R funciona",
         "Fator R = (folha CLT + pró-labore dos últimos 12 meses) ÷ receita bruta dos últimos 12 meses. "
         "Igual ou acima de 28%, a empresa fica no Anexo III do Simples (alíquota efetiva a partir de "
         "6%). Abaixo, cai no Anexo V (a partir de 15,5%). Pagamento a equipe PJ NÃO entra na conta."),
        ("Por que isso importa para a Oppa",
         "Com o quadro atual — equipe PJ e pró-labore modesto — o Fator R fica bem abaixo de 28%, o que "
         "leva ao Anexo V. Migrar parte do time para CLT ou aumentar o pró-labore eleva o custo de "
         "folha, mas pode reduzir mais imposto do que custa. Use a linha de exemplo CLT e a aba "
         "Tributos para comparar os dois efeitos antes de decidir."),
        ("O que está de fato contratado",
         "Só a equipe PJ do MVP, de setembro a outubro de 2026, R$ 21.000 no total. Todo o resto é "
         "plano — desligue no controle 1 de Premissas para ver o modelo sem nenhuma contratação futura."),
    ]:
        c = ws.cell(r, 1, "▸  " + t_); c.font = f(10, True, "1F3864")
        c.alignment = Alignment(indent=1); r += 1
        c = ws.cell(r, 1, txt); c.font = f(9, False, CINZA)
        c.alignment = Alignment(wrap_text=True, vertical="top", indent=2)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=10)
        ws.row_dimensions[r].height = 44
        r += 1

    return r
