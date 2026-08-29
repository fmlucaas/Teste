# -*- coding: utf-8 -*-
"""Estilos, formatos e utilitarios comuns do modelo Oppa."""
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as gcl
import datetime as dt

FONT = "Arial"

# ---- Formatos numericos -------------------------------------------------
BRL   = '"R$ "#,##0;("R$ "#,##0);"–"'
BRL2  = '"R$ "#,##0.00;("R$ "#,##0.00);"–"'
INT   = '#,##0;(#,##0);"–"'
PCT   = '0.0%;(0.0%);"–"'
PCT2  = '0.00%;(0.00%);"–"'
MULT  = '0.0"x";(0.0"x");"–"'
MES   = 'mmm/yy'
DATA  = 'dd/mm/yyyy'
TXT   = '@'

# ---- Cores --------------------------------------------------------------
AZUL_INPUT   = "0000FF"   # entrada digitada (hardcode)
PRETO        = "000000"   # formula na propria aba
VERDE_LINK   = "008000"   # link para outra aba
CINZA        = "595959"
BRANCO       = "FFFFFF"

FILL_TITULO  = PatternFill("solid", fgColor="1F3864")
FILL_SECAO   = PatternFill("solid", fgColor="2E5C8A")
FILL_SUB     = PatternFill("solid", fgColor="D9E2F3")
FILL_TOTAL   = PatternFill("solid", fgColor="BDD7EE")
FILL_INPUT   = PatternFill("solid", fgColor="FFFF00")   # premissa-chave a preencher
FILL_ALERTA  = PatternFill("solid", fgColor="FCE4D6")
FILL_OK      = PatternFill("solid", fgColor="E2EFDA")
FILL_ZEBRA   = PatternFill("solid", fgColor="F2F2F2")
FILL_HEADER  = PatternFill("solid", fgColor="404040")

THIN = Side(style="thin", color="BFBFBF")
MED  = Side(style="medium", color="1F3864")
BORDA_FINA = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
BORDA_TOPO = Border(top=Side(style="thin", color="404040"))


def f(size=10, bold=False, color=PRETO, italic=False):
    return Font(name=FONT, size=size, bold=bold, color=color, italic=italic)


def titulo(ws, texto, subtitulo=None, largura=62):
    """Faixa de titulo padrao nas linhas 1-2."""
    ws["A1"] = texto
    ws["A1"].font = f(14, True, BRANCO)
    ws["A1"].fill = FILL_TITULO
    ws["A1"].alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[1].height = 26
    for c in range(1, largura + 1):
        ws.cell(1, c).fill = FILL_TITULO
    if subtitulo:
        ws["A2"] = subtitulo
        ws["A2"].font = f(9, False, CINZA, italic=True)
        ws["A2"].alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[2].height = 16


def secao(ws, linha, texto, largura=62):
    """Linha de secao destacada."""
    ws.cell(linha, 1, texto).font = f(10, True, BRANCO)
    for c in range(1, largura + 1):
        ws.cell(linha, c).fill = FILL_SECAO
    ws.row_dimensions[linha].height = 18


def rotulo(ws, linha, texto, nivel=0, bold=False, color=PRETO, italic=False):
    c = ws.cell(linha, 1, texto)
    c.font = f(10, bold, color, italic)
    c.alignment = Alignment(indent=nivel, vertical="center")
    return c


def nota(ws, cel, texto):
    """Nota curta em celula, cinza italico."""
    ws[cel] = texto
    ws[cel].font = f(8, False, CINZA, italic=True)
    ws[cel].alignment = Alignment(vertical="center", wrap_text=False)


# ---- Linha do tempo -----------------------------------------------------
INICIO = dt.date(2026, 1, 1)
N_MESES = 60
COL0 = 2                       # coluna B = mes 1
COL_TOT = COL0 + N_MESES       # 62 -> BJ
COL_PCT = COL_TOT + 1          # 63 -> BK

def col_mes(i):
    """Letra da coluna do mes i (1-based)."""
    return gcl(COL0 + i - 1)

def idx_mes(ano, mes):
    """Indice 1-based do mes (ano, mes) na linha do tempo."""
    return (ano - 2026) * 12 + mes

def col_am(ano, mes):
    return col_mes(idx_mes(ano, mes))

def data_mes(i):
    ano = 2026 + (i - 1) // 12
    mes = (i - 1) % 12 + 1
    return dt.datetime(ano, mes, 1)

L_TOT = gcl(COL_TOT)
L_PCT = gcl(COL_PCT)
L_FIM = gcl(COL0 + N_MESES - 1)   # BI


def cabecalho_tempo(ws, linha_data=4, congelar="B7", anos_row=None):
    """Escreve datas (linha 4), indice (5), ano (6) e indice do ano (7)."""
    ws.cell(linha_data, 1, "Linha do tempo").font = f(9, True, CINZA)
    for i in range(1, N_MESES + 1):
        L = col_mes(i)
        c = ws[f"{L}{linha_data}"]
        c.value = data_mes(i)
        c.number_format = MES
        c.font = f(9, True, BRANCO)
        c.fill = FILL_HEADER
        c.alignment = Alignment(horizontal="center")
        ci = ws[f"{L}{linha_data+1}"]
        ci.value = i
        ci.number_format = INT
        ci.font = f(8, False, CINZA)
        ci.alignment = Alignment(horizontal="center")
        ca = ws[f"{L}{linha_data+2}"]
        ca.value = f"=YEAR({L}{linha_data})"
        ca.number_format = '0'
        ca.font = f(8, False, CINZA)
        ca.alignment = Alignment(horizontal="center")
        if anos_row:
            cy = ws[f"{L}{linha_data+3}"]
            cy.value = (f"=MATCH({L}{linha_data+2},"
                        f"Premissas!$C${anos_row}:$G${anos_row},0)")
            cy.number_format = '0'
            cy.font = f(8, False, "BFBFBF")
            cy.alignment = Alignment(horizontal="center")
    for L, txt in ((L_TOT, "TOTAL"), (L_PCT, "% / Nota")):
        c = ws[f"{L}{linha_data}"]
        c.value = txt
        c.font = f(9, True, BRANCO)
        c.fill = FILL_TITULO
        c.alignment = Alignment(horizontal="center")
    ws.cell(linha_data + 1, 1, "Mês nº").font = f(8, False, CINZA)
    ws.cell(linha_data + 2, 1, "Ano").font = f(8, False, CINZA)
    if anos_row:
        ws.cell(linha_data + 3, 1, "Índice do ano (auxiliar)").font = f(8, False, "BFBFBF")
    ws.freeze_panes = congelar


def largura_padrao(ws, col_a=46, col_mes_w=11, col_tot_w=15):
    ws.column_dimensions["A"].width = col_a
    for i in range(1, N_MESES + 1):
        ws.column_dimensions[col_mes(i)].width = col_mes_w
    ws.column_dimensions[L_TOT].width = col_tot_w
    ws.column_dimensions[L_PCT].width = 34


def preencher_linha(ws, linha, gerador, fmt=BRL, color=PRETO, bold=False,
                    total=True, fmt_total=None, pct_base=None):
    """Escreve B..BI com gerador(i, L) -> valor/formula. Total em BJ."""
    for i in range(1, N_MESES + 1):
        L = col_mes(i)
        v = gerador(i, L)
        if v is None:
            continue
        c = ws[f"{L}{linha}"]
        c.value = v
        c.number_format = fmt
        c.font = f(9, bold, color)
    if total:
        c = ws[f"{L_TOT}{linha}"]
        c.value = f"=SUM(B{linha}:{L_FIM}{linha})"
        c.number_format = fmt_total or fmt
        c.font = f(9, True, PRETO)
        c.fill = FILL_TOTAL
    if pct_base:
        c = ws[f"{L_PCT}{linha}"]
        c.value = f"=IFERROR({L_TOT}{linha}/{L_TOT}${pct_base},0)"
        c.number_format = PCT
        c.font = f(9, False, CINZA)
