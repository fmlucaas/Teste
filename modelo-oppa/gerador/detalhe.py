# -*- coding: utf-8 -*-
"""Abas de detalhe do cenario selecionado: Usuarios, Faturamento, Analise Fluxo, DRE."""
from openpyxl.styles import Alignment
from common import *
import engine as EN

TEMPO, LIDX, LANO, LIDXA = 4, 5, 6, 7


def mk_sel(P, BLOCOS):
    """Fabricas de formulas que apontam para o cenario selecionado."""
    def sel(off, L):
        alvos = ",".join(f"'Cenários'!{L}{S+off}" for S in BLOCOS)
        return f'CHOOSE(Premissas!$B${P["cenario"]},{alvos})'

    def yr(key, L, ano_row=None):
        r = P[key]
        return f'INDEX(Premissas!$C${r}:$G${r},{L}${LIDXA})'

    def yr_cen(prefixo, L, ano_row=None):
        partes = ",".join(yr(f"{prefixo}_{s}", L, ano_row) for s in (1, 2, 3))
        return f'CHOOSE(Premissas!$B${P["cenario"]},{partes})'

    return sel, yr, yr_cen


def _sub(ws, r):
    ws.cell(r, 1).font = f(10, True, BRANCO)
    for c in range(1, COL_PCT + 1):
        ws.cell(r, c).fill = FILL_SUB
        ws.cell(r, c).font = f(10, True, "1F3864")


# =====================================================================
def build_usuarios(wb, P, BLOCOS):
    ws = wb.create_sheet("Usuários")
    sel, yr, yr_cen = mk_sel(P, BLOCOS)
    titulo(ws, "USUÁRIOS E INDICADORES DE UNIDADE",
           "Funil de usuários e métricas de unidade do cenário selecionado em Premissas.")
    largura_padrao(ws, col_a=46)
    cabecalho_tempo(ws, linha_data=TEMPO, congelar="B9", anos_row=P["anos"])
    ws["A2"] = ('=" Cenário ativo: "&Premissas!B' + str(P["cenario_nome"]) +
                '&"  ·  todas as linhas recalculam ao trocar o cenário."')
    ws["A2"].font = f(9, True, VERDE_LINK)

    r = 8
    secao(ws, r, "FUNIL DE USUÁRIOS", COL_PCT); r += 1
    U = {}
    linhas = [
        ("Novos usuários captados", 1, INT, "novos", False),
        ("Base de usuários — início do mês", 2, INT, "base_ini", False),
        ("(−) Usuários perdidos (churn)", 3, INT, "churn", False),
        ("Base de usuários — fim do mês", 4, INT, "base", True),
        ("Taxa de conversão", 5, PCT, "conv", False),
        ("Usuários pagantes", 6, INT, "pag", True),
        ("Pagantes — Plano Básico", 7, INT, "pag_b", False),
        ("Pagantes — Plano Intermediário", 8, INT, "pag_i", False),
        ("Pagantes — Plano Premium", 9, INT, "pag_p", False),
    ]
    for lbl, off, fmt, key, bold in linhas:
        rotulo(ws, r, lbl, 0 if bold else 1, bold=bold)
        preencher_linha(ws, r, lambda i, L, o=off: f'={sel(o, L)}', fmt, bold=bold,
                        total=False, color=VERDE_LINK if not bold else PRETO)
        c = ws[f"{L_TOT}{r}"]
        c.value = (f"=SUM(B{r}:{L_FIM}{r})" if off in (1, 3) else f"={L_FIM}{r}")
        c.number_format = fmt; c.font = f(9, True); c.fill = FILL_TOTAL
        if off not in (1, 3):
            nota(ws, f"{L_PCT}{r}", "Saldo em dez/2030.")
        U[key] = r; r += 1
    rotulo(ws, r, "Usuários gratuitos", 1)
    preencher_linha(ws, r, lambda i, L: f'={L}{U["base"]}-{L}{U["pag"]}', INT, total=False)
    ws[f"{L_TOT}{r}"] = f"={L_FIM}{r}"
    ws[f"{L_TOT}{r}"].number_format = INT
    ws[f"{L_TOT}{r}"].font = f(9, True); ws[f"{L_TOT}{r}"].fill = FILL_TOTAL
    U["free"] = r; r += 2

    secao(ws, r, "INDICADORES DE UNIDADE", COL_PCT); r += 1
    rotulo(ws, r, "Investimento em mídia de aquisição", 1)
    preencher_linha(ws, r, lambda i, L:
        f'={L}{U["novos"]}*(1-{yr("organico", L)})*{yr_cen("cac", L)}', BRL)
    U["midia"] = r; r += 1
    rotulo(ws, r, "CAC — custo por usuário adquirido", 1)
    preencher_linha(ws, r, lambda i, L: f'=IFERROR({L}{U["midia"]}/{L}{U["novos"]},0)',
                    BRL2, total=False)
    ws[f"{L_TOT}{r}"] = f'=IFERROR({L_TOT}{U["midia"]}/{L_TOT}{U["novos"]},0)'
    ws[f"{L_TOT}{r}"].number_format = BRL2; ws[f"{L_TOT}{r}"].font = f(9, True)
    ws[f"{L_TOT}{r}"].fill = FILL_TOTAL
    nota(ws, f"{L_PCT}{r}", "Média do período.")
    U["cac"] = r; r += 1
    rotulo(ws, r, "CAC por usuário PAGANTE conquistado", 1, bold=True)
    preencher_linha(ws, r, lambda i, L:
        f'=IFERROR({L}{U["midia"]}/({L}{U["novos"]}*{L}{U["conv"]}),0)', BRL2, bold=True, total=False)
    nota(ws, f"{L_PCT}{r}", "É este o CAC que a receita de assinatura precisa pagar.")
    U["cac_pag"] = r; r += 1
    rotulo(ws, r, "ARPPU — receita média por assinante pagante", 1)
    preencher_linha(ws, r, lambda i, L:
        f'=IFERROR({sel(13, L)}/{L}{U["pag"]},0)', BRL2, total=False)
    U["arppu"] = r; r += 1
    rotulo(ws, r, "ARPU — receita média por usuário da base", 1)
    preencher_linha(ws, r, lambda i, L:
        f'=IFERROR({sel(16, L)}/{L}{U["base"]},0)', BRL2, total=False)
    U["arpu"] = r; r += 1
    rotulo(ws, r, "Margem de contribuição por pagante (R$/mês)", 1)
    preencher_linha(ws, r, lambda i, L:
        f'=IFERROR(({sel(22, L)}+{sel(23, L)}+{sel(24, L)})/{L}{U["pag"]},0)', BRL2, total=False)
    nota(ws, f"{L_PCT}{r}", "Receita líquida menos nuvem e suporte, por assinante.")
    U["mc"] = r; r += 1
    rotulo(ws, r, "LTV — valor do assinante ao longo da vida", 1, bold=True)
    preencher_linha(ws, r, lambda i, L:
        f'=IFERROR({L}{U["mc"]}/{yr_cen("churn", L)},0)', BRL2, bold=True, total=False)
    nota(ws, f"{L_PCT}{r}", "Margem de contribuição ÷ churn mensal.")
    U["ltv"] = r; r += 1
    rotulo(ws, r, "LTV ÷ CAC por pagante", 1, bold=True)
    preencher_linha(ws, r, lambda i, L:
        f'=IFERROR({L}{U["ltv"]}/{L}{U["cac_pag"]},0)', MULT, bold=True, total=False)
    nota(ws, f"{L_PCT}{r}", "Referência de mercado: ≥ 3,0x é saudável; < 1,0x destrói valor.")
    U["ltv_cac"] = r; r += 1
    rotulo(ws, r, "Payback do CAC (meses)", 1)
    preencher_linha(ws, r, lambda i, L:
        f'=IFERROR({L}{U["cac_pag"]}/{L}{U["mc"]},0)', '0.0;(0.0);"–"', total=False)
    nota(ws, f"{L_PCT}{r}", "Meses de assinatura para recuperar o custo de aquisição. Ideal: < 12.")
    U["pb_cac"] = r; r += 2

    secao(ws, r, "CAIXA E FÔLEGO", COL_PCT); r += 1
    rotulo(ws, r, "Caixa acumulado", 1, bold=True)
    preencher_linha(ws, r, lambda i, L: f'={sel(37, L)}', BRL, bold=True, total=False)
    U["caixa"] = r; r += 1
    rotulo(ws, r, "Queima de caixa do mês (burn)", 1)
    preencher_linha(ws, r, lambda i, L: f'=MAX(0,-{sel(30, L)})', BRL, total=False)
    U["burn"] = r; r += 1
    rotulo(ws, r, "Runway (meses de caixa restantes)", 1, bold=True)
    def _runway(i, L):
        j = max(1, i - 2)
        ini = col_mes(j)
        return (f'=IF({L}{U["caixa"]}<=0,0,'
                f'IF(AVERAGE({ini}{U["burn"]}:{L}{U["burn"]})<=0,999,'
                f'{L}{U["caixa"]}/AVERAGE({ini}{U["burn"]}:{L}{U["burn"]})))')
    preencher_linha(ws, r, _runway, '0;(0);"–"', bold=True, total=False)
    nota(ws, f"{L_PCT}{r}", "Caixa ÷ queima média dos últimos 3 meses. 999 = sem queima.")
    U["runway"] = r
    return ws, U


# =====================================================================
def build_faturamento(wb, P, BLOCOS):
    ws = wb.create_sheet("Faturamento")
    sel, yr, yr_cen = mk_sel(P, BLOCOS)
    titulo(ws, "FATURAMENTO PROJETADO",
           "Receita por linha de negócio e deduções, no cenário selecionado.")
    largura_padrao(ws, col_a=46)
    cabecalho_tempo(ws, linha_data=TEMPO, congelar="B9", anos_row=P["anos"])
    ws["A2"] = '=" Cenário ativo: "&Premissas!B' + str(P["cenario_nome"])
    ws["A2"].font = f(9, True, VERDE_LINK)

    F = {}
    r = 8
    secao(ws, r, "RECEITA BRUTA", COL_PCT); r += 1
    for lbl, off, key in [
        ("Assinaturas — Plano Básico", 10, "b"),
        ("Assinaturas — Plano Intermediário", 11, "i"),
        ("Assinaturas — Plano Premium", 12, "p"),
        ("Marketplace — comissões sobre dispositivos", 14, "mp"),
        ("B2B — Boston Scientific (se acordo ativo)", 15, "b2b"),
    ]:
        rotulo(ws, r, lbl, 1)
        preencher_linha(ws, r, lambda i, L, o=off: f'={sel(o, L)}', BRL, color=VERDE_LINK)
        F[key] = r; r += 1
    rotulo(ws, r, "RECEITA BRUTA TOTAL", 0, bold=True)
    preencher_linha(ws, r, lambda i, L: f'=SUM({L}{F["b"]}:{L}{F["b2b"]})', BRL, bold=True)
    for i in range(1, N_MESES + 1):
        ws[f"{col_mes(i)}{r}"].fill = FILL_TOTAL
    F["bruta"] = r
    for k in ("b", "i", "p", "mp", "b2b"):
        ws[f"{L_PCT}{F[k]}"] = f'=IFERROR({L_TOT}{F[k]}/{L_TOT}${F["bruta"]},0)'
        ws[f"{L_PCT}{F[k]}"].number_format = PCT
        ws[f"{L_PCT}{F[k]}"].font = f(9, False, CINZA)
    r += 2

    secao(ws, r, "DEDUÇÕES DA RECEITA", COL_PCT); r += 1
    for lbl, off, key in [
        ("(−) Impostos sobre a receita", 19, "imp"),
        ("(−) Comissão das lojas de aplicativos", 20, "loja"),
        ("(−) Taxas de meios de pagamento", 21, "pgto"),
    ]:
        rotulo(ws, r, lbl, 1)
        preencher_linha(ws, r, lambda i, L, o=off: f'={sel(o, L)}', BRL, color=VERDE_LINK,
                        pct_base=F["bruta"])
        F[key] = r; r += 1
    rotulo(ws, r, "RECEITA LÍQUIDA", 0, bold=True)
    preencher_linha(ws, r, lambda i, L: f'={L}{F["bruta"]}+SUM({L}{F["imp"]}:{L}{F["pgto"]})',
                    BRL, bold=True, pct_base=F["bruta"])
    for i in range(1, N_MESES + 1):
        ws[f"{col_mes(i)}{r}"].fill = FILL_TOTAL
    F["liquida"] = r; r += 1
    rotulo(ws, r, "Alíquota efetiva de impostos", 1, italic=True)
    preencher_linha(ws, r, lambda i, L: f'={sel(18, L)}', PCT2, color=CINZA, total=False)
    nota(ws, f"{L_PCT}{r}", "Simples Nacional progressivo; migra para Lucro Presumido acima de R$ 4,8 mi.")
    F["aliq"] = r; r += 2

    secao(ws, r, "RESUMO ANUAL", COL_PCT); r += 1
    hdr = r
    ws.cell(r, 1, "Ano").font = f(9, True, BRANCO)
    for j, a in enumerate([2026, 2027, 2028, 2029, 2030]):
        c = ws.cell(r, 2 + j, a)
        c.font = f(10, True, BRANCO); c.fill = FILL_HEADER
        c.alignment = Alignment(horizontal="center"); c.number_format = '0'
    c = ws.cell(r, 7, "TOTAL"); c.font = f(10, True, BRANCO); c.fill = FILL_TITULO
    c.alignment = Alignment(horizontal="center")
    r += 1
    for lbl, src in [("Receita bruta", F["bruta"]), ("Receita líquida", F["liquida"]),
                     ("Impostos sobre receita", F["imp"]),
                     ("Comissão de lojas", F["loja"])]:
        rotulo(ws, r, lbl, 1, bold=(lbl == "Receita bruta"))
        for j in range(5):
            L = gcl(2 + j)
            c = ws.cell(r, 2 + j,
                f'=SUMIFS($B${src}:${L_FIM}${src},$B${LANO}:${L_FIM}${LANO},{L}${hdr})')
            c.number_format = BRL; c.font = f(9, lbl == "Receita bruta")
        c = ws.cell(r, 7, f'=SUM(B{r}:F{r})')
        c.number_format = BRL; c.font = f(9, True); c.fill = FILL_TOTAL
        r += 1
    F["anual_hdr"] = hdr
    return ws, F


# =====================================================================
def build_fluxo(wb, P, A, I, BLOCOS):
    ws = wb.create_sheet("Análise Fluxo")
    sel, yr, yr_cen = mk_sel(P, BLOCOS)
    titulo(ws, "ANÁLISE DE FLUXO DE CAIXA E VIABILIDADE",
           "Estrutura herdada da planilha de Análise de Investimento: Projeto (investimentos) → "
           "Operacionais → Fluxo Líquido → VPL / TIR / Payback.")
    largura_padrao(ws, col_a=46)
    cabecalho_tempo(ws, linha_data=TEMPO, congelar="B9", anos_row=P["anos"])
    ws["A2"] = '=" Cenário ativo: "&Premissas!B' + str(P["cenario_nome"])
    ws["A2"].font = f(9, True, VERDE_LINK)

    F = {}
    r = 8
    # ---------------- PROJETO ------------------------------------------
    secao(ws, r, "PROJETO — INVESTIMENTOS", COL_PCT); r += 1
    proj_ini = r
    for lbl, cat in [("Desenvolvimento do app (MVP e evoluções)", "Desenvolvimento"),
                     ("Equipamentos", "Equipamentos"),
                     ("Constituição, marca e jurídico", "Marca e jurídico"),
                     ("Certificações e conformidade", "Certificações")]:
        rotulo(ws, r, lbl, 1)
        preencher_linha(ws, r, lambda i, L, c=cat: f'=-Investimentos!{L}{I[c]}', BRL,
                        color=VERDE_LINK)
        F[cat] = r; r += 1
    rotulo(ws, r, "TOTAL PROJETO", 0, bold=True)
    preencher_linha(ws, r, lambda i, L: f'=SUM({L}{proj_ini}:{L}{r-1})', BRL, bold=True)
    for i in range(1, N_MESES + 1):
        ws[f"{col_mes(i)}{r}"].fill = FILL_TOTAL
    F["projeto"] = r
    for cat in ["Desenvolvimento", "Equipamentos", "Marca e jurídico", "Certificações"]:
        ws[f"{L_PCT}{F[cat]}"] = f'=IFERROR({L_TOT}{F[cat]}/{L_TOT}${F["projeto"]},0)'
        ws[f"{L_PCT}{F[cat]}"].number_format = PCT
        ws[f"{L_PCT}{F[cat]}"].font = f(9, False, CINZA)
    r += 2

    # ---------------- OPERACIONAIS -------------------------------------
    secao(ws, r, "OPERACIONAIS", COL_PCT); r += 1
    op_ini = r
    ops = [
        ("Infraestrutura e nuvem",              lambda i, L: f'={sel(23, L)}',                     "nuvem"),
        ("Suporte ao cliente",                  lambda i, L: f'={sel(24, L)}',                     "suporte"),
        ("Equipe PJ (produto e engenharia)",
         lambda i, L: f'=-IF({L}${TEMPO}>=Premissas!$B${P["equipe_ini"]},{yr("equipe", L)},0)',    "equipe"),
        ("Pró-labore dos sócios",
         lambda i, L: f'=-IF({L}${TEMPO}>=Premissas!$B${P["prolab_ini"]},{yr("prolab", L)},0)',    "prolab"),
        ("Marketing de performance (aquisição)",
         lambda i, L: f'=-{sel(1, L)}*(1-{yr("organico", L)})*{yr_cen("cac", L)}',                 "mkt_perf"),
        ("Marketing recorrente (agência e conteúdo)",
         lambda i, L: f'=-IF({L}${TEMPO}>=Premissas!$B${P["mkt_ini"]},{yr("mkt_rec", L)},0)',      "mkt_rec"),
        ("Propaganda de lançamento",
         lambda i, L: (f'=-IF({L}${TEMPO}=DATE(YEAR(Premissas!$B${P["mkt_lanc_mes"]}),'
                       f'MONTH(Premissas!$B${P["mkt_lanc_mes"]}),1),Premissas!$B${P["mkt_lanc"]},0)'),
         "mkt_lanc"),
        ("Comissão das lojas de aplicativos",   lambda i, L: f'={sel(20, L)}',                     "loja"),
        ("Taxas de meios de pagamento",         lambda i, L: f'={sel(21, L)}',                     "pgto"),
        ("Impostos sobre a receita",            lambda i, L: f'={sel(19, L)}',                     "imp"),
        ("Despesas administrativas",
         lambda i, L: f'=-IF({L}${TEMPO}>=Premissas!$B${P["adm_ini"]},{yr("admin", L)},0)',        "adm"),
    ]
    for lbl, fn, key in ops:
        rotulo(ws, r, lbl, 1)
        preencher_linha(ws, r, fn, BRL)
        F[key] = r; r += 1
    rotulo(ws, r, "TOTAL OPERACIONAIS", 0, bold=True)
    preencher_linha(ws, r, lambda i, L: f'=SUM({L}{op_ini}:{L}{r-1})', BRL, bold=True)
    for i in range(1, N_MESES + 1):
        ws[f"{col_mes(i)}{r}"].fill = FILL_TOTAL
    F["opex"] = r
    for _, _, key in ops:
        ws[f"{L_PCT}{F[key]}"] = f'=IFERROR({L_TOT}{F[key]}/{L_TOT}${F["opex"]},0)'
        ws[f"{L_PCT}{F[key]}"].number_format = PCT
        ws[f"{L_PCT}{F[key]}"].font = f(9, False, CINZA)
    r += 2

    # ---------------- FLUXO --------------------------------------------
    secao(ws, r, "FLUXO DE CAIXA", COL_PCT); r += 1
    rotulo(ws, r, "Receita bruta total", 1)
    preencher_linha(ws, r, lambda i, L: f'={sel(16, L)}', BRL, color=VERDE_LINK)
    F["receita"] = r; r += 1
    rotulo(ws, r, "FLUXO DE CAIXA LIVRE (FCL)", 0, bold=True)
    preencher_linha(ws, r, lambda i, L:
        f'={L}{F["receita"]}+{L}{F["projeto"]}+{L}{F["opex"]}', BRL, bold=True)
    for i in range(1, N_MESES + 1):
        ws[f"{col_mes(i)}{r}"].fill = FILL_TOTAL
    nota(ws, f"{L_PCT}{r}", "Sem aportes — aporte é financiamento, não geração de caixa.")
    F["fcl"] = r; r += 1
    rotulo(ws, r, "FCL descontado", 1)
    preencher_linha(ws, r, lambda i, L:
        f'={L}{F["fcl"]}/(1+Premissas!$B${P["tma_m"]})^{L}${LIDX}', BRL)
    F["fcld"] = r; r += 1
    rotulo(ws, r, "FCL descontado acumulado", 1)
    preencher_linha(ws, r, lambda i, L:
        f'={L}{F["fcld"]}' if i == 1 else f'={gcl(COL0+i-2)}{r}+{L}{F["fcld"]}',
        BRL, total=False)
    F["fcld_ac"] = r; r += 1
    rotulo(ws, r, "FCL acumulado (nominal)", 1)
    preencher_linha(ws, r, lambda i, L:
        f'={L}{F["fcl"]}' if i == 1 else f'={gcl(COL0+i-2)}{r}+{L}{F["fcl"]}',
        BRL, total=False)
    F["fcl_ac"] = r; r += 2

    rotulo(ws, r, "(+) Aportes de sócios", 1)
    preencher_linha(ws, r, lambda i, L: f'=Aportes!{L}{A["total"]}', BRL, color=VERDE_LINK)
    F["aporte"] = r; r += 1
    rotulo(ws, r, "Fluxo de caixa do período (FCL + aportes)", 1, bold=True)
    preencher_linha(ws, r, lambda i, L: f'={L}{F["fcl"]}+{L}{F["aporte"]}', BRL, bold=True)
    F["fc_total"] = r; r += 1
    rotulo(ws, r, "CAIXA ACUMULADO", 0, bold=True)
    preencher_linha(ws, r, lambda i, L:
        f'={L}{F["fc_total"]}' if i == 1 else f'={gcl(COL0+i-2)}{r}+{L}{F["fc_total"]}',
        BRL, bold=True, total=False)
    for i in range(1, N_MESES + 1):
        ws[f"{col_mes(i)}{r}"].fill = FILL_TOTAL
    F["caixa"] = r; r += 1
    rotulo(ws, r, "✔ Verificação: FCL desta aba − FCL do motor (deve ser zero)", 1, italic=True)
    preencher_linha(ws, r, lambda i, L: f'=ROUND({L}{F["fcl"]}-{sel(30, L)},2)', BRL2,
                    color=CINZA, total=False)
    ws[f"{L_TOT}{r}"] = f'=SUM(B{r}:{L_FIM}{r})'
    ws[f"{L_TOT}{r}"].number_format = BRL2
    ws[f"{L_TOT}{r}"].font = f(9, True, "C00000")
    nota(ws, f"{L_PCT}{r}", "Zero confirma que o detalhamento reconcilia com o motor de cenários.")
    F["check"] = r; r += 2

    # ---------------- INDICADORES --------------------------------------
    secao(ws, r, "INDICADORES DE VIABILIDADE — CENÁRIO SELECIONADO", COL_PCT); r += 1
    ind = [
        ("VPL — Valor Presente Líquido",
         f'=NPV(Premissas!$B${P["tma_m"]},B{F["fcl"]}:{L_FIM}{F["fcl"]})', BRL,
         "Soma do FCL trazido a valor presente pela TMA. Positivo = cria valor."),
        ("TIR mensal", f'=IFERROR(IRR(B{F["fcl"]}:{L_FIM}{F["fcl"]}),"n/d")', PCT2,
         "Taxa que zera o VPL."),
        ("TIR anual", f'=IFERROR((1+$B{r+1})^12-1,"n/d")', PCT,
         "Compare com a TMA. Acima dela, o projeto remunera o risco."),
        ("Payback simples (meses)",
         f'=IF({L_FIM}{F["fcl_ac"]}<0,"> 60 meses",COUNTIF(B{F["fcl_ac"]}:{L_FIM}{F["fcl_ac"]},"<0")+1)',
         INT, "Mês em que o FCL acumulado vira positivo."),
        ("Payback descontado (meses)",
         f'=IF({L_FIM}{F["fcld_ac"]}<0,"> 60 meses",COUNTIF(B{F["fcld_ac"]}:{L_FIM}{F["fcld_ac"]},"<0")+1)',
         INT, "Idem, com o dinheiro trazido a valor presente."),
        ("Capital requerido (pico de caixa negativo)",
         f'=-MIN(B{F["fcl_ac"]}:{L_FIM}{F["fcl_ac"]})', BRL,
         "Quanto de capital o projeto precisa antes de se sustentar."),
        ("Caixa mínimo ao longo do período",
         f'=MIN(B{F["caixa"]}:{L_FIM}{F["caixa"]})', BRL,
         "Se negativo, o capital planejado não é suficiente."),
        ("Investimento total (CAPEX)", f'=-{L_TOT}{F["projeto"]}', BRL, ""),
        ("Receita bruta acumulada", f'={L_TOT}{F["receita"]}', BRL, ""),
        ("FCL acumulado no período", f'={L_TOT}{F["fcl"]}', BRL, ""),
        ("ROI sobre o capital requerido",
         f'=IFERROR({L_TOT}{F["fcl"]}/$B{r+5},0)', PCT, "FCL total ÷ capital requerido."),
        ("TMA utilizada (a.a.)", f'=Premissas!$B${P["tma"]}', PCT, ""),
        ("Alíquota efetiva média de impostos",
         f'=IFERROR(-{L_TOT}{F["imp"]}/{L_TOT}{F["receita"]},0)', PCT2, ""),
    ]
    ind_ini = r
    for lbl, fx, fmt, obs in ind:
        rotulo(ws, r, lbl, 1)
        c = ws.cell(r, 2, fx)
        c.number_format = fmt; c.font = f(10, True, PRETO)
        c.fill = FILL_TOTAL; c.border = BORDA_FINA
        c.alignment = Alignment(horizontal="center")
        if obs:
            nota(ws, f"D{r}", obs)
        r += 1
    F["ind_ini"] = ind_ini
    F["vpl"], F["tir_m"], F["tir_a"] = ind_ini, ind_ini + 1, ind_ini + 2
    F["pb"], F["pbd"], F["capreq"], F["caixa_min"] = ind_ini + 3, ind_ini + 4, ind_ini + 5, ind_ini + 6
    r += 1

    # ---------------- PREMISSAS DE CUSTO --------------------------------
    secao(ws, r, "PREMISSAS — RESUMO (editar na aba Premissas)", COL_PCT); r += 1
    for lbl, key, fmt in [("Plano Básico (R$/mês)", "p_basico", BRL2),
                          ("Plano Intermediário (R$/mês)", "p_inter", BRL2),
                          ("Plano Premium (R$/mês)", "p_premium", BRL2),
                          ("Taxa das lojas de aplicativos", "taxa_loja", PCT),
                          ("Comissão do marketplace", "com_mp", PCT),
                          ("TMA (a.a.)", "tma", PCT),
                          ("Payback desejado (meses)", "payback_alvo", INT)]:
        rotulo(ws, r, lbl, 1)
        c = ws.cell(r, 2, f'=Premissas!$B${P[key]}')
        c.number_format = fmt; c.font = f(10, False, VERDE_LINK)
        c.alignment = Alignment(horizontal="center"); c.border = BORDA_FINA
        r += 1
    return ws, F


# =====================================================================
def build_dre(wb, P, I, BLOCOS, FL):
    ws = wb.create_sheet("DRE Projetada")
    sel, yr, yr_cen = mk_sel(P, BLOCOS)
    titulo(ws, "DRE PROJETADA — DEMONSTRAÇÃO DO RESULTADO DO EXERCÍCIO",
           "Regime de competência. Investimentos NÃO aparecem como despesa: entram como depreciação "
           "e amortização em 60 meses. Por isso o EBITDA aqui difere do fluxo de caixa livre.")
    largura_padrao(ws, col_a=48)
    cabecalho_tempo(ws, linha_data=TEMPO, congelar="B9", anos_row=P["anos"])
    ws["A2"] = '=" Cenário ativo: "&Premissas!B' + str(P["cenario_nome"])
    ws["A2"].font = f(9, True, VERDE_LINK)

    D = {}
    r = 8

    def linha(lbl, fn, key, fmt=BRL, bold=False, nivel=1, destaque=False, obs=None,
              cor=PRETO, pct_base=None):
        nonlocal r
        rotulo(ws, r, lbl, 0 if bold else nivel, bold=bold, color=cor)
        preencher_linha(ws, r, fn, fmt, bold=bold, color=cor,
                        pct_base=pct_base)
        if destaque:
            for i in range(1, N_MESES + 1):
                ws[f"{col_mes(i)}{r}"].fill = FILL_TOTAL
        if obs:
            nota(ws, f"{L_PCT}{r}", obs)
        D[key] = r
        r += 1
        return D[key]

    secao(ws, r, "RECEITA OPERACIONAL BRUTA", COL_PCT); r += 1
    linha("Assinaturas — Plano Básico",        lambda i, L: f'={sel(10, L)}', "rb_b", cor=VERDE_LINK)
    linha("Assinaturas — Plano Intermediário", lambda i, L: f'={sel(11, L)}', "rb_i", cor=VERDE_LINK)
    linha("Assinaturas — Plano Premium",       lambda i, L: f'={sel(12, L)}', "rb_p", cor=VERDE_LINK)
    linha("Marketplace — comissões",           lambda i, L: f'={sel(14, L)}', "rb_mp", cor=VERDE_LINK)
    linha("B2B — Boston Scientific",           lambda i, L: f'={sel(15, L)}', "rb_b2b", cor=VERDE_LINK)
    linha("(=) RECEITA BRUTA", lambda i, L: f'=SUM({L}{D["rb_b"]}:{L}{D["rb_b2b"]})',
          "bruta", bold=True, destaque=True)
    r += 1

    secao(ws, r, "DEDUÇÕES DA RECEITA BRUTA", COL_PCT); r += 1
    linha("(−) Impostos sobre a receita (Simples / Lucro Presumido)",
          lambda i, L: f'={sel(19, L)}', "ded_imp", cor=VERDE_LINK, pct_base=D["bruta"])
    linha("(−) Comissão das lojas de aplicativos",
          lambda i, L: f'={sel(20, L)}', "ded_loja", cor=VERDE_LINK, pct_base=D["bruta"],
          obs="Apple/Google. A maior dedução isolada da receita de assinatura.")
    linha("(−) Taxas de meios de pagamento",
          lambda i, L: f'={sel(21, L)}', "ded_pgto", cor=VERDE_LINK, pct_base=D["bruta"])
    linha("(=) RECEITA OPERACIONAL LÍQUIDA",
          lambda i, L: f'={L}{D["bruta"]}+SUM({L}{D["ded_imp"]}:{L}{D["ded_pgto"]})',
          "liquida", bold=True, destaque=True, pct_base=D["bruta"])
    r += 1

    secao(ws, r, "CUSTO DOS SERVIÇOS PRESTADOS", COL_PCT); r += 1
    linha("(−) Infraestrutura e nuvem", lambda i, L: f'={sel(23, L)}', "cs_nuvem",
          cor=VERDE_LINK, pct_base=D["bruta"])
    linha("(−) Suporte ao cliente", lambda i, L: f'={sel(24, L)}', "cs_sup",
          cor=VERDE_LINK, pct_base=D["bruta"])
    linha("(=) LUCRO BRUTO",
          lambda i, L: f'={L}{D["liquida"]}+{L}{D["cs_nuvem"]}+{L}{D["cs_sup"]}',
          "bruto", bold=True, destaque=True)
    rotulo(ws, r, "Margem bruta (% da receita líquida)", 2, italic=True)
    preencher_linha(ws, r, lambda i, L: f'=IFERROR({L}{D["bruto"]}/{L}{D["liquida"]},0)',
                    PCT, color=CINZA, total=False)
    ws[f"{L_TOT}{r}"] = f'=IFERROR({L_TOT}{D["bruto"]}/{L_TOT}{D["liquida"]},0)'
    ws[f"{L_TOT}{r}"].number_format = PCT; ws[f"{L_TOT}{r}"].font = f(9, True)
    ws[f"{L_TOT}{r}"].fill = FILL_TOTAL
    D["mg_bruta"] = r; r += 2

    secao(ws, r, "DESPESAS OPERACIONAIS", COL_PCT); r += 1
    linha("(−) Pessoal — equipe PJ e pró-labore",
          lambda i, L: f"='Análise Fluxo'!{L}{FL['equipe']}+'Análise Fluxo'!{L}{FL['prolab']}",
          "do_pes", cor=VERDE_LINK, pct_base=D["bruta"])
    linha("(−) Marketing e aquisição",
          lambda i, L: (f"='Análise Fluxo'!{L}{FL['mkt_perf']}+'Análise Fluxo'!{L}{FL['mkt_rec']}"
                        f"+'Análise Fluxo'!{L}{FL['mkt_lanc']}"),
          "do_mkt", cor=VERDE_LINK, pct_base=D["bruta"],
          obs="Inclui a propaganda de lançamento, tratada como despesa e não como investimento.")
    linha("(−) Despesas administrativas",
          lambda i, L: f"='Análise Fluxo'!{L}{FL['adm']}", "do_adm",
          cor=VERDE_LINK, pct_base=D["bruta"])
    linha("(=) EBITDA",
          lambda i, L: f'={L}{D["bruto"]}+SUM({L}{D["do_pes"]}:{L}{D["do_adm"]})',
          "ebitda", bold=True, destaque=True)
    rotulo(ws, r, "Margem EBITDA (% da receita líquida)", 2, italic=True)
    preencher_linha(ws, r, lambda i, L: f'=IFERROR({L}{D["ebitda"]}/{L}{D["liquida"]},0)',
                    PCT, color=CINZA, total=False)
    ws[f"{L_TOT}{r}"] = f'=IFERROR({L_TOT}{D["ebitda"]}/{L_TOT}{D["liquida"]},0)'
    ws[f"{L_TOT}{r}"].number_format = PCT; ws[f"{L_TOT}{r}"].font = f(9, True)
    ws[f"{L_TOT}{r}"].fill = FILL_TOTAL
    D["mg_ebitda"] = r; r += 2

    secao(ws, r, "RESULTADO", COL_PCT); r += 1
    linha("(−) Depreciação e amortização",
          lambda i, L: f'=-Investimentos!{L}{I["da"]}', "da", cor=VERDE_LINK,
          obs="Investimentos amortizados linearmente em 60 meses.")
    linha("(=) EBIT — resultado operacional",
          lambda i, L: f'={L}{D["ebitda"]}+{L}{D["da"]}', "ebit", bold=True, destaque=True)
    linha("(+/−) Resultado financeiro", lambda i, L: '=0', "fin", cor=CINZA,
          obs="Rendimento do caixa não considerado (premissa conservadora).")
    linha("(=) RESULTADO ANTES DO IR",
          lambda i, L: f'={L}{D["ebit"]}+{L}{D["fin"]}', "lair", bold=True)
    linha("(−) IRPJ e CSLL", lambda i, L: '=0', "ir", cor=CINZA,
          obs="Já incluídos na linha de impostos sobre a receita (DAS do Simples ou Lucro Presumido).")
    linha("(=) LUCRO LÍQUIDO DO EXERCÍCIO",
          lambda i, L: f'={L}{D["lair"]}+{L}{D["ir"]}', "ll", bold=True, destaque=True)
    rotulo(ws, r, "Margem líquida (% da receita líquida)", 2, italic=True)
    preencher_linha(ws, r, lambda i, L: f'=IFERROR({L}{D["ll"]}/{L}{D["liquida"]},0)',
                    PCT, color=CINZA, total=False)
    ws[f"{L_TOT}{r}"] = f'=IFERROR({L_TOT}{D["ll"]}/{L_TOT}{D["liquida"]},0)'
    ws[f"{L_TOT}{r}"].number_format = PCT; ws[f"{L_TOT}{r}"].font = f(9, True)
    ws[f"{L_TOT}{r}"].fill = FILL_TOTAL
    D["mg_liq"] = r; r += 1
    rotulo(ws, r, "Lucro líquido acumulado", 1, italic=True)
    preencher_linha(ws, r, lambda i, L:
        f'={L}{D["ll"]}' if i == 1 else f'={gcl(COL0+i-2)}{r}+{L}{D["ll"]}',
        BRL, color=CINZA, total=False)
    D["ll_ac"] = r; r += 2

    # ---------------- DRE ANUAL -----------------------------------------
    secao(ws, r, "DRE ANUAL CONSOLIDADA", COL_PCT); r += 1
    hdr = r
    ws.cell(r, 1, "Demonstração do Resultado").font = f(10, True, BRANCO)
    ws.cell(r, 1).fill = FILL_HEADER
    for j, a in enumerate([2026, 2027, 2028, 2029, 2030]):
        c = ws.cell(r, 2 + j, a)
        c.font = f(10, True, BRANCO); c.fill = FILL_HEADER
        c.alignment = Alignment(horizontal="center"); c.number_format = '0'
    c = ws.cell(r, 7, "5 ANOS")
    c.font = f(10, True, BRANCO); c.fill = FILL_TITULO
    c.alignment = Alignment(horizontal="center")
    r += 1
    D["anual_hdr"] = hdr

    anual = [
        ("RECEITA BRUTA", "bruta", "soma", True),
        ("(−) Impostos sobre a receita", "ded_imp", "soma", False),
        ("(−) Comissão das lojas de aplicativos", "ded_loja", "soma", False),
        ("(−) Taxas de meios de pagamento", "ded_pgto", "soma", False),
        ("(=) RECEITA LÍQUIDA", "liquida", "soma", True),
        ("(−) Infraestrutura e nuvem", "cs_nuvem", "soma", False),
        ("(−) Suporte ao cliente", "cs_sup", "soma", False),
        ("(=) LUCRO BRUTO", "bruto", "soma", True),
        ("Margem bruta", None, "mg_bruto", False),
        ("(−) Pessoal", "do_pes", "soma", False),
        ("(−) Marketing e aquisição", "do_mkt", "soma", False),
        ("(−) Despesas administrativas", "do_adm", "soma", False),
        ("(=) EBITDA", "ebitda", "soma", True),
        ("Margem EBITDA", None, "mg_ebitda", False),
        ("(−) Depreciação e amortização", "da", "soma", False),
        ("(=) EBIT", "ebit", "soma", True),
        ("(=) LUCRO LÍQUIDO", "ll", "soma", True),
        ("Margem líquida", None, "mg_liq", False),
    ]
    ref = {}
    for lbl, key, modo, bold in anual:
        rotulo(ws, r, lbl, 0 if bold else 1, bold=bold)
        for j in range(5):
            L = gcl(2 + j)
            if modo == "soma":
                fx = (f'=SUMIFS($B${D[key]}:${L_FIM}${D[key]},'
                      f'$B${LANO}:${L_FIM}${LANO},{L}${hdr})')
                fmt = BRL
            elif modo == "mg_bruto":
                fx = f'=IFERROR({L}{ref["bruto"]}/{L}{ref["liquida"]},0)'; fmt = PCT
            elif modo == "mg_ebitda":
                fx = f'=IFERROR({L}{ref["ebitda"]}/{L}{ref["liquida"]},0)'; fmt = PCT
            else:
                fx = f'=IFERROR({L}{ref["ll"]}/{L}{ref["liquida"]},0)'; fmt = PCT
            c = ws.cell(r, 2 + j, fx)
            c.number_format = fmt
            c.font = f(9, bold, PRETO if modo == "soma" else CINZA)
            c.border = BORDA_FINA
            if bold:
                c.fill = FILL_SUB
        if modo == "soma":
            c = ws.cell(r, 7, f'=SUM(B{r}:F{r})'); c.number_format = BRL
        elif modo == "mg_bruto":
            c = ws.cell(r, 7, f'=IFERROR(G{ref["bruto"]}/G{ref["liquida"]},0)'); c.number_format = PCT
        elif modo == "mg_ebitda":
            c = ws.cell(r, 7, f'=IFERROR(G{ref["ebitda"]}/G{ref["liquida"]},0)'); c.number_format = PCT
        else:
            c = ws.cell(r, 7, f'=IFERROR(G{ref["ll"]}/G{ref["liquida"]},0)'); c.number_format = PCT
        c.font = f(9, True, PRETO); c.fill = FILL_TOTAL; c.border = BORDA_FINA
        if key:
            ref[key] = r
        r += 1
    D["anual_ref"] = ref
    return ws, D
