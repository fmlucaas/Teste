# -*- coding: utf-8 -*-
"""Aba Preco Minimo: a conta de tras pra frente.

O resto do modelo vai do preco para o resultado. Esta aba vai do custo para o
preco: dado o que a Oppa gasta, quanto a assinatura precisa custar.

A identidade que sustenta tudo:

    EBITDA = preco_essencial x unidades_equivalentes x rho + outras_receitas_liq - custo_operacional

onde rho e quanto sobra de cada real cobrado depois de imposto, loja e meio de
pagamento, e "unidades equivalentes" converte Premium e Familia em multiplos do
Essencial pela razao de preco. Zerando o EBITDA sai o preco de equilibrio. A
linha de reconciliacao no fim da aba confere essa identidade contra o motor.

Nada aqui realimenta a projecao. Fechar o laco seria circular: preco muda
receita, receita muda a faixa do Simples, a faixa muda rho, rho muda o preco.
"""
from openpyxl.styles import Alignment, PatternFill
from common import *
from engine import O
from ajuste import ov
from detalhe import mk_sel

TEMPO, LIDX, LANO, LIDXA = 4, 5, 6, 7
COLS = ["B", "C", "D"]   # colunas de cenario nos blocos de celula unica
NOMES = ["CONSERVADOR", "PROVÁVEL", "AGRESSIVO"]
CORES = ["C00000", "1F3864", "1F7040"]


def build(wb, P, BLOCOS, FL, D):
    ws = wb.create_sheet("Preço Mínimo")
    sel, yr, yr_cen, ov_cen = mk_sel(P, BLOCOS)
    titulo(ws, "PREÇO MÍNIMO — DO CUSTO PARA O PREÇO",
           "Quanto a assinatura precisa custar para pagar a operação. Mês a mês, no cenário "
           "ativo. Esta aba não altera a projeção: é régua, não motor.")
    largura_padrao(ws, col_a=54, col_mes_w=13)
    cabecalho_tempo(ws, linha_data=TEMPO, congelar="B9", anos_row=P["anos"])
    ws["A2"] = '=" Cenário ativo: "&Premissas!B' + str(P["cenario_nome"])
    ws["A2"].font = f(9, True, VERDE_LINK)

    M = {}
    r = 8

    def linha(lbl, fn, key, fmt=BRL, bold=False, nivel=1, destaque=False,
              obs=None, cor=PRETO, total=True):
        nonlocal r
        rotulo(ws, r, lbl, 0 if bold else nivel, bold=bold, color=cor)
        preencher_linha(ws, r, fn, fmt, bold=bold, color=cor, total=total)
        if destaque:
            for i in range(1, N_MESES + 1):
                ws[f"{col_mes(i)}{r}"].fill = FILL_TOTAL
        if obs:
            nota(ws, f"{L_PCT}{r}", obs)
        M[key] = r
        r += 1
        return M[key]

    pb, pi, pp = P["p_basico"], P["p_inter"], P["p_premium"]

    # =================================================================
    secao(ws, r, "1. QUEM PAGA, E QUANTO PAGA HOJE", COL_PCT); r += 1
    linha("Assinantes pagantes", lambda i, L: f'={sel(O.pag, L)}', "pag", INT,
          cor=VERDE_LINK, total=False)
    linha("Pagantes — Essencial", lambda i, L: f'={sel(O.pag_ess, L)}', "p_ess", INT,
          cor=VERDE_LINK, total=False)
    linha("Pagantes — Premium", lambda i, L: f'={sel(O.pag_pre, L)}', "p_pre", INT,
          cor=VERDE_LINK, total=False,
          obs="Zero até o mês de entrada do plano, definido em Premissas.")
    linha("Pagantes — Família", lambda i, L: f'={sel(O.pag_fam, L)}', "p_fam", INT,
          cor=VERDE_LINK, total=False)
    linha("Preço do Essencial vigente no mês",
          lambda i, L: f'=Premissas!$B${pb}*(1+Premissas!$B${P["reajuste"]})^({L}${LANO}-2026)',
          "pe", BRL2, total=False)
    linha("Premium ÷ Essencial", lambda i, L: f'=IFERROR(Premissas!$B${pi}/Premissas!$B${pb},0)',
          "k_pre", MULT, total=False, cor=CINZA,
          obs="A razão entre os preços. É o que permite traduzir todo mundo em Essencial.")
    linha("Família ÷ Essencial", lambda i, L: f'=IFERROR(Premissas!$B${pp}/Premissas!$B${pb},0)',
          "k_fam", MULT, total=False, cor=CINZA)
    linha("Assinantes equivalentes a Essencial (unidades equivalentes)",
          lambda i, L: (f'={L}{M["p_ess"]}+{L}{M["p_pre"]}*{L}{M["k_pre"]}'
                        f'+{L}{M["p_fam"]}*{L}{M["k_fam"]}'), "ue", INT, bold=True, total=False,
          obs="Receita de assinatura = preço do Essencial × esta linha. É a unidade da conta.")
    linha("Receita de assinaturas (conferência: deve bater com o motor)",
          lambda i, L: f'={L}{M["pe"]}*{L}{M["ue"]}-{sel(O.assinaturas, L)}', "chk_assin",
          BRL2, cor=CINZA, obs="Deve ser zero em todos os meses.")
    linha("ARPPU real do mês (receita de assinatura ÷ pagantes)",
          lambda i, L: f'=IFERROR({sel(O.assinaturas, L)}/{L}{M["pag"]},0)', "arppu", BRL2,
          total=False, obs="Durante a fase só-Essencial, igual ao preço do Essencial.")
    r += 1

    # =================================================================
    secao(ws, r, "2. QUANTO SOBRA DE CADA REAL COBRADO", COL_PCT); r += 1
    linha("(−) Imposto sobre a receita (alíquota efetiva)",
          lambda i, L: f'={sel(O.aliquota, L)}', "aliq", PCT2, cor=VERDE_LINK, total=False)
    linha("(−) Comissão das lojas, efetiva",
          lambda i, L: f'=Premissas!$B${P["taxa_loja"]}*{yr("share_loja", L)}', "c_loja",
          PCT2, total=False,
          obs="15% só sobre a parte cobrada dentro do app. Fora dela não há comissão.")
    linha("(−) Meio de pagamento, efetivo",
          lambda i, L: f'=Premissas!$B${P["taxa_pgto"]}*(1-{yr("share_loja", L)})', "c_pgto",
          PCT2, total=False)
    linha("(=) Fica para a empresa, de cada real de assinatura",
          lambda i, L: f'=1-{L}{M["aliq"]}-{L}{M["c_loja"]}-{L}{M["c_pgto"]}', "rho",
          PCT2, bold=True, destaque=True, total=False,
          obs="É o divisor de toda a aba: cobrar R$ 1 não é receber R$ 1.")
    r += 1

    # =================================================================
    secao(ws, r, "3. O CUSTO A COBRIR", COL_PCT); r += 1
    linha("Nuvem por assinante pagante", lambda i, L: f'={yr("cloud_pag", L)}', "cv_nuvem",
          BRL2, total=False)
    linha("Suporte por assinante pagante", lambda i, L: f'={yr("sup_var", L)}', "cv_sup",
          BRL2, total=False)
    linha("Nuvem dos usuários gratuitos, rateada por pagante",
          lambda i, L: (f'=IFERROR(({sel(O.base, L)}-{sel(O.pag, L)})*{yr("cloud_free", L)}'
                        f'/{L}{M["pag"]},0)'), "cv_free", BRL2, total=False,
          obs="O plano gratuito não paga, mas consome. Quem paga sustenta quem não paga.")
    linha("Painel de dados, rateado por pagante",
          lambda i, L: (f'=IFERROR({sel(O.painel, L)}*Premissas!$B${P["custo_painel"]}'
                        f'/{L}{M["pag"]},0)'), "cv_painel", BRL2, total=False)
    linha("(=) Custo variável por pagante",
          lambda i, L: f'=SUM({L}{M["cv_nuvem"]}:{L}{M["cv_painel"]})', "cvp", BRL2,
          bold=True, total=False)
    linha("CAC por assinante conquistado",
          lambda i, L: (f'=IFERROR((1-{yr("organico", L)})*{ov_cen("cac", L)}'
                        f'/{sel(O.conv, L)},0)'), "cac_pag", BRL2, total=False, cor=CINZA,
          obs="CAC por usuário captado, dividido pela conversão: o custo de conquistar um PAGANTE.")
    linha("CAC diluído por mês de vida do assinante",
          lambda i, L: f'={L}{M["cac_pag"]}*{ov_cen("churn", L)}', "cac_mes", BRL2,
          total=False, cor=CINZA,
          obs="Vida esperada = 1 ÷ churn. Multiplicar por churn é diluir o CAC nessa vida.")
    r += 1
    custo_op = "+".join(f"'Análise Fluxo'!{{L}}{FL[k]}" for k in
                        ("nuvem", "suporte", "equipe", "clt", "prolab",
                         "mkt_perf", "mkt_rec", "mkt_lanc", "adm"))
    linha("Custo operacional total do mês",
          lambda i, L: "=-(" + custo_op.replace("{L}", L) + ")", "custo_op", bold=True,
          cor=VERDE_LINK,
          obs="Nuvem, suporte, pessoas, marketing e administrativas. Sem imposto, loja e meio "
              "de pagamento: esses já estão descontados no fator acima.")
    linha("Outras receitas, líquidas (marketplace e B2B)",
          lambda i, L: (f'=({sel(O.marketplace, L)})*(1-{L}{M["aliq"]}-Premissas!$B${P["taxa_pgto"]})'
                        f'+({sel(O.parcerias, L)}+{sel(O.dados, L)}+{sel(O.outras_b2b, L)})'
                        f'*(1-{L}{M["aliq"]})'), "outras_liq", cor=VERDE_LINK,
          obs="Entram como abatimento: o que o marketplace e o B2B pagam, a assinatura não "
              "precisa pagar. Venda de dados não passa por loja nem por adquirente.")
    linha("(=) Custo que a assinatura tem de cobrir",
          lambda i, L: f'={L}{M["custo_op"]}-{L}{M["outras_liq"]}', "custo_liq", bold=True,
          destaque=True)
    linha("Investimento do mês (CAPEX) — fora da conta de preço",
          lambda i, L: f"=-'Análise Fluxo'!{L}{FL['projeto']}", "capex", cor=CINZA,
          obs="Equipamento e certificação são investimento, pagos por capital, não por "
              "assinatura. Ficam de fora de propósito — aparecem aqui só para você ver o tamanho.")
    r += 1

    # =================================================================
    secao(ws, r, "4. OS PISOS DE PREÇO DO ESSENCIAL, MÊS A MÊS", COL_PCT); r += 1
    linha("Piso 1 — cobre só o custo variável (margem de contribuição zero)",
          lambda i, L: f'=IFERROR({L}{M["cvp"]}/{L}{M["rho"]},0)', "piso1", BRL2,
          total=False, obs="Abaixo disto, cada assinante novo dá prejuízo. É o piso absoluto.")
    linha("Piso 2 — cobre o variável e o CAC diluído",
          lambda i, L: f'=IFERROR(({L}{M["cvp"]}+{L}{M["cac_mes"]})/{L}{M["rho"]},0)', "piso2",
          BRL2, total=False, obs="Abaixo disto, crescer destrói valor: o assinante não devolve o que custou conquistar.")
    linha("Piso 3 — EQUILIBRA O MÊS (EBITDA zero)",
          lambda i, L: f'=IFERROR({L}{M["custo_liq"]}/({L}{M["ue"]}*{L}{M["rho"]}),0)', "piso3",
          BRL2, bold=True, destaque=True, total=False,
          obs="Com a base daquele mês. Nos primeiros meses fica absurdo — e é esse o ponto: "
              "não é o preço que está errado, é a base que ainda é pequena.")
    linha("Preço-alvo — entrega a margem EBITDA desejada",
          lambda i, L: (f'=IFERROR(({L}{M["custo_op"]}/(1-Premissas!$B${P["margem_alvo"]})'
                        f'-{L}{M["outras_liq"]})/({L}{M["ue"]}*{L}{M["rho"]}),0)'), "alvo",
          BRL2, total=False)
    r += 1
    rotulo(ws, r, "Os mesmos pisos, traduzidos para os outros dois planos", 0, bold=True); r += 1
    for lbl, base, k, key in [
        ("Premium — piso de equilíbrio do mês", "piso3", "k_pre", "piso3_pre"),
        ("Família — piso de equilíbrio do mês", "piso3", "k_fam", "piso3_fam"),
        ("Premium — preço-alvo", "alvo", "k_pre", "alvo_pre"),
        ("Família — preço-alvo", "alvo", "k_fam", "alvo_fam"),
    ]:
        linha(lbl, lambda i, L, b=base, kk=k: f'={L}{M[b]}*{L}{M[kk]}', key, BRL2, total=False)
    nota(ws, f"{L_PCT}{M['piso3_pre']}",
         "Mantém a razão de preço atual entre os planos. Mude os preços em Premissas e a razão muda com eles.")
    r += 1

    # =================================================================
    secao(ws, r, "5. O OUTRO LADO DA MESMA CONTA: QUANTOS ASSINANTES", COL_PCT); r += 1
    linha("Assinantes equivalentes necessários para equilibrar o mês",
          lambda i, L: f'=IFERROR({L}{M["custo_liq"]}/({L}{M["pe"]}*{L}{M["rho"]}),0)', "ue_nec",
          INT, bold=True, total=False,
          obs="Ao preço de tabela. É o piso 3 lido pelo outro lado.")
    linha("Assinantes equivalentes projetados", lambda i, L: f'={L}{M["ue"]}', "ue2", INT,
          total=False, cor=CINZA)
    linha("Faltam (ou sobram)", lambda i, L: f'={L}{M["ue2"]}-{L}{M["ue_nec"]}', "ue_dif",
          INT, total=False,
          obs="Positivo = o mês se paga no preço atual.")
    linha("Marcador — mês em que a projeção cobre o custo",
          lambda i, L: (f'=IF(AND({L}{M["ue2"]}>0,{L}{M["ue2"]}>={L}{M["ue_nec"]}),'
                        f'{L}${LIDX},9999)'), "mk_eq", INT, total=False, cor=CINZA)
    linha("✔ Reconciliação com o motor (deve ser zero em todos os meses)",
          lambda i, L: (f'={L}{M["pe"]}*{L}{M["ue"]}*{L}{M["rho"]}+{L}{M["outras_liq"]}'
                        f'-{L}{M["custo_op"]}-{sel(O.ebitda, L)}'), "chk_ident", BRL2,
          cor=CINZA,
          obs="Prova a identidade da aba contra o EBITDA do motor. Se der diferente de zero, "
              "a decomposição de custo ou o fator de retenção estão errados.")
    r += 2

    # =================================================================
    # Bloco da janela inicial — celulas unicas, nao mensais
    # =================================================================
    secao(ws, r, "6. A PERGUNTA DO LANÇAMENTO: O ESSENCIAL COBRE OS PRIMEIROS MESES?", COL_PCT)
    r += 1
    W = {}

    def unica(lbl, fx, key=None, fmt=BRL, bold=False, obs=None, cor=PRETO):
        nonlocal r
        c = ws.cell(r, 1, lbl)
        c.font = f(10, bold, cor)
        c.alignment = Alignment(indent=1, vertical="center")
        cc = ws.cell(r, 2, fx)
        cc.number_format = fmt
        cc.font = f(10, bold, cor)
        cc.alignment = Alignment(horizontal="center")
        cc.border = BORDA_FINA
        if bold:
            cc.fill = FILL_TOTAL
        if obs:
            c2 = ws.cell(r, 4, obs)
            c2.font = f(8, False, CINZA, italic=True)
            c2.alignment = Alignment(wrap_text=True, vertical="center", indent=1)
            ws.merge_cells(start_row=r, start_column=4, end_row=r, end_column=12)
            ws.row_dimensions[r].height = max(ws.row_dimensions[r].height or 15,
                                              15 + 11 * (len(obs) // 110))
        if key:
            W[key] = r
        r += 1
        return r - 1

    idx = f'$B${LIDX}:${L_FIM}${LIDX}'
    jan = (f'({idx}>=Premissas!$B${P["lanc_idx"]})'
           f'*({idx}<=Premissas!$B${P["lanc_idx"]}+Premissas!$B${P["meses_cob"]}-1)')
    pre = f'({idx}<Premissas!$B${P["lanc_idx"]})'

    def _rng(key):
        return f'$B${M[key]}:${L_FIM}${M[key]}'

    unica("Primeiro mês da janela", f'=Premissas!$B${P["lancamento"]}', "ini", MES)
    unica("Último mês da janela",
          f'=EDATE(Premissas!$B${P["lancamento"]},Premissas!$B${P["meses_cob"]}-1)', "fim", MES,
          obs="O número de meses está em Premissas, seção 3.")
    unica("Custo operacional dentro da janela",
          f'=SUMPRODUCT({jan}*{_rng("custo_op")})', "c_jan")
    unica("Custo operacional anterior ao lançamento",
          f'=SUMPRODUCT({pre}*{_rng("custo_op")})*Premissas!$B${P["cob_prelanc"]}', "c_pre",
          obs="Nuvem desde maio, administrativas desde abril, a equipe PJ de setembro e outubro "
              "e a propaganda de lançamento. Entra ou não conforme o interruptor em Premissas.")
    unica("(−) Outras receitas líquidas na janela",
          f'=-SUMPRODUCT({jan}*{_rng("outras_liq")})', "o_jan")
    unica("(=) Total que a assinatura precisa cobrir",
          f'=B{W["c_jan"]}+B{W["c_pre"]}+B{W["o_jan"]}', "tot", bold=True)
    unica("Assinantes equivalentes × fator de retenção, somados na janela",
          f'=SUMPRODUCT({jan}*{_rng("ue")}*{_rng("rho")})', "uer", fmt='#,##0.0')
    unica("PREÇO ÚNICO DO ESSENCIAL QUE ZERA A JANELA",
          f'=IFERROR(B{W["tot"]}/B{W["uer"]},0)', "preco", BRL2, bold=True,
          obs="Um só preço, cobrado em todos os meses da janela, que faz a conta fechar no fim dela.")
    unica("Premium equivalente",
          f'=B{W["preco"]}*IFERROR(Premissas!$B${pi}/Premissas!$B${pb},0)', "preco_pre", BRL2)
    unica("Família equivalente",
          f'=B{W["preco"]}*IFERROR(Premissas!$B${pp}/Premissas!$B${pb},0)', "preco_fam", BRL2)
    unica("Preço do Essencial hoje", f'=Premissas!$B${pb}', "atual", BRL2)
    unica("Quantas vezes o preço de hoje",
          f'=IFERROR(B{W["preco"]}/B{W["atual"]},0)', "mult", MULT, bold=True,
          obs="Acima de ~1,5x o preço perde o mercado: é sinal de que a resposta não é preço.")
    unica("Assinantes necessários na janela, mantendo o preço de hoje",
          f'=IFERROR(B{W["tot"]}/(B{W["atual"]}*SUMPRODUCT({jan}*{_rng("rho")})),0)', "n_nec",
          INT, bold=True,
          obs="Média de assinantes equivalentes ao longo da janela, não no último mês.")
    unica("Assinantes equivalentes projetados no último mês da janela",
          f'=IFERROR(INDEX({_rng("ue")},Premissas!$B${P["lanc_idx"]}'
          f'+Premissas!$B${P["meses_cob"]}-1),0)', "n_proj", INT)
    unica("Média de assinantes equivalentes projetada na janela",
          f'=IFERROR(SUMPRODUCT({jan}*{_rng("ue")})/Premissas!$B${P["meses_cob"]},0)',
          "n_med", INT)
    unica("Quanto falta de base para o preço de hoje bastar",
          f'=IFERROR(B{W["n_nec"]}/B{W["n_med"]},0)', "gap", MULT, bold=True,
          obs="Quantas vezes a base projetada teria de ser maior. É a mesma pergunta do preço, "
              "do lado do volume.")
    unica("VEREDICTO",
          f'=IF(B{W["n_med"]}<=0,"Sem base na janela — nenhum preço cobre",'
          f'IF(B{W["mult"]}<=1,"O preço de hoje cobre a janela",'
          f'IF(B{W["mult"]}<=1.5,"Dá para fechar com reajuste: "&TEXT(B{W["preco"]},"R$ 0.00"),'
          f'"Nenhum preço de mercado cobre: a janela é financiada por capital")))',
          "veredicto", TXT, bold=True)
    r += 1
    c = ws.cell(r, 1, "A leitura honesta: assinatura de ticket baixo não paga custo fixo com "
                      "base pequena. O preço cobre o custo variável desde o primeiro assinante "
                      "— é o piso 1 e ele é baixo. O custo fixo dos primeiros meses é pago por "
                      "capital, e é exatamente isso que a aba NECESSIDADE DE CAPITAL dimensiona. "
                      "Preço alto demais no lançamento não resolve o caixa: reduz a conversão e "
                      "piora o caixa.")
    c.font = f(9, False, "C00000", italic=True)
    c.alignment = Alignment(wrap_text=True, vertical="top", indent=1)
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=12)
    ws.row_dimensions[r].height = 58
    r += 2

    # =================================================================
    secao(ws, r, "7. RESUMO POR CENÁRIO", COL_PCT); r += 1
    hdr = r
    ws.cell(r, 1, "INDICADOR").font = f(10, True, BRANCO)
    ws.cell(r, 1).fill = FILL_HEADER
    ws.cell(r, 1).alignment = Alignment(indent=1, vertical="center")
    for k, col in enumerate(COLS):
        c = ws[f"{col}{r}"]
        c.value = NOMES[k]
        c.font = f(10, True, BRANCO)
        c.fill = PatternFill("solid", fgColor=CORES[k])
        c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[r].height = 22
    r += 1

    def Bk(k, off):
        return f"'Cenários'!$B${BLOCOS[k]+off}:${L_FIM}${BLOCOS[k]+off}"

    S = {}

    def cen(lbl, fx, key=None, fmt=BRL, bold=False, obs=None):
        nonlocal r
        c = ws.cell(r, 1, lbl)
        c.font = f(10, bold, PRETO)
        c.alignment = Alignment(indent=1, vertical="center")
        for k, col in enumerate(COLS):
            cc = ws[f"{col}{r}"]
            cc.value = fx(k, col)
            cc.number_format = fmt
            cc.font = f(10, bold, PRETO)
            cc.alignment = Alignment(horizontal="center")
            cc.border = BORDA_FINA
            if bold:
                cc.fill = FILL_TOTAL
        if obs:
            c2 = ws.cell(r, 6, obs)
            c2.font = f(8, False, CINZA, italic=True)
            c2.alignment = Alignment(wrap_text=True, vertical="center", indent=1)
            ws.merge_cells(start_row=r, start_column=6, end_row=r, end_column=14)
            ws.row_dimensions[r].height = max(ws.row_dimensions[r].height or 15,
                                              15 + 11 * (len(obs) // 95))
        if key:
            S[key] = r
        r += 1
        return r - 1

    cen("Assinantes equivalentes no último mês da janela",
        lambda k, col: (f'=IFERROR(INDEX({Bk(k, O.pag_ess)},Premissas!$B${P["lanc_idx"]}'
                        f'+Premissas!$B${P["meses_cob"]}-1)'
                        f'+INDEX({Bk(k, O.pag_pre)},Premissas!$B${P["lanc_idx"]}'
                        f'+Premissas!$B${P["meses_cob"]}-1)*Premissas!$B${pi}/Premissas!$B${pb}'
                        f'+INDEX({Bk(k, O.pag_fam)},Premissas!$B${P["lanc_idx"]}'
                        f'+Premissas!$B${P["meses_cob"]}-1)*Premissas!$B${pp}/Premissas!$B${pb},0)'),
        "ue_fim", INT)
    cen("EBITDA acumulado na janela",
        lambda k, col: (f'=SUMPRODUCT({jan}*{Bk(k, O.ebitda)})'), "eb_jan",
        obs="Negativo significa que a janela não se paga: a diferença vem de capital.")
    cen("EBITDA acumulado até o fim da janela, somando o que veio antes",
        lambda k, col: (f'=SUMPRODUCT(({idx}<=Premissas!$B${P["lanc_idx"]}'
                        f'+Premissas!$B${P["meses_cob"]}-1)*{Bk(k, O.ebitda)})'), "eb_tot",
        bold=True)
    cen("Primeiro mês em que o EBITDA cobre o custo",
        lambda k, col: (f'=IF(MAX({Bk(k, O.flag_eb)})=0,"não atinge",'
                        f'CHOOSE(MOD(60-SUM({Bk(k, O.flag_eb)}),12)+1,'
                        f'"jan","fev","mar","abr","mai","jun","jul","ago","set","out","nov","dez")'
                        f'&"/"&(2026+INT((60-SUM({Bk(k, O.flag_eb)}))/12)))'), "mes_eq", TXT,
        obs="Mesmo mês do break-even no Resumo. Se divergir, há inconsistência entre as abas.")
    return ws, M, W, S
