# -*- coding: utf-8 -*-
"""Abas Tributos, Aportes e Investimentos."""
import datetime as dt
from openpyxl.styles import Alignment, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment
from common import *

# =====================================================================
#  TRIBUTOS
# =====================================================================
ANEXO_III = [  # faixa, de, ate, aliquota nominal, parcela a deduzir
    (1,        0.00,  180000.00, 0.0600,      0.00),
    (2,   180000.01,  360000.00, 0.1120,   9360.00),
    (3,   360000.01,  720000.00, 0.1350,  17640.00),
    (4,   720000.01, 1800000.00, 0.1600,  35640.00),
    (5,  1800000.01, 3600000.00, 0.2100, 125640.00),
    (6,  3600000.01, 4800000.00, 0.3300, 648000.00),
]
ANEXO_V = [
    (1,        0.00,  180000.00, 0.1550,      0.00),
    (2,   180000.01,  360000.00, 0.1800,   4500.00),
    (3,   360000.01,  720000.00, 0.1950,   9900.00),
    (4,   720000.01, 1800000.00, 0.2050,  17100.00),
    (5,  1800000.01, 3600000.00, 0.2300,  62100.00),
    (6,  3600000.01, 4800000.00, 0.3050, 540000.00),
]

T = {}

def build_tributos(wb, P):
    ws = wb.create_sheet("Tributos")
    titulo(ws, "TRIBUTAÇÃO — SIMPLES NACIONAL E LUCRO PRESUMIDO",
           "Tabelas da LC 123/2006. A 'Tabela Ativa' é a que o modelo usa, conforme o regime "
           "escolhido em Premissas.", largura=8)
    for col, w in zip("ABCDEFGH", (10, 20, 20, 18, 20, 6, 78, 6)):
        ws.column_dimensions[col].width = w

    def tabela(lin, nome, dados, cor):
        secao(ws, lin, nome, 8)
        hdr = ["Faixa", "RBT12 — de", "RBT12 — até", "Alíquota nominal", "Parcela a deduzir"]
        for j, h in enumerate(hdr):
            c = ws.cell(lin + 1, 1 + j, h)
            c.font = f(9, True, BRANCO); c.fill = FILL_HEADER
            c.alignment = Alignment(horizontal="center", wrap_text=True)
        for k, (fx, de, ate, al, ded) in enumerate(dados):
            rr = lin + 2 + k
            ws.cell(rr, 1, fx).number_format = INT
            ws.cell(rr, 2, de).number_format = BRL
            ws.cell(rr, 3, ate).number_format = BRL
            ws.cell(rr, 4, al).number_format = PCT2
            ws.cell(rr, 5, ded).number_format = BRL
            for j in range(1, 6):
                cc = ws.cell(rr, j)
                cc.font = f(9, False, cor); cc.border = BORDA_FINA
                cc.alignment = Alignment(horizontal="center")
        return lin + 2   # primeira linha de dados

    r = 4
    T["a3"] = tabela(r, "ANEXO III — Serviços em geral (alíquota efetiva de 6% a ~19,5%)", ANEXO_III, AZUL_INPUT)
    nota(ws, "G5", "Aplica-se a software/SaaS SOMENTE se o Fator R ≥ 28%.")
    r += 9
    T["a5"] = tabela(r, "ANEXO V — Licenciamento/cessão de uso de software (Fator R < 28%)", ANEXO_V, AZUL_INPUT)
    nota(ws, "G14", "Regra padrão para licenciamento de software quando o Fator R fica abaixo de 28%.")
    r += 9

    # --- Tabela ativa ---------------------------------------------------
    secao(ws, r, "TABELA ATIVA — selecionada pelo regime definido em Premissas", 8)
    hdr = ["Faixa", "RBT12 — de", "Alíquota nominal", "Parcela a deduzir"]
    for j, h in enumerate(hdr):
        c = ws.cell(r + 1, 1 + j, h)
        c.font = f(9, True, BRANCO); c.fill = FILL_HEADER
        c.alignment = Alignment(horizontal="center", wrap_text=True)
    base = r + 2
    for k in range(6):
        rr = base + k
        ws.cell(rr, 1, k + 1).number_format = INT
        ws.cell(rr, 2, f'=IF(Premissas!$B${P["regime"]}=2,B{T["a5"]+k},B{T["a3"]+k})').number_format = BRL
        ws.cell(rr, 3, f'=IF(Premissas!$B${P["regime"]}=2,D{T["a5"]+k},D{T["a3"]+k})').number_format = PCT2
        ws.cell(rr, 4, f'=IF(Premissas!$B${P["regime"]}=2,E{T["a5"]+k},E{T["a3"]+k})').number_format = BRL
        for j in range(1, 5):
            cc = ws.cell(rr, j)
            cc.font = f(9, True, VERDE_LINK); cc.border = BORDA_FINA
            cc.fill = FILL_SUB
            cc.alignment = Alignment(horizontal="center")
    T["ativa_ini"] = base
    T["ativa_fim"] = base + 5
    r = base + 7

    # --- Fator R --------------------------------------------------------
    secao(ws, r, "FATOR R — o teste que decide entre Anexo III e Anexo V", 8); r += 1
    txt = [
        ("Fator R = (folha de salários + pró-labore + encargos dos últimos 12 meses) ÷ RBT12", True),
        ("Fator R ≥ 28%  →  Anexo III (alíquota efetiva a partir de 6%)", False),
        ("Fator R < 28%   →  Anexo V (alíquota efetiva a partir de 15,5%)", False),
        ("", False),
        ("⚠ RISCO MATERIAL PARA A OPPA: pagamento a equipe PJ (prestadores de serviço) NÃO entra no "
         "cálculo do Fator R. Só contam folha CLT e pró-labore dos sócios. Com a estrutura descrita no "
         "brief — equipe PJ e pró-labore zero em 2026 — o Fator R fica próximo de 0% e a empresa cai no "
         "Anexo V, com alíquota inicial de 15,5% em vez de 6%.", False),
        ("", False),
        ("Como reverter: (a) registrar pró-labore relevante para os sócios; (b) contratar parte da equipe "
         "em regime CLT. Ambos aumentam o custo de folha, mas podem valer a pena pela economia tributária. "
         "O quadro abaixo compara os dois anexos no faturamento projetado.", False),
    ]
    for t, b in txt:
        c = ws.cell(r, 1, t)
        c.font = f(9, b, "C00000" if t.startswith("⚠") else CINZA, italic=not b)
        c.alignment = Alignment(wrap_text=True, vertical="top", indent=1)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=7)
        ws.row_dimensions[r].height = 15 if len(t) < 90 else 44
        r += 1
    T["fator_r_calc"] = r + 1
    r += 1

    # --- Comparativo de anexos ------------------------------------------
    secao(ws, r, "COMPARATIVO — alíquota efetiva por faixa de faturamento anual", 8); r += 1
    for j, h in enumerate(["RBT12", "Anexo III efetivo", "Anexo V efetivo", "Diferença",
                           "Custo anual da diferença"]):
        c = ws.cell(r, 1 + j, h)
        c.font = f(9, True, BRANCO); c.fill = FILL_HEADER
        c.alignment = Alignment(horizontal="center", wrap_text=True)
    r += 1
    for k, rbt in enumerate([180000, 360000, 720000, 1800000, 3600000, 4800000]):
        rr = r + k
        ws.cell(rr, 1, rbt).number_format = BRL
        ws.cell(rr, 2, f'=($A{rr}*INDEX($D${T["a3"]}:$D${T["a3"]+5},MATCH($A{rr},$B${T["a3"]}:$B${T["a3"]+5},1))'
                       f'-INDEX($E${T["a3"]}:$E${T["a3"]+5},MATCH($A{rr},$B${T["a3"]}:$B${T["a3"]+5},1)))/$A{rr}'
               ).number_format = PCT2
        ws.cell(rr, 3, f'=($A{rr}*INDEX($D${T["a5"]}:$D${T["a5"]+5},MATCH($A{rr},$B${T["a5"]}:$B${T["a5"]+5},1))'
                       f'-INDEX($E${T["a5"]}:$E${T["a5"]+5},MATCH($A{rr},$B${T["a5"]}:$B${T["a5"]+5},1)))/$A{rr}'
               ).number_format = PCT2
        ws.cell(rr, 4, f'=$C{rr}-$B{rr}').number_format = PCT2
        ws.cell(rr, 5, f'=$D{rr}*$A{rr}').number_format = BRL
        for j in range(1, 6):
            cc = ws.cell(rr, j)
            cc.font = f(9, False, PRETO); cc.border = BORDA_FINA
            cc.alignment = Alignment(horizontal="center")
        ws.cell(rr, 5).font = f(9, True, "C00000")
    nota(ws, f"G{r}", "Quanto a Oppa paga a mais por ano se cair no Anexo V em vez do Anexo III.")
    return ws, T


# =====================================================================
#  APORTES E INVESTIDORES
# =====================================================================
APORTES = [
    # data, socio, valor, status, tipo, observacao
    (dt.date(2026, 4, 16), "Paulo Kapp",        10000, "Realizado",  "Fundador", ""),
    (dt.date(2026, 7,  4), "Lucas Medina",       5000, "Realizado",  "Fundador", ""),
    (dt.date(2026, 8, 20), "Paulo Kapp",        10000, "Realizado",  "Fundador",
     "Completa os R$ 20.000 de compromisso."),
    (dt.date(2026, 8, 20), "Felipe Toescher",    3000, "Realizado",  "Fundador", ""),
    (dt.date(2026, 8, 20), "Lucas Iung",         2000, "Realizado",  "Fundador", ""),
    (dt.date(2026, 8, 29), "Lucas Iung",         3000, "Realizado",  "Fundador", ""),
    (dt.date(2026,12, 15), "Lucas Medina",       5000, "Compromisso","Fundador",
     "Fecha os R$ 10.000 exigidos ainda em 2026."),
    (dt.date(2026,12, 15), "Felipe Toescher",    7000, "Compromisso","Fundador",
     "Fecha os R$ 10.000 exigidos ainda em 2026."),
    (dt.date(2026,12, 15), "Lucas Iung",         5000, "Compromisso","Fundador",
     "Fecha os R$ 10.000 exigidos ainda em 2026."),
    (dt.date(2027, 6, 15), "Lucas Medina",       5000, "Compromisso","Fundador", ""),
    (dt.date(2027, 6, 15), "Felipe Toescher",    5000, "Compromisso","Fundador", ""),
    (dt.date(2027, 6, 15), "Lucas Iung",         5000, "Compromisso","Fundador", ""),
    (dt.date(2027,12, 15), "Lucas Medina",       5000, "Compromisso","Fundador",
     "Completa os R$ 20.000."),
    (dt.date(2027,12, 15), "Felipe Toescher",    5000, "Compromisso","Fundador",
     "Completa os R$ 20.000."),
    (dt.date(2027,12, 15), "Lucas Iung",         5000, "Compromisso","Fundador",
     "Completa os R$ 20.000."),
    (dt.date(2026,10, 15), "C-ioT",             40000, "Previsto",   "Novo sócio",
     "Negociação em curso, sem assinatura."),
    (dt.date(2026,10, 15), "Felipe Martinelli", 40000, "Previsto",   "Novo sócio",
     "Negociação em curso, sem assinatura."),
    (dt.date(2026,11, 15), "Shaiane",           40000, "Previsto",   "Novo sócio",
     "Sócia responsável pelo marketing."),
]
VAZIAS_AP = 14

# Investidores estrategicos: aporte + receita recorrente de parceria
ESTRATEGICOS = [
    ("Boston Scientific", "Em negociação", 500000, dt.date(2027, 7, 1),
     80000, dt.date(2027, 7, 1), dt.date(2030, 12, 1),
     "Monitoramento de pacientes pós-operatórios. Aporte na casa dos seis dígitos, não "
     "confirmado — adotado o ponto médio da faixa. Contrapartida esperada: exclusividade "
     "B2B hospitalar por um período longo."),
]
VAZIAS_EST = 7

FUNDADORES = ["Paulo Kapp", "Lucas Medina", "Felipe Toescher", "Lucas Iung"]
SOCIOS_CAP = FUNDADORES + ["C-ioT", "Felipe Martinelli", "Shaiane"]
VAZIAS_CAP = 6

A = {}


def build_aportes(wb, P):
    ws = wb.create_sheet("Aportes")
    titulo(ws, "APORTES E INVESTIDORES",
           "Todo o capital da Oppa entra por esta aba. Há linhas em branco em todas as tabelas: "
           "para incluir um investidor novo, basta preencher uma delas — o modelo inteiro "
           "recalcula sozinho, sem mexer em fórmula nenhuma.")
    largura_padrao(ws, col_a=30, col_mes_w=11)
    for col, w in zip("ABCDEFGH", (16, 24, 16, 16, 16, 16, 16, 58)):
        ws.column_dimensions[col].width = w

    # ---------------- 1. lançamentos ------------------------------------
    r = 4
    secao(ws, r, "1. LANÇAMENTOS DE APORTE — quem colocou (ou colocará) dinheiro, e quando", 8)
    r += 1
    for j, h in enumerate(["Data", "Sócio / Investidor", "Valor", "Status", "Tipo", "Observação"]):
        c = ws.cell(r, 1 + j, h)
        c.font = f(9, True, BRANCO); c.fill = FILL_HEADER
        c.alignment = Alignment(horizontal="center")
    ws.merge_cells(start_row=r, start_column=6, end_row=r, end_column=8)
    r += 1
    A["ini"] = r
    linhas = sorted(APORTES, key=lambda x: x[0]) + [
        (dt.date(2027, 1, 15), "", 0, "Previsto", "Novo sócio", "") for _ in range(VAZIAS_AP)]
    for (d, s, v, st, tp, obs) in linhas:
        vazia = not s
        ws.cell(r, 1, dt.datetime(d.year, d.month, d.day)).number_format = DATA
        ws.cell(r, 2, s)
        ws.cell(r, 3, v).number_format = BRL
        ws.cell(r, 4, st)
        ws.cell(r, 5, tp)
        ws.cell(r, 6, obs)
        ws.merge_cells(start_row=r, start_column=6, end_row=r, end_column=8)
        cor = AZUL_INPUT if st == "Realizado" else ("000000" if st == "Compromisso" else "BF8F00")
        for j in range(1, 9):
            cc = ws.cell(r, j)
            cc.font = f(9, st == "Realizado" and not vazia, CINZA if vazia else cor)
            cc.border = BORDA_FINA
            cc.alignment = Alignment(horizontal="center" if j not in (2, 6) else "left",
                                     indent=1 if j in (2, 6) else 0)
            if vazia:
                cc.fill = PatternFill()
            elif st == "Realizado":
                cc.fill = FILL_OK
            elif st == "Previsto":
                cc.fill = FILL_ALERTA
        r += 1
    A["fim"] = r - 1
    ws.cell(r, 2, "TOTAL LANÇADO (conforme os controles de status)").font = f(10, True)
    A["total_lanc"] = r
    r += 2

    # ---------------- 2. investidores estratégicos ----------------------
    secao(ws, r, "2. INVESTIDORES ESTRATÉGICOS E PARCERIAS — aporte mais receita recorrente", 8)
    r += 1
    c = ws.cell(r, 1, "Use esta tabela para investidor que, além de aportar, traz receita — como a "
                      "Boston Scientific. Preencha uma linha em branco para incluir outro. O aporte "
                      "entra no caixa no mês indicado; a receita recorrente entra na DRE como "
                      "receita B2B, todo mês, entre as datas de início e fim.")
    c.font = f(9, False, CINZA, italic=True)
    c.alignment = Alignment(wrap_text=True, vertical="top", indent=1)
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8)
    ws.row_dimensions[r].height = 28
    r += 1
    for j, h in enumerate(["Investidor", "Status", "Aporte (R$)", "Mês do aporte",
                           "Receita B2B (R$/mês)", "Início da receita", "Fim da receita",
                           "Observação"]):
        c = ws.cell(r, 1 + j, h)
        c.font = f(9, True, BRANCO); c.fill = FILL_HEADER
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[r].height = 28
    r += 1
    A["est_ini"] = r
    ests = list(ESTRATEGICOS) + [
        ("", "Em negociação", 0, dt.date(2028, 1, 1), 0, dt.date(2028, 1, 1),
         dt.date(2030, 12, 1), "") for _ in range(VAZIAS_EST)]
    for (nome, st, ap, mes, rec, ini_r, fim_r, obs) in ests:
        vazia = not nome
        ws.cell(r, 1, nome)
        ws.cell(r, 2, st)
        ws.cell(r, 3, ap).number_format = BRL
        ws.cell(r, 4, dt.datetime(mes.year, mes.month, 1)).number_format = MES
        ws.cell(r, 5, rec).number_format = BRL
        ws.cell(r, 6, dt.datetime(ini_r.year, ini_r.month, 1)).number_format = MES
        ws.cell(r, 7, dt.datetime(fim_r.year, fim_r.month, 1)).number_format = MES
        ws.cell(r, 8, obs)
        for j in range(1, 9):
            cc = ws.cell(r, j)
            cc.font = f(9, False, CINZA if vazia else AZUL_INPUT)
            cc.border = BORDA_FINA
            cc.alignment = Alignment(horizontal="center" if j not in (1, 8) else "left",
                                     indent=1 if j in (1, 8) else 0,
                                     wrap_text=(j == 8))
            if not vazia:
                cc.fill = FILL_ALERTA
        if not vazia:
            ws.row_dimensions[r].height = 44
        r += 1
    A["est_fim"] = r - 1
    ws.cell(r, 1, "TOTAL").font = f(10, True)
    A["total_est"] = r
    r += 2
    # portões de status: o que entra e o que não entra na conta
    i, fi = A["ini"], A["fim"]
    ei, ef = A["est_ini"], A["est_fim"]
    G = (f'(($D${i}:$D${fi}="Realizado")+($D${i}:$D${fi}="Compromisso")'
         f'+($D${i}:$D${fi}="Previsto")*Premissas!$B${P["inc_previsto"]}'
         f'+($D${i}:$D${fi}="Em negociação")*Premissas!$B${P["inc_negoc"]})')
    GE = (f'(($B${ei}:$B${ef}="Confirmado")'
          f'+($B${ei}:$B${ef}="Em negociação")*Premissas!$B${P["inc_negoc"]})')
    A["gate_est"] = GE

    c = ws.cell(A["total_lanc"], 3, f'=SUMPRODUCT({G}*$C${i}:$C${fi})')
    c.number_format = BRL; c.font = f(10, True); c.fill = FILL_TOTAL
    c = ws.cell(A["total_est"], 3, f'=SUMPRODUCT({GE}*$C${ei}:$C${ef})')
    c.number_format = BRL; c.font = f(10, True); c.fill = FILL_TOTAL
    c = ws.cell(A["total_est"], 5, f'=SUMPRODUCT({GE}*$E${ei}:$E${ef})')
    c.number_format = BRL; c.font = f(10, True); c.fill = FILL_TOTAL
    nota(ws, f"I{A['total_est']}", "Receita B2B mensal somada de todas as parcerias ativas.")

    # --- validações -----------------------------------------------------
    for col, opts, rng in (("D", '"Realizado,Compromisso,Previsto,Em negociação"', (i, fi)),
                           ("E", '"Fundador,Novo sócio,Investidor estratégico"', (i, fi)),
                           ("B", '"Confirmado,Em negociação"', (ei, ef))):
        dv = DataValidation(type="list", formula1=opts, allow_blank=True, showDropDown=False)
        ws.add_data_validation(dv)
        dv.add(f"{col}{rng[0]}:{col}{rng[1]}")

    # --- 3. resumo do capital --------------------------------------------
    secao(ws, r, "3. RESUMO DO CAPITAL", 8); r += 1
    linhas_res = [
        ("Realizado — dinheiro que já entrou", f'=SUMIFS($C${i}:$C${fi},$D${i}:$D${fi},"Realizado")', False),
        ("Compromisso dos fundadores a integralizar",
         f'=SUMIFS($C${i}:$C${fi},$D${i}:$D${fi},"Compromisso")', False),
        ("Previsto — negociação em curso, sem assinatura",
         f'=SUMIFS($C${i}:$C${fi},$D${i}:$D${fi},"Previsto")*Premissas!$B${P["inc_previsto"]}', False),
        ("Em negociação — lançamentos",
         f'=SUMIFS($C${i}:$C${fi},$D${i}:$D${fi},"Em negociação")*Premissas!$B${P["inc_negoc"]}', False),
        ("Investidores estratégicos", f'=SUMPRODUCT({GE}*$C${ei}:$C${ef})', False),
        ("CAPITAL TOTAL CONSIDERADO NO MODELO",
         f'=C{A["total_lanc"]}+C{A["total_est"]}', True),
    ]
    for lbl, fx, bold in linhas_res:
        ws.cell(r, 2, lbl).font = f(10 if bold else 9, bold)
        c = ws.cell(r, 3, fx)
        c.number_format = BRL; c.font = f(10 if bold else 9, True)
        if bold:
            c.fill = FILL_TOTAL
        r += 1
    A["capital_total"] = r - 1
    nota(ws, f"E{A['capital_total']}",
         "Os controles que ligam e desligam 'Previsto' e 'Em negociação' estão na aba Premissas, seção 7.")
    r += 2

    # --- 4. compromisso dos fundadores -----------------------------------
    secao(ws, r, "4. COMPROMISSO DOS FUNDADORES — R$ 20.000 cada, sendo R$ 10.000 ainda em 2026", 8)
    r += 1
    for j, h in enumerate(["Fundador", "Meta total", "Aportado até 2026", "Aportado total",
                           "Falta em 2026", "Falta total"]):
        c = ws.cell(r, 1 + j, h)
        c.font = f(9, True, BRANCO); c.fill = FILL_HEADER
        c.alignment = Alignment(horizontal="center", wrap_text=True)
    r += 1
    A["meta"] = r
    for nome in FUNDADORES:
        ws.cell(r, 1, nome).font = f(9, True)
        ws.cell(r, 2, 20000).number_format = BRL
        ws.cell(r, 2).font = f(9, False, AZUL_INPUT)
        ws.cell(r, 3, f'=SUMIFS($C${i}:$C${fi},$B${i}:$B${fi},$A{r},'
                      f'$A${i}:$A${fi},"<="&DATE(2026,12,31))').number_format = BRL
        ws.cell(r, 4, f'=SUMIFS($C${i}:$C${fi},$B${i}:$B${fi},$A{r})').number_format = BRL
        ws.cell(r, 5, f'=MAX(0,10000-$C{r})').number_format = BRL
        ws.cell(r, 6, f'=MAX(0,$B{r}-$D{r})').number_format = BRL
        for j in range(1, 7):
            cc = ws.cell(r, j)
            if j > 2:
                cc.font = f(9, False, PRETO)
            cc.border = BORDA_FINA
            cc.alignment = Alignment(horizontal="center" if j > 1 else "left", indent=1)
        r += 1
    nota(ws, f"H{A['meta']}", "Paulo Kapp já integralizou os R$ 20.000. Os demais têm saldo a aportar.")
    r += 2

    # --- 5. participação proporcional ------------------------------------
    secao(ws, r, "5. PARTICIPAÇÃO PROPORCIONAL AO CAPITAL APORTADO — inclua aqui o nome de todo investidor novo", 8); r += 1
    c = ws.cell(r, 1, "Leia com cuidado: isto é a proporção do dinheiro colocado, não a participação "
                      "societária. Percentual de sociedade depende do valuation acordado em cada "
                      "rodada e de outras condições — no caso de um investidor estratégico como a "
                      "Boston, a contrapartida pode ser exclusividade comercial em vez de mais "
                      "equity. Use esta tabela como ponto de partida da conversa, não como resultado.")
    c.font = f(9, False, "C00000", italic=True)
    c.alignment = Alignment(wrap_text=True, vertical="top", indent=1)
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8)
    ws.row_dimensions[r].height = 32
    r += 1
    for j, h in enumerate(["Sócio / Investidor", "Capital aportado", "Proporção"]):
        c = ws.cell(r, 1 + j, h)
        c.font = f(9, True, BRANCO); c.fill = FILL_HEADER
        c.alignment = Alignment(horizontal="center")
    r += 1
    cap_ini = r
    for nome in SOCIOS_CAP + [""] * VAZIAS_CAP:
        ws.cell(r, 1, nome).font = f(9, True if nome else False, PRETO if nome else CINZA)
        ws.cell(r, 2, f'=IF($A{r}="",0,SUMPRODUCT({G}*($B${i}:$B${fi}=$A{r})*$C${i}:$C${fi})'
                      f'+SUMPRODUCT({GE}*($A${ei}:$A${ef}=$A{r})*$C${ei}:$C${ef}))'
               ).number_format = BRL
        ws.cell(r, 3, f'=IFERROR($B{r}/$B${{TOTCAP}},0)').number_format = PCT
        for j in range(1, 4):
            cc = ws.cell(r, j)
            cc.border = BORDA_FINA
            cc.alignment = Alignment(horizontal="center" if j > 1 else "left", indent=1)
            if j > 1:
                cc.font = f(9, False, PRETO)
        r += 1
    cap_fim = r - 1
    ws.cell(r, 1, "TOTAL").font = f(10, True)
    c = ws.cell(r, 2, f'=SUM(B{cap_ini}:B{cap_fim})')
    c.number_format = BRL; c.font = f(10, True); c.fill = FILL_TOTAL
    c = ws.cell(r, 3, f'=SUM(C{cap_ini}:C{cap_fim})')
    c.number_format = PCT; c.font = f(10, True); c.fill = FILL_TOTAL
    for rr in range(cap_ini, cap_fim + 1):
        ws.cell(rr, 3).value = ws.cell(rr, 3).value.replace("{TOTCAP}", str(r))
    A["cap_total"] = r
    r += 1
    ws.cell(r, 1, "Capital não atribuído nesta tabela").font = f(10, True, "C00000")
    c = ws.cell(r, 2, f'=C{A["total_lanc"]}+C{A["total_est"]}-B{A["cap_total"]}')
    c.number_format = BRL; c.font = f(10, True, "C00000")
    c.border = BORDA_FINA
    c.alignment = Alignment(horizontal="center")
    c2 = ws.cell(r, 3, f'=IF(ROUND(B{r},2)=0,"✔ tudo atribuído",'
                       f'"⚠ acrescente o nome que falta")')
    c2.font = f(10, True, "C00000")
    c2.alignment = Alignment(horizontal="center")
    nota(ws, f"E{r}", "Se aparecer valor diferente de zero, algum investidor foi lançado nas tabelas "
                      "1 ou 2 e não foi acrescentado aqui. Digite o nome dele numa linha em branco "
                      "acima — exatamente como está escrito no lançamento.")
    A["cap_dif"] = r
    r += 2

    # --- 6. cronograma mensal --------------------------------------------
    secao(ws, r, "6. CRONOGRAMA MENSAL DE ENTRADA DE CAPITAL", COL_PCT); r += 1
    cabecalho_tempo(ws, linha_data=r, congelar=f"B{r+5}", anos_row=P["anos"])
    A["tempo"] = r
    t = r
    r += 4

    def janela(L):
        return f'($A${i}:$A${fi}>={L}${t})*($A${i}:$A${fi}<=EOMONTH({L}${t},0))'

    rotulo(ws, r, "Aportes de fundadores", 1)
    preencher_linha(ws, r, lambda k, L: (
        f'=SUMPRODUCT(($E${i}:$E${fi}="Fundador")*{G}*{janela(L)}*$C${i}:$C${fi})'), BRL)
    A["fund"] = r; r += 1
    rotulo(ws, r, "Aportes de novos sócios", 1)
    preencher_linha(ws, r, lambda k, L: (
        f'=SUMPRODUCT(($E${i}:$E${fi}="Novo sócio")*{G}*{janela(L)}*$C${i}:$C${fi})'), BRL)
    A["novos"] = r; r += 1
    rotulo(ws, r, "Aportes de investidores estratégicos (lançamentos)", 1)
    preencher_linha(ws, r, lambda k, L: (
        f'=SUMPRODUCT(($E${i}:$E${fi}="Investidor estratégico")*{G}*{janela(L)}*$C${i}:$C${fi})'), BRL)
    A["estr_lanc"] = r; r += 1
    rotulo(ws, r, "Aportes da tabela de parcerias estratégicas", 1)
    preencher_linha(ws, r, lambda k, L: (
        f'=SUMPRODUCT({GE}*($D${ei}:$D${ef}>={L}${t})'
        f'*($D${ei}:$D${ef}<=EOMONTH({L}${t},0))*$C${ei}:$C${ef})'), BRL, color=VERDE_LINK)
    A["estr"] = r; r += 1
    rotulo(ws, r, "TOTAL DE APORTES NO MÊS", 0, bold=True)
    preencher_linha(ws, r, lambda k, L: f'=SUM({L}{A["fund"]}:{L}{A["estr"]})', BRL, bold=True)
    for k in range(1, N_MESES + 1):
        ws[f"{col_mes(k)}{r}"].fill = FILL_TOTAL
    A["total"] = r; r += 1
    rotulo(ws, r, "Capital acumulado", 1, italic=True)
    preencher_linha(ws, r, lambda k, L: (
        f'={L}{A["total"]}' if k == 1 else f'={gcl(COL0+k-2)}{r}+{L}{A["total"]}'),
        BRL, color=CINZA, total=False)
    A["acum"] = r; r += 1
    rotulo(ws, r, "Receita B2B das parcerias estratégicas", 1, italic=True)
    preencher_linha(ws, r, lambda k, L: (
        f'=SUMPRODUCT({GE}*($F${ei}:$F${ef}<={L}${t})'
        f'*($G${ei}:$G${ef}>={L}${t})*$E${ei}:$E${ef})'), BRL, color=VERDE_LINK)
    A["b2b"] = r
    nota(ws, f"{L_PCT}{r}", "Entra na DRE como receita B2B. Some outra parceria na tabela 2 e ela aparece aqui.")
    return ws, A
