# -*- coding: utf-8 -*-
"""Aba Ajuste Mensal: override mes a mes dos drivers volateis.

Os drivers em Premissas sao anuais — um numero para 2026, outro para 2027.
Isso serve para os anos de regime, mas nao para os primeiros meses, em que a
realidade muda a cada semana. Esta aba deixa digitar o valor de um mes
especifico sem mexer na premissa do ano.

Regra unica: celula em branco = usa o driver anual. Celula preenchida manda.
Zero digitado tambem manda — e como se forca um mes sem captacao, por exemplo.
"""
from openpyxl.styles import Alignment, PatternFill
from common import *

ABA = "Ajuste Mensal"
TEMPO = 4
AJ = {}

# chave          rotulo                                               formato
DRIVERS_CEN = [
    ("novos", "Novos usuários captados no mês",                        INT),
    ("churn", "Churn mensal da base",                                  PCT),
    ("conv",  "Taxa de conversão — patamar, antes da rampa",           PCT),
    ("cac",   "CAC — custo por usuário captado",                       BRL2),
]
DRIVERS_FIX = [
    ("cloud_fixo", "Nuvem — custo fixo mensal",                        BRL),
    ("mkt_rec",    "Marketing recorrente",                             BRL),
    ("admin",      "Despesas administrativas",                         BRL),
]
NOMES = ["CONSERVADOR", "PROVÁVEL", "AGRESSIVO"]
CORES = ["C00000", "1F3864", "1F7040"]


def ov(key, L, base):
    """Formula de override: em branco cai no `base` (o driver anual)."""
    cel = f"'{ABA}'!{L}${AJ[key]}"
    return f'IF({cel}="",{base},{cel})'


def _faixa(ws, r, fmt):
    """Escreve a faixa editavel de 60 meses, vazia."""
    for i in range(1, N_MESES + 1):
        c = ws[f"{col_mes(i)}{r}"]
        c.number_format = fmt
        c.font = f(9, False, AZUL_INPUT)
        c.fill = FILL_SUB
        c.border = BORDA_FINA
        c.alignment = Alignment(horizontal="center")
    c = ws[f"{L_TOT}{r}"]
    c.value = f'=COUNT(B{r}:{L_FIM}{r})'
    c.number_format = INT
    c.font = f(9, True, PRETO)
    c.fill = FILL_TOTAL


def build(wb, P):
    ws = wb.create_sheet(ABA)
    titulo(ws, "AJUSTE MENSAL — SOBREPOR O DRIVER ANUAL EM UM MÊS ESPECÍFICO",
           "Tudo em branco: o modelo roda só com as premissas anuais. Digite em um mês e "
           "aquele mês passa a usar o que você digitou. A coluna TOTAL conta quantos meses "
           "estão sobrepostos — é o jeito rápido de ver se sobrou alguma célula esquecida.")
    largura_padrao(ws, col_a=46)
    cabecalho_tempo(ws, linha_data=TEMPO, congelar="B9", anos_row=P["anos"])

    r = 8
    for t in ["Serve para os primeiros meses, quando o número do ano não descreve o mês: a "
              "captação do mês do lançamento, a conversão do primeiro mês de loja, o CAC de "
              "uma campanha pontual, o mês em que a nuvem muda de plano.",
              "Também serve para fechar o passado: conforme os meses acontecem, digite o "
              "realizado e o modelo projeta só o que falta.",
              "Zero é um valor, não um branco. Digitar 0 em 'novos usuários' força um mês sem "
              "captação; deixar vazio usa o driver do ano."]:
        c = ws.cell(r, 1, "•  " + t)
        c.font = f(9, False, CINZA)
        c.alignment = Alignment(wrap_text=True, vertical="top", indent=1)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=14)
        ws.row_dimensions[r].height = 26
        r += 1
    r += 1

    # ---- um bloco por cenario: estes drivers mudam de cenario para cenario ----
    for s, (nome, cor) in enumerate(zip(NOMES, CORES), start=1):
        c = ws.cell(r, 1, f"CENÁRIO {s} — {nome}")
        c.font = f(11, True, BRANCO)
        fill = PatternFill("solid", fgColor=cor)
        for cc in range(1, COL_PCT + 1):
            ws.cell(r, cc).fill = fill
        ws.row_dimensions[r].height = 18
        r += 1
        for key, lbl, fmt in DRIVERS_CEN:
            rotulo(ws, r, lbl, 1)
            AJ[f"{key}_{s}"] = r
            _faixa(ws, r, fmt)
            nota(ws, f"{L_PCT}{r}", "Meses sobrepostos (coluna TOTAL à esquerda).")
            r += 1
        r += 1

    # ---- custos: valem para os tres cenarios ----
    c = ws.cell(r, 1, "CUSTOS — VALEM PARA OS TRÊS CENÁRIOS")
    c.font = f(11, True, BRANCO)
    fill = PatternFill("solid", fgColor="808080")
    for cc in range(1, COL_PCT + 1):
        ws.cell(r, cc).fill = fill
    ws.row_dimensions[r].height = 18
    r += 1
    for key, lbl, fmt in DRIVERS_FIX:
        rotulo(ws, r, lbl, 1)
        AJ[key] = r
        _faixa(ws, r, fmt)
        r += 1
    nota(ws, f"{L_PCT}{AJ['cloud_fixo']}",
         "Continua respeitando a data de início dos custos de nuvem, em Premissas.")
    nota(ws, f"{L_PCT}{AJ['mkt_rec']}",
         "Continua respeitando a data de início do marketing recorrente.")
    nota(ws, f"{L_PCT}{AJ['admin']}",
         "Continua respeitando a data de início das despesas administrativas.")
    r += 2

    # ---- guarda-corpo: override digitado fora da janela do driver nao vale ----
    # A regra do modelo e uma so: o override entra no lugar da premissa anual, mas
    # as datas de inicio continuam valendo. Isso cria uma armadilha — digitar a
    # captacao de agosto quando o app lanca em outubro nao faz nada. Estas linhas
    # contam esses casos em vez de deixar o numero sumir em silencio.
    secao(ws, r, "CONFERÊNCIA — SOBREPOSIÇÕES QUE O MODELO VAI IGNORAR", COL_PCT); r += 1
    # O prefixo "⚠ " nao e enfeite: sem ele, "Marketing recorrente digitado..."
    # comeca igual ao rotulo do driver "Marketing recorrente" e qualquer leitor
    # que procure por prefixo confunde a linha de checagem com uma de dados.
    checks = [("novos", "lancamento", "⚠ Novos usuários digitados antes do lançamento"),
              ("conv",  "lancamento", "⚠ Conversão digitada antes do lançamento"),
              ("cloud_fixo", "cloud_ini", "⚠ Nuvem fixa digitada antes do início dos custos de nuvem"),
              ("mkt_rec", "mkt_ini", "⚠ Marketing recorrente digitado antes do início previsto"),
              ("admin", "adm_ini", "⚠ Administrativas digitadas antes do início previsto")]
    for key, pkey, lbl in checks:
        rotulo(ws, r, lbl, 1)
        linhas = ([AJ[f"{key}_{k}"] for k in (1, 2, 3)] if f"{key}_1" in AJ else [AJ[key]])
        termos = "+".join(
            f'SUMPRODUCT((B{x}:{L_FIM}{x}<>"")*(B${TEMPO}:{L_FIM}${TEMPO}'
            f'<Premissas!$B${P[pkey]}))' for x in linhas)
        c = ws.cell(r, 2, f'={termos}')
        c.number_format = INT
        c.font = f(10, True, PRETO)
        c.alignment = Alignment(horizontal="center")
        c2 = ws.cell(r, 3, f'=IF(B{r}=0,"✔ nada ignorado",'
                           f'"⚠ "&B{r}&" célula(s) sem efeito — mude a data de início '
                           f'em Premissas ou apague o valor")')
        c2.font = f(9, False, PRETO)
        c2.alignment = Alignment(indent=1, vertical="center")
        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=10)
        AJ[f"chk_{key}"] = r
        r += 1
    r += 1

    c = ws.cell(r, 1, "O pessoal não está aqui de propósito: a aba PESSOAS já é linha a linha e "
                      "mês a mês, com data de entrada e de saída de cada pessoa. É lá que se "
                      "ajusta equipe.")
    c.font = f(9, False, CINZA, italic=True)
    c.alignment = Alignment(wrap_text=True, vertical="top", indent=1)
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=14)
    ws.row_dimensions[r].height = 26

    return ws, AJ
