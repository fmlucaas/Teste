# -*- coding: utf-8 -*-
"""Aba Sócios: quanto sobra para distribuir e quanto cabe a cada um."""
from openpyxl.styles import Alignment
from openpyxl.comments import Comment
from openpyxl.worksheet.datavalidation import DataValidation
from common import *

ANOS = [2026, 2027, 2028, 2029, 2030]
CA = {a: gcl(2 + i) for i, a in enumerate(ANOS)}   # B..F
TOT = "G"

# nome, tipo, participacao %, fatia do pro-labore %
QUADRO = [
    ("Paulo Kapp",        "Fundador",   0.13, 0.25),
    ("Lucas Medina",      "Fundador",   0.13, 0.25),
    ("Felipe Toescher",   "Fundador",   0.13, 0.25),
    ("Lucas Iung",        "Fundador",   0.13, 0.25),
    ("C-ioT",             "Investidor", 0.16, 0.00),
    ("Felipe Martinelli", "Investidor", 0.16, 0.00),
    ("Shaiane",           "Investidor", 0.16, 0.00),
]
VAZIAS = 6
S = {}


def build(wb, P, A, FL, D, PE, BLOCOS):
    ws = wb.create_sheet("Sócios")
    titulo(ws, "SÓCIOS — QUANTO SOBRA PARA DIVIDIR",
           "Quanto o negócio gera de lucro, quanto disso pode virar dinheiro no bolso dos sócios, "
           "e quanto cabe a cada um. Os percentuais de participação são PLACEHOLDER — troque pelos "
           "reais quando o quadro societário estiver fechado.", largura=9)
    for col, w in zip("ABCDEFGHI", (46, 17, 17, 17, 17, 17, 18, 3, 66)):
        ws.column_dimensions[col].width = w
    ws.sheet_view.showGridLines = False

    r = 4
    # ==================================================================
    secao(ws, r, "1. QUADRO SOCIETÁRIO — preencha quando os percentuais estiverem fechados", 9)
    r += 1
    c = ws.cell(r, 1, "⚠ Os percentuais abaixo são um PLACEHOLDER que apenas reflete o que você "
                      "disse: fundadores majoritários (52% somados) e investidores com fatias "
                      "menores. Não vieram de negociação nenhuma. Participação societária depende do "
                      "valuation acordado em cada rodada, e é diferente da proporção do capital "
                      "aportado — a coluna da direita mostra o tamanho dessa diferença.")
    c.font = f(9, False, "C00000")
    c.alignment = Alignment(wrap_text=True, vertical="top", indent=1)
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=7)
    ws.row_dimensions[r].height = 42
    r += 1
    for j, h in enumerate(["Sócio", "Tipo", "Participação\nsocietária",
                           "Fatia do\npró-labore", "Capital\naportado",
                           "Proporção do\ncapital", "Prêmio sobre\no capital"]):
        c = ws.cell(r, 1 + j, h)
        c.font = f(9, True, BRANCO); c.fill = FILL_HEADER
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[r].height = 30
    r += 1
    S["q_ini"] = r
    for nome, tipo, part, prol in QUADRO + [("", "Investidor", 0.0, 0.0)] * VAZIAS:
        vazia = not nome
        ws.cell(r, 1, nome)
        ws.cell(r, 2, tipo)
        ws.cell(r, 3, part).number_format = PCT
        ws.cell(r, 4, prol).number_format = PCT
        ws.cell(r, 5, f'=IF($A{r}="",0,SUMIFS(Aportes!$C${A["ini"]}:$C${A["fim"]},'
                      f'Aportes!$B${A["ini"]}:$B${A["fim"]},$A{r})'
                      f'+SUMIFS(Aportes!$C${A["est_ini"]}:$C${A["est_fim"]},'
                      f'Aportes!$A${A["est_ini"]}:$A${A["est_fim"]},$A{r}))').number_format = BRL
        ws.cell(r, 6, f'=IFERROR($E{r}/$E${{TOTQ}},0)').number_format = PCT
        ws.cell(r, 7, f'=$C{r}-$F{r}').number_format = PCT
        for j in range(1, 8):
            cc = ws.cell(r, j)
            cc.border = BORDA_FINA
            cc.alignment = Alignment(horizontal="left" if j == 1 else "center", indent=1 if j == 1 else 0)
            cc.font = f(9, False, AZUL_INPUT if j in (1, 2, 3, 4) else PRETO)
            if j in (3, 4) and not vazia:
                cc.fill = FILL_INPUT
            elif vazia:
                cc.fill = PatternFill()
        r += 1
    S["q_fim"] = r - 1
    ws.cell(r, 1, "TOTAL").font = f(10, True)
    for col, fx, fmt in (("C", f'=SUM(C{S["q_ini"]}:C{S["q_fim"]})', PCT),
                         ("D", f'=SUM(D{S["q_ini"]}:D{S["q_fim"]})', PCT),
                         ("E", f'=SUM(E{S["q_ini"]}:E{S["q_fim"]})', BRL),
                         ("F", f'=SUM(F{S["q_ini"]}:F{S["q_fim"]})', PCT)):
        c = ws.cell(r, cidx(col), fx); c.number_format = fmt
        c.font = f(10, True); c.fill = FILL_TOTAL; c.border = BORDA_FINA
        c.alignment = Alignment(horizontal="center")
    S["q_tot"] = r
    for rr in range(S["q_ini"], S["q_fim"] + 1):
        ws.cell(rr, 6).value = ws.cell(rr, 6).value.replace("{TOTQ}", str(r))
    r += 1
    c = ws.cell(r, 1, "Verificação — a participação soma 100%?")
    c.font = f(9, True); c.alignment = Alignment(indent=1)
    c = ws.cell(r, 3, f'=IF(ROUND(C{S["q_tot"]},6)=1,"✔ SIM","⚠ NÃO — ajuste")')
    c.font = f(10, True, "C00000"); c.alignment = Alignment(horizontal="center")
    nota(ws, "I" + str(r), "Os fundadores somam 52% neste placeholder. Se a participação seguisse "
                           "o capital aportado, eles teriam 40% e seriam minoria.")
    dv = DataValidation(type="list", formula1='"Fundador,Investidor"', allow_blank=True,
                        showDropDown=False)
    ws.add_data_validation(dv); dv.add(f"B{S['q_ini']}:B{S['q_fim']}")
    r += 2

    # ==================================================================
    secao(ws, r, "2. POLÍTICA DE DISTRIBUIÇÃO", 9); r += 1
    rotulo(ws, r, "Meses de custo operacional a manter em caixa", 1)
    c = ws.cell(r, 2, 6); c.number_format = INT; c.font = f(10, False, AZUL_INPUT)
    c.fill = FILL_INPUT; c.border = BORDA_FINA
    c.alignment = Alignment(horizontal="center")
    c.comment = Comment("Reserva de segurança. Nenhuma empresa distribui até o último real: "
                        "sazonalidade, inadimplência e imprevisto exigem colchão. Seis meses de "
                        "custo operacional é uma referência conservadora e comum.", "Modelo Oppa",
                        height=110, width=340)
    S["meses_res"] = r; r += 1
    rotulo(ws, r, "% do lucro disponível distribuído no ano", 1)
    for i, a in enumerate(ANOS):
        c = ws.cell(r, 2 + i, [0.0, 0.0, 0.30, 0.50, 0.50][i])
        c.number_format = PCT; c.font = f(10, False, AZUL_INPUT)
        c.fill = FILL_INPUT; c.border = BORDA_FINA
        c.alignment = Alignment(horizontal="center")
    S["pol"] = r
    nota(ws, "I" + str(r), "Zero nos primeiros anos: o caixa é para financiar o crescimento. "
                           "A partir de 2028 distribui parte, e retém o resto.")
    r += 2

    # ==================================================================
    secao(ws, r, "3. APURAÇÃO — CENÁRIO ATIVO", 9); r += 1
    ws.cell(r, 1, "").font = f(9)
    for i, a in enumerate(ANOS):
        c = ws.cell(r, 2 + i, a)
        c.font = f(10, True, BRANCO); c.fill = FILL_HEADER
        c.alignment = Alignment(horizontal="center"); c.number_format = '0'
    c = ws.cell(r, 7, "TOTAL"); c.font = f(10, True, BRANCO); c.fill = FILL_TITULO
    c.alignment = Alignment(horizontal="center")
    S["hdr"] = r; r += 1

    def linha(lbl, fx, key, fmt=BRL, bold=False, obs=None, total=True, cor=PRETO):
        nonlocal r
        rotulo(ws, r, lbl, 0 if bold else 1, bold=bold, color=cor)
        for i, a in enumerate(ANOS):
            c = ws.cell(r, 2 + i, fx(i, CA[a], a))
            c.number_format = fmt; c.font = f(10 if bold else 9, bold, cor)
            c.border = BORDA_FINA
            if bold: c.fill = FILL_TOTAL
        if total:
            c = ws.cell(r, 7, f'=SUM(B{r}:F{r})')
            c.number_format = fmt; c.font = f(10, True); c.fill = FILL_TOTAL
            c.border = BORDA_FINA
        if obs: nota(ws, f"I{r}", obs)
        S[key] = r; r += 1

    ll_an = D["anual_ref"]["ll"]
    linha("Lucro líquido do exercício",
          lambda i, L, a: f"='DRE Projetada'!{L}{ll_an}", "ll", cor=VERDE_LINK,
          obs="Vem da DRE anual. É lucro contábil, não dinheiro em caixa.")
    linha("Lucro acumulado (compensa prejuízo dos anos anteriores)",
          lambda i, L, a: f'=B{S["ll"]}' if i == 0 else f'={gcl(1+i)}{r}+{L}{S["ll"]}',
          "ll_ac", total=False,
          obs="Prejuízo de um ano abate o lucro do seguinte antes de qualquer distribuição.")
    # reservada agora, preenchida depois que a linha de acumulado existir
    rotulo(ws, r, "Distribuições já feitas em anos anteriores", 1, color=CINZA)
    S["dist_ant"] = r
    r += 1
    linha("Lucro ainda não distribuído",
          lambda i, L, a: f'=MAX(0,{L}{S["ll_ac"]}-{L}{S["dist_ant"]})', "ll_disp", total=False)

    linha("Caixa em dezembro (antes de distribuir)",
          lambda i, L, a: f"='Análise Fluxo'!{col_am(a,12)}{FL['caixa']}", "caixa",
          cor=VERDE_LINK, total=False)
    linha("(−) Distribuições de anos anteriores",
          lambda i, L, a: f'=-{L}{S["dist_ant"]}', "caixa_dist", total=False, cor=CINZA)
    linha("(−) Reserva de caixa",
          lambda i, L, a: (f'=-$B${S["meses_res"]}*ABS(SUMIFS(\'Análise Fluxo\'!$B${FL["opex"]}:'
                           f'${L_FIM}${FL["opex"]},\'Análise Fluxo\'!$B$6:${L_FIM}$6,{L}${S["hdr"]}))/12'),
          "reserva", total=False,
          obs="Meses de custo operacional médio do ano, conforme a política acima.")
    linha("Caixa livre",
          lambda i, L, a: f'=MAX(0,{L}{S["caixa"]}+{L}{S["caixa_dist"]}+{L}{S["reserva"]})',
          "caixa_livre", total=False)

    linha("LUCRO DISTRIBUÍVEL (o menor entre lucro e caixa livre)",
          lambda i, L, a: f'=MIN({L}{S["ll_disp"]},{L}{S["caixa_livre"]})', "distvel",
          bold=True, total=False,
          obs="Só se distribui o que é lucro E está em caixa. Um dos dois limita.")
    linha("% aplicado",
          lambda i, L, a: f'={L}{S["pol"]}', "pct", PCT, total=False, cor=CINZA)
    linha("DISTRIBUÍDO NO ANO",
          lambda i, L, a: f'={L}{S["distvel"]}*{L}{S["pct"]}', "dist", bold=True)
    linha("Distribuições acumuladas",
          lambda i, L, a: f'=B{S["dist"]}' if i == 0 else f'={gcl(1+i)}{r}+{L}{S["dist"]}',
          "dist_ac", total=False)
    linha("Pró-labore pago aos sócios no ano",
          lambda i, L, a: (f'=SUMIFS(Pessoas!$B${PE["socio"]}:${L_FIM}${PE["socio"]},'
                           f'Pessoas!$B${PE["tempo"]+2}:${L_FIM}${PE["tempo"]+2},{L}${S["hdr"]})'),
          "prol", cor=VERDE_LINK,
          obs="Já está na DRE como despesa — é dinheiro que os sócios recebem trabalhando, "
              "separado da distribuição de lucros.")
    # agora que "Distribuições acumuladas" existe, preenche a linha reservada
    for i, a in enumerate(ANOS):
        c = ws.cell(S["dist_ant"], 2 + i,
                    '=0' if i == 0 else f'={gcl(1+i)}{S["dist_ac"]}')
        c.number_format = BRL; c.font = f(9, False, CINZA); c.border = BORDA_FINA
    nota(ws, f'I{S["dist_ant"]}', "O que já foi distribuído em anos anteriores não pode ser "
                                  "distribuído de novo, nem continua em caixa.")

    linha("Lucro retido na empresa",
          lambda i, L, a: f'={L}{S["ll_ac"]}-{L}{S["dist_ac"]}', "retido", total=False, cor=CINZA,
          obs="Fica na empresa para financiar crescimento — e continua sendo dos sócios.")
    r += 1

    # ==================================================================
    secao(ws, r, "4. QUANTO CABE A CADA SÓCIO", 9); r += 1
    ws.cell(r, 1, "Distribuição de lucros por sócio").font = f(10, True)
    for i, a in enumerate(ANOS):
        c = ws.cell(r, 2 + i, a)
        c.font = f(10, True, BRANCO); c.fill = FILL_HEADER
        c.alignment = Alignment(horizontal="center"); c.number_format = '0'
    c = ws.cell(r, 7, "TOTAL 5 ANOS")
    c.font = f(10, True, BRANCO); c.fill = FILL_TITULO
    c.alignment = Alignment(horizontal="center")
    hdr2 = r; r += 1
    S["p_ini"] = r
    for k in range(S["q_ini"], S["q_fim"] + 1):
        ws.cell(r, 1, f'=IF($A{k}="","",$A{k})').font = f(9, True)
        ws.cell(r, 1).alignment = Alignment(indent=1)
        for i, a in enumerate(ANOS):
            c = ws.cell(r, 2 + i, f'=$C{k}*{CA[a]}${S["dist"]}')
            c.number_format = BRL; c.font = f(9); c.border = BORDA_FINA
        c = ws.cell(r, 7, f'=SUM(B{r}:F{r})')
        c.number_format = BRL; c.font = f(9, True); c.fill = FILL_TOTAL; c.border = BORDA_FINA
        r += 1
    S["p_fim"] = r - 1
    ws.cell(r, 1, "TOTAL DISTRIBUÍDO").font = f(10, True)
    for i, a in enumerate(ANOS):
        c = ws.cell(r, 2 + i, f'=SUM({CA[a]}{S["p_ini"]}:{CA[a]}{S["p_fim"]})')
        c.number_format = BRL; c.font = f(10, True); c.fill = FILL_TOTAL; c.border = BORDA_FINA
    c = ws.cell(r, 7, f'=SUM(G{S["p_ini"]}:G{S["p_fim"]})')
    c.number_format = BRL; c.font = f(10, True); c.fill = FILL_TOTAL; c.border = BORDA_FINA
    r += 2

    # --- total recebido, somando pró-labore ------------------------------
    ws.cell(r, 1, "Total recebido por sócio nos 5 anos").font = f(10, True)
    for j, h in enumerate(["Distribuição de lucros", "Pró-labore", "TOTAL RECEBIDO",
                           "Por 1% de participação"]):
        c = ws.cell(r, 2 + j, h)
        c.font = f(9, True, BRANCO); c.fill = FILL_HEADER
        c.alignment = Alignment(horizontal="center", wrap_text=True)
    ws.row_dimensions[r].height = 28
    r += 1
    S["r_ini"] = r
    for idx, k in enumerate(range(S["q_ini"], S["q_fim"] + 1)):
        ws.cell(r, 1, f'=IF($A{k}="","",$A{k})').font = f(9, True)
        ws.cell(r, 1).alignment = Alignment(indent=1)
        ws.cell(r, 2, f'=$G{S["p_ini"]+idx}').number_format = BRL
        ws.cell(r, 3, f'=$D{k}*$G${S["prol"]}').number_format = BRL
        ws.cell(r, 4, f'=$B{r}+$C{r}').number_format = BRL
        ws.cell(r, 5, f'=IFERROR($B{r}/$C{k}/100,0)').number_format = BRL
        for j in range(1, 6):
            cc = ws.cell(r, j); cc.border = BORDA_FINA
            cc.font = f(9, j == 4)
            cc.alignment = Alignment(horizontal="left" if j == 1 else "center", indent=1 if j == 1 else 0)
        ws.cell(r, 4).fill = FILL_TOTAL
        r += 1
    S["r_fim"] = r - 1
    r += 1
    return ws, S
