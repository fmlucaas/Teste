# -*- coding: utf-8 -*-
"""Aba Investimentos: cronograma de CAPEX e depreciacao/amortizacao."""
import datetime as dt
from openpyxl.styles import Alignment
from common import *

CAPEX = [
    (dt.date(2026, 5,  4), "Constituição da empresa, registro de marca e assessoria jurídica",
     "Marca e jurídico", 6000.00, "Previsto"),
    (dt.date(2026, 8, 20), "2 notebooks Lenovo", "Equipamentos", 9127.82, "Realizado"),
    (dt.date(2026, 8, 31), "Desenvolvimento do MVP — parcela 1/3", "Desenvolvimento", 7000.00, "Contratado"),
    (dt.date(2026, 9, 30), "Desenvolvimento do MVP — parcela 2/3", "Desenvolvimento", 7000.00, "Contratado"),
    (dt.date(2026,10, 31), "Desenvolvimento do MVP — parcela 3/3", "Desenvolvimento", 7000.00, "Contratado"),
    (dt.date(2027, 1, 31), "Release v2 — atualizações de janeiro", "Desenvolvimento", 15000.00, "Previsto"),
    (dt.date(2027, 3, 31), "Adequação LGPD e segurança da informação", "Certificações", 12000.00, "Previsto"),
    (dt.date(2027, 6, 30), "Equipamentos — expansão da equipe", "Equipamentos", 6000.00, "Previsto"),
    (dt.date(2027, 7, 31), "Certificação de conformidade em saúde (SBIS / ANVISA se houver device)",
     "Certificações", 25000.00, "Previsto"),
    (dt.date(2028, 1, 31), "Evolução da plataforma — módulo marketplace", "Desenvolvimento", 45000.00, "Previsto"),
    (dt.date(2028, 6, 30), "Equipamentos", "Equipamentos", 15000.00, "Previsto"),
    (dt.date(2029, 1, 31), "Evolução da plataforma", "Desenvolvimento", 60000.00, "Previsto"),
    (dt.date(2029, 6, 30), "Equipamentos", "Equipamentos", 25000.00, "Previsto"),
    (dt.date(2030, 1, 31), "Evolução da plataforma", "Desenvolvimento", 70000.00, "Previsto"),
    (dt.date(2030, 6, 30), "Equipamentos", "Equipamentos", 30000.00, "Previsto"),
]
CATS = ["Desenvolvimento", "Equipamentos", "Marca e jurídico", "Certificações"]
I = {}


def build(wb, P):
    ws = wb.create_sheet("Investimentos")
    titulo(ws, "INVESTIMENTOS (CAPEX) — CRONOGRAMA E AMORTIZAÇÃO",
           "Desembolsos de capital. Saem do caixa no mês do lançamento (aba Análise Fluxo) e são "
           "amortizados linearmente em 60 meses na DRE. Verde = já realizado.")
    largura_padrao(ws, col_a=46)
    for col, w in zip("ABCDE", (16, 56, 20, 16, 14)):
        ws.column_dimensions[col].width = w

    r = 4
    secao(ws, r, "LANÇAMENTOS DE INVESTIMENTO", 6); r += 1
    for j, h in enumerate(["Data", "Item", "Categoria", "Valor", "Status"]):
        c = ws.cell(r, 1 + j, h)
        c.font = f(9, True, BRANCO); c.fill = FILL_HEADER
        c.alignment = Alignment(horizontal="center")
    r += 1
    I["ini"] = r
    for (d, item, cat, val, st) in sorted(CAPEX, key=lambda x: x[0]):
        ws.cell(r, 1, dt.datetime(d.year, d.month, d.day)).number_format = DATA
        ws.cell(r, 2, item)
        ws.cell(r, 3, cat)
        ws.cell(r, 4, val).number_format = BRL2
        ws.cell(r, 5, st)
        cor = AZUL_INPUT if st == "Realizado" else PRETO
        for j in range(1, 6):
            cc = ws.cell(r, j)
            cc.font = f(9, st == "Realizado", cor)
            cc.border = BORDA_FINA
            cc.alignment = Alignment(horizontal="center" if j != 2 else "left", indent=1)
            if st == "Realizado":
                cc.fill = FILL_OK
        r += 1
    I["fim"] = r - 1
    ws.cell(r, 3, "TOTAL").font = f(10, True)
    c = ws.cell(r, 4, f'=SUM(D{I["ini"]}:D{I["fim"]})')
    c.number_format = BRL2; c.font = f(10, True); c.fill = FILL_TOTAL
    nota(ws, f"F{I['ini']+1}", "Fatos do brief: 2 notebooks Lenovo por R$ 9.127,82.")
    nota(ws, f"F{I['ini']+2}", "Orçamento de R$ 21.000 da equipe PJ para entregar o MVP até outubro.")
    r += 2

    # --- Cronograma mensal ------------------------------------------------
    cabecalho_tempo(ws, linha_data=r, congelar=f"B{r+5}")
    I["tempo"] = r
    r += 4
    secao(ws, r, "CRONOGRAMA MENSAL DE DESEMBOLSO", COL_PCT); r += 1
    primeiro = r
    for cat in CATS:
        rotulo(ws, r, cat, 1)
        preencher_linha(ws, r, lambda i, L, c=cat: (
            f'=SUMIFS($D${I["ini"]}:$D${I["fim"]},$C${I["ini"]}:$C${I["fim"]},"{c}",'
            f'$A${I["ini"]}:$A${I["fim"]},">="&{L}${I["tempo"]},'
            f'$A${I["ini"]}:$A${I["fim"]},"<="&EOMONTH({L}${I["tempo"]},0))'), BRL)
        I[cat] = r
        r += 1
    rotulo(ws, r, "TOTAL DE INVESTIMENTOS", 0, bold=True)
    preencher_linha(ws, r, lambda i, L: f'=SUM({L}{primeiro}:{L}{r-1})', BRL, bold=True)
    for i in range(1, N_MESES + 1):
        ws[f"{col_mes(i)}{r}"].fill = FILL_TOTAL
    I["total"] = r
    for cat in CATS:
        ws[f"{L_PCT}{I[cat]}"] = f'=IFERROR({L_TOT}{I[cat]}/{L_TOT}${I["total"]},0)'
        ws[f"{L_PCT}{I[cat]}"].number_format = PCT
        ws[f"{L_PCT}{I[cat]}"].font = f(9, False, CINZA)
    r += 2

    # --- Amortizacao ------------------------------------------------------
    secao(ws, r, "DEPRECIAÇÃO E AMORTIZAÇÃO (linear, 60 meses)", COL_PCT); r += 1
    rotulo(ws, r, "Vida útil adotada (meses)", 1)
    c = ws.cell(r, 2, 60); c.number_format = INT; c.font = f(10, False, AZUL_INPUT)
    c.fill = FILL_SUB; c.border = BORDA_FINA
    c.alignment = Alignment(horizontal="center")
    nota(ws, "D" + str(r), "Equipamentos de informática: 5 anos (20% a.a., IN RFB 1.700). "
                           "Software e intangíveis amortizados no mesmo prazo.")
    I["vida"] = r; r += 1

    rotulo(ws, r, "Investimento acumulado (base amortizável)", 1)
    preencher_linha(ws, r, lambda i, L: f'=SUM($B{I["total"]}:{L}{I["total"]})', BRL, color=CINZA)
    I["acum"] = r; r += 1

    rotulo(ws, r, "Depreciação e amortização do mês", 1, bold=True)
    preencher_linha(ws, r, lambda i, L: (
        '=0' if i == 1 else f'={gcl(COL0+i-2)}{I["acum"]}/$B${I["vida"]}'), BRL, bold=True)
    I["da"] = r; r += 1
    nota(ws, f"{L_PCT}{I['da']}", "Cada item começa a amortizar no mês seguinte ao desembolso.")

    rotulo(ws, r, "Saldo contábil do ativo (imobilizado + intangível)", 1, italic=True)
    preencher_linha(ws, r, lambda i, L: (
        f'={L}{I["total"]}' if i == 1 else f'={gcl(COL0+i-2)}{r}+{L}{I["total"]}-{L}{I["da"]}'),
        BRL, color=CINZA, total=False)
    I["saldo"] = r
    return ws, I
