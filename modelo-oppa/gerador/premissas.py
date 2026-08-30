# -*- coding: utf-8 -*-
"""Aba Premissas: todos os inputs do modelo, com registro de linhas."""
import datetime as dt
from openpyxl.styles import Alignment, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment
from common import *

ANOS = [2026, 2027, 2028, 2029, 2030]
C_ANO = {a: gcl(3 + i) for i, a in enumerate(ANOS)}   # C..G

P = {}


def _in(ws, lin, col, val, fmt, chave=False, fonte=None):
    c = ws.cell(lin, col, val)
    c.number_format = fmt
    c.font = f(10, False, AZUL_INPUT)
    c.fill = FILL_INPUT if chave else FILL_SUB
    c.border = BORDA_FINA
    c.alignment = Alignment(horizontal="center")
    if fonte:
        c.comment = Comment(fonte, "Modelo Oppa", height=120, width=340)
    return c


def linha_input(ws, lin, label, valores, fmt, key, nivel=1, chave=False,
                fonte=None, obs=None):
    rotulo(ws, lin, label, nivel)
    for i in range(5):
        v = valores[i] if isinstance(valores, (list, tuple)) else valores
        _in(ws, lin, 3 + i, v, fmt, chave=chave, fonte=fonte)
    if obs:
        nota(ws, f"I{lin}", obs)
    P[key] = lin
    return lin


def linha_evolucao(ws, lin, label, base, evol, fmt, key, chave=False,
                   fonte=None, obs=None, teto=None):
    """Coluna B = evolução anual · C = patamar de 2026 · D..G calculados."""
    rotulo(ws, lin, label, 1)
    _in(ws, lin, 2, evol, PCT, chave=False)
    _in(ws, lin, 3, base, fmt, chave=chave, fonte=fonte)
    for i in range(1, 5):
        col = 3 + i
        ant = gcl(col - 1)
        fx = f'={ant}{lin}*(1+$B{lin})'
        if teto:
            fx = f'=MIN({teto},{ant}{lin}*(1+$B{lin}))'
        c = ws.cell(lin, col, fx)
        c.number_format = fmt
        c.font = f(10, False, PRETO)
        c.border = BORDA_FINA
        c.alignment = Alignment(horizontal="center")
    if obs:
        nota(ws, f"I{lin}", obs)
    P[key] = lin
    return lin


def linha_unica(ws, lin, label, valor, fmt, key, nivel=1, chave=False,
                fonte=None, obs=None, formula=False):
    rotulo(ws, lin, label, nivel)
    c = ws.cell(lin, 2, valor)
    c.number_format = fmt
    c.font = f(10, formula, PRETO if formula else AZUL_INPUT)
    c.fill = PatternFill() if formula else (FILL_INPUT if chave else FILL_SUB)
    c.border = BORDA_FINA
    c.alignment = Alignment(horizontal="center")
    if fonte:
        c.comment = Comment(fonte, "Modelo Oppa", height=120, width=340)
    if obs:
        nota(ws, f"D{lin}", obs)
    P[key] = lin
    return lin


def build(wb):
    ws = wb.create_sheet("Premissas")
    titulo(ws, "OPPA — PREMISSAS DO MODELO",
           "AZUL = você digita · fundo AMARELO = premissa-chave · PRETO = calculado pelo modelo. "
           "Células com marca vermelha no canto têm a fonte do número em comentário. "
           "O quadro de pessoas fica na aba Pessoas.",
           largura=12)
    ws.column_dimensions["A"].width = 54
    ws.column_dimensions["B"].width = 17
    for a in ANOS:
        ws.column_dimensions[C_ANO[a]].width = 14
    ws.column_dimensions["H"].width = 3
    ws.column_dimensions["I"].width = 72

    r = 4
    # ==================================================================
    secao(ws, r, "1. CONTROLES DO MODELO", 12); r += 1
    rotulo(ws, r, "Ano →", 0, bold=True)
    ws.cell(r, 2, "Evolução a.a.").font = f(9, True, BRANCO)
    ws.cell(r, 2).fill = FILL_HEADER
    ws.cell(r, 2).alignment = Alignment(horizontal="center")
    for i, a in enumerate(ANOS):
        c = ws.cell(r, 3 + i, a)
        c.font = f(10, True, BRANCO); c.fill = FILL_HEADER
        c.alignment = Alignment(horizontal="center"); c.number_format = '0'
    P["anos"] = r
    r += 1

    linha_unica(ws, r, "Cenário ativo  (1 = Conservador · 2 = Provável · 3 = Agressivo)",
                2, '0', "cenario", chave=True,
                obs="Troque este número e a planilha inteira recalcula no cenário escolhido."); r += 1
    rotulo(ws, r, "Cenário ativo — nome", 1)
    c = ws.cell(r, 2, f'=CHOOSE($B${P["cenario"]},"CONSERVADOR","PROVÁVEL","AGRESSIVO")')
    c.font = f(10, True, VERDE_LINK); c.alignment = Alignment(horizontal="center")
    P["cenario_nome"] = r; r += 1

    linha_unica(ws, r, "Considerar o plano de contratações?  (1 = Sim · 0 = Não)", 1, '0',
                "plano", chave=True,
                fonte=("Aba Pessoas. Com 0, só entram no modelo as pessoas com status 'Contratado' — "
                       "hoje, apenas a equipe do MVP de set–out/2026. Com 1, entra também o plano de "
                       "contratações futuras, que é sugestão a validar, não compromisso."),
                obs="Com 0, o modelo considera só quem já está contratado (equipe do MVP)."); r += 2

    rotulo(ws, r, "LANÇAMENTO DO APP", 0, bold=True); r += 1
    linha_unica(ws, r, "Ano de lançamento", 2026, '0', "lanc_ano", chave=True); r += 1
    linha_unica(ws, r, "Mês de lançamento  (1 = janeiro … 12 = dezembro)", 10, '0',
                "lanc_mes", chave=True,
                obs="Brief: app majoritariamente pronto em outubro."); r += 1
    rotulo(ws, r, "Data de lançamento (calculada)", 1)
    c = ws.cell(r, 2, f'=DATE($B${P["lanc_ano"]},$B${P["lanc_mes"]},1)')
    c.number_format = MES; c.font = f(10, True, PRETO)
    c.alignment = Alignment(horizontal="center"); c.border = BORDA_FINA
    P["lancamento"] = r; r += 1
    rotulo(ws, r, "Mês nº do lançamento na linha do tempo", 1, italic=True)
    c = ws.cell(r, 2, f'=($B${P["lanc_ano"]}-2026)*12+$B${P["lanc_mes"]}')
    c.number_format = INT; c.font = f(9, False, CINZA)
    c.alignment = Alignment(horizontal="center")
    P["lanc_idx"] = r; r += 1
    linha_unica(ws, r, "Rampa de maturação da conversão (meses)", 6, INT, "rampa_conv", chave=True,
                fonte=("Um app recém-lançado não converte no patamar de regime no primeiro mês. "
                       "A conversão cresce linearmente do lançamento até atingir, em N meses, o "
                       "patamar definido para o ano. Com 6 meses e patamar de 5%, outubro/2026 "
                       "converte a 0,83%, novembro a 1,67% e dezembro a 2,50%."),
                obs="É isto que torna 2026 mais tímido. Coloque 1 para converter no patamar cheio "
                    "desde o primeiro mês."); r += 2

    rotulo(ws, r, "PARÂMETROS FINANCEIROS", 0, bold=True); r += 1
    linha_unica(ws, r, "TMA — taxa mínima de atratividade (a.a.)", 0.25, PCT, "tma", chave=True,
                fonte=("Taxa de desconto. A planilha original usava 16% a.a. Para startup pré-receita "
                       "a faixa usual é 25%–40% a.a. Ajustável."),
                obs="Original usava 16% a.a.; 25% reflete risco de startup pré-receita."); r += 1
    rotulo(ws, r, "TMA equivalente mensal", 1)
    c = ws.cell(r, 2, f'=(1+$B${P["tma"]})^(1/12)-1'); c.number_format = PCT2
    c.font = f(10, True, PRETO); c.alignment = Alignment(horizontal="center")
    P["tma_m"] = r; r += 1
    linha_unica(ws, r, "Payback desejado (meses) — critério de decisão", 36, INT, "payback_alvo",
                obs="Critério Go/No-Go, herdado da planilha de viabilidade da EMS."); r += 2

    # ==================================================================
    secao(ws, r, "2. TRIBUTAÇÃO", 12); r += 1
    linha_unica(ws, r, "Regime  (1 = Simples Anexo III · 2 = Simples Anexo V · 3 = Lucro Presumido)",
                2, '0', "regime", chave=True,
                fonte=("ATENÇÃO — Fator R. Licenciamento/cessão de uso de software é Anexo V, migrando "
                       "para Anexo III só se o Fator R (folha + pró-labore dos últimos 12 meses ÷ RBT12) "
                       "for >= 28%. Pagamento a equipe PJ NÃO conta como folha. A aba Pessoas calcula "
                       "o Fator R do seu quadro."),
                obs="⚠ Padrão = 2 (Anexo V). O Fator R do quadro atual fica muito abaixo de 28%, então o Anexo III não está disponível. Veja a conferência na aba Pessoas."); r += 1
    linha_unica(ws, r, "Alíquota Lucro Presumido / acima do teto do Simples", 0.1633, PCT2, "lp",
                fonte=("Serviços: PIS 0,65% + COFINS 3,00% + ISS 5,00% + IRPJ 4,80% (32% de presunção × "
                       "15%) + CSLL 2,88% (32% × 9%) = 16,33% da receita bruta. ISS varia de 2% a 5% "
                       "conforme o município."),
                obs="PIS 0,65 + COFINS 3,00 + ISS 5,00 + IRPJ 4,80 + CSLL 2,88 = 16,33%"); r += 1
    linha_unica(ws, r, "Teto do Simples Nacional (RBT12)", 4800000, BRL, "teto_simples",
                obs="LC 123/2006. Acima disso o modelo passa sozinho ao Lucro Presumido."); r += 2

    # ==================================================================
    secao(ws, r, "3. PREÇOS E MIX DE PLANOS", 12); r += 1
    c = ws.cell(r, 1, "•  São quatro planos, conforme o backlog do produto: Free (gratuito, entra na "
                      "base de usuários mas não gera receita de assinatura), Essencial, Premium e "
                      "Família. Os três pagos estão abaixo; o Free é a diferença entre a base ativa "
                      "e os pagantes, na aba Usuários.")
    c.font = f(9, False, CINZA)
    c.alignment = Alignment(wrap_text=True, vertical="top", indent=1)
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=9)
    ws.row_dimensions[r].height = 28
    r += 1
    linha_unica(ws, r, "Plano Essencial (R$/mês)  ·  1 familiar, até 3 cuidadores",
                19.90, BRL2, "p_basico", chave=True); r += 1
    linha_unica(ws, r, "Plano Premium (R$/mês)  ·  1 familiar, cuidadores ilimitados",
                39.90, BRL2, "p_inter", chave=True); r += 1
    linha_unica(ws, r, "Plano Família (R$/mês)  ·  até 4 familiares, cuidadores ilimitados",
                79.90, BRL2, "p_premium", chave=True); r += 1
    linha_unica(ws, r, "Reajuste anual de preços (a.a.)", 0.045, PCT, "reajuste",
                obs="Aplicado a partir de 2027 (IPCA projetado ~4,5%)."); r += 1

    rotulo(ws, r, "Mix de pagantes por plano (% dos assinantes)", 0, bold=True); r += 1
    linha_input(ws, r, "Plano Essencial", [0.75, 0.72, 0.68, 0.64, 0.60], PCT, "mix_b", chave=True,
        fonte=("Em assinatura B2C de ticket baixo a decisão é 'pagar ou não pagar', não 'qual plano' — "
               "o degrau de entrada concentra a base. Referências de mercado põem o plano de entrada em "
               "60–75% dos assinantes. O efeito compromisso (Simonson & Tversky, 1992) puxa parte da "
               "base para o plano do meio ao longo do tempo. Daí o mix migrar de 75/20/5 para 60/28/12."),
        obs="Efeito compromisso (Simonson & Tversky, 1992) + padrão de assinatura B2C de ticket baixo."); r += 1
    linha_input(ws, r, "Plano Premium", [0.20, 0.22, 0.24, 0.26, 0.28], PCT, "mix_i", chave=True,
        obs="Plano-ponte: cresce conforme o Premium ganha valor percebido."); r += 1
    linha_input(ws, r, "Plano Família", [0.05, 0.06, 0.08, 0.10, 0.12], PCT, "mix_p", chave=True,
        obs="Fatia menor mas crescente — ancoragem de preço que sustenta o Intermediário."); r += 1
    rotulo(ws, r, "Verificação do mix (deve somar 100%)", 2, italic=True)
    for i, a in enumerate(ANOS):
        L = C_ANO[a]
        c = ws.cell(r, 3 + i, f'=SUM({L}{P["mix_b"]}:{L}{P["mix_p"]})')
        c.number_format = PCT; c.font = f(9, True, PRETO)
        c.alignment = Alignment(horizontal="center")
    P["mix_chk"] = r; r += 1
    rotulo(ws, r, "ARPPU implícita (receita média por assinante pagante)", 2, italic=True)
    for i, a in enumerate(ANOS):
        L = C_ANO[a]
        c = ws.cell(r, 3 + i,
            f'=({L}{P["mix_b"]}*$B${P["p_basico"]}+{L}{P["mix_i"]}*$B${P["p_inter"]}'
            f'+{L}{P["mix_p"]}*$B${P["p_premium"]})*(1+$B${P["reajuste"]})^({L}${P["anos"]}-2026)')
        c.number_format = BRL2; c.font = f(9, True, VERDE_LINK)
        c.alignment = Alignment(horizontal="center")
    P["arppu"] = r; r += 2

    # ==================================================================
    secao(ws, r, "4. DRIVERS POR CENÁRIO", 12); r += 1
    for t in ["Cada driver tem 3 linhas — Conservador, Provável e Agressivo. O modelo usa a do cenário ativo.",
              "Digite o patamar de 2026 na coluna C e a taxa de evolução anual na coluna B: 2027 a 2030 "
              "são calculados sozinhos. Quer um ano específico fora da curva? Digite por cima — a célula "
              "vira entrada manual e os anos seguintes continuam a partir dela."]:
        c = ws.cell(r, 1, "•  " + t)
        c.font = f(9, False, CINZA)
        c.alignment = Alignment(wrap_text=True, vertical="top", indent=1)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=9)
        ws.row_dimensions[r].height = 26
        r += 1

    rotulo(ws, r, "Novos usuários captados por mês (média do ano)", 0, bold=True); r += 1
    linha_evolucao(ws, r, "Conservador", 10860, 0.40, INT, "novos_1"); r += 1
    linha_evolucao(ws, r, "Provável", 18100, 0.50, INT, "novos_2", chave=True,
        fonte=("Calibrado para a meta do brief: ~50.000 usuários na base ao final de 2026. Com "
               "lançamento em out/2026 e churn de 8%/mês, 18.100 novos/mês em out+nov+dez levam a "
               "uma base de ~50.100 em dez/2026."),
        obs="Calibrado para a meta de ~50.000 usuários em dez/2026."); r += 1
    linha_evolucao(ws, r, "Agressivo", 25340, 0.60, INT, "novos_3"); r += 1

    rotulo(ws, r, "Churn mensal da base de usuários", 0, bold=True); r += 1
    linha_evolucao(ws, r, "Conservador", 0.100, -0.04, PCT, "churn_1"); r += 1
    linha_evolucao(ws, r, "Provável", 0.080, -0.06, PCT, "churn_2", chave=True,
        fonte=("Benchmarks 2025/2026 de apps de assinatura em Saúde & Fitness: churn mensal de 7% a 10% "
               "(mediana ~9,2%); os melhores ficam abaixo de 5%. Adotado 8% em 2026 melhorando 6% ao ano, "
               "assumindo que monitoramento de saúde retém melhor que fitness (necessidade médica, não "
               "motivação). Fontes: RevenueCat State of Subscription Apps 2025; Business of Apps."),
        obs="Saúde & Fitness: 7–10%/mês (mediana 9,2%). Evolução negativa = churn melhorando."); r += 1
    linha_evolucao(ws, r, "Agressivo", 0.060, -0.08, PCT, "churn_3"); r += 1

    rotulo(ws, r, "Taxa de conversão (pagantes ÷ base ativa) — patamar de regime", 0, bold=True); r += 1
    linha_evolucao(ws, r, "Conservador", 0.030, 0.08, PCT, "conv_1", teto=0.25); r += 1
    linha_evolucao(ws, r, "Provável", 0.050, 0.08, PCT, "conv_2", chave=True, teto=0.25,
        fonte=("Premissa do brief: 5% de pagantes. Benchmarks de freemium: mediana 2,2%–2,6%, faixa "
               "típica 2%–5%, quartil superior 5%–8%. Este é o patamar de REGIME: nos primeiros meses "
               "após o lançamento a conversão sobe pela rampa definida no controle 1, então 2026 sai "
               "bem abaixo de 5%. Fontes: RevenueCat 2025; Userpilot; Geneo."),
        obs="Patamar de regime. A rampa (controle 1) segura a conversão nos primeiros meses."); r += 1
    linha_evolucao(ws, r, "Agressivo", 0.070, 0.10, PCT, "conv_3", teto=0.25); r += 1

    rotulo(ws, r, "CAC — custo de aquisição por novo usuário (R$)", 0, bold=True); r += 1
    linha_evolucao(ws, r, "Conservador", 4.00, 0.18, BRL2, "cac_1"); r += 1
    linha_evolucao(ws, r, "Provável", 2.50, 0.15, BRL2, "cac_2", chave=True,
        fonte=("CPI (custo por instalação) no Brasil em Meta/Google Ads para apps de saúde: R$ 1,50–4,00. "
               "Sobe ao longo do tempo conforme se esgotam as audiências baratas. A agência da Shaiane "
               "reduz o custo de produção criativa, não o custo de mídia."),
        obs="CPI Brasil saúde: R$ 1,50–4,00. Sobe com a escala."); r += 1
    linha_evolucao(ws, r, "Agressivo", 1.50, 0.15, BRL2, "cac_3"); r += 1

    linha_input(ws, r, "Aquisição orgânica (% dos novos usuários, sem custo de mídia)",
                [0.30, 0.32, 0.35, 0.38, 0.40], PCT, "organico", nivel=0,
        obs="Boca a boca + conteúdo da agência da Shaiane. Reduz a fatia paga com mídia."); r += 2

    # ==================================================================
    secao(ws, r, "5. MONETIZAÇÃO — CANAIS, TAXAS E RECEITA B2B", 12); r += 1
    linha_unica(ws, r, "Taxa das lojas de aplicativos (App Store / Google Play)", 0.15, PCT,
        "taxa_loja", chave=True,
        fonte=("Google Play cobra 15% em todas as assinaturas auto-renováveis desde o 1º dia. Apple "
               "cobra 30%, reduzidos a 15% no App Store Small Business Program (receita anual < US$ 1 mi) "
               "ou a partir do 2º ano de assinatura. Acima de US$ 1 mi/ano a Apple volta a 30%."),
        obs="Google Play 15% desde o 1º dia; Apple 15% no Small Business Program."); r += 1
    linha_input(ws, r, "% da receita de assinaturas cobrada pelas lojas",
                [0.85, 0.80, 0.75, 0.72, 0.70], PCT, "share_loja",
        obs="Migração gradual para checkout web/PIX próprio, que não paga comissão de loja."); r += 1
    linha_unica(ws, r, "Taxa de meio de pagamento (venda direta web/PIX)", 0.045, PCT, "taxa_pgto",
        obs="Média ponderada de cartão (~4,99%) e PIX (~1,0%)."); r += 1


    rotulo(ws, r, "Receita B2B de dados — o painel", 0, bold=True); r += 1
    for t in ["O dado só vale se houver escala e consentimento. Estas três premissas definem o "
              "tamanho do painel comercializável; a tabela logo abaixo define o que se cobra por ele. "
              "Enquanto o painel não atingir o mínimo, nenhum contrato de dados gera receita, por "
              "mais que a data de início já tenha passado."]:
        c = ws.cell(r, 1, "•  " + t)
        c.font = f(9, False, CINZA)
        c.alignment = Alignment(wrap_text=True, vertical="top", indent=1)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=9)
        ws.row_dimensions[r].height = 28
        r += 1
    linha_input(ws, r, "% da base que consente compartilhar dados",
                [0.25, 0.28, 0.30, 0.32, 0.35], PCT, "consent", chave=True,
        fonte=("LGPD: dado de saúde é dado sensível (art. 11). Vender exige anonimização efetiva "
               "(art. 12, que tira o dado anonimizado do alcance da lei) OU consentimento "
               "ESPECÍFICO e DESTACADO para esta finalidade — separado do aceite de uso do app, "
               "revogável a qualquer tempo. Quem revoga sai do painel. Taxas de opt-in em apps de "
               "saúde com pedido bem desenhado ficam na casa de 20% a 40%."),
        obs="⚠ Consentimento específico e destacado, separado dos termos do app. Revogável."); r += 1
    linha_unica(ws, r, "Base mínima do painel para ter valor comercial (usuários)", 75000, INT,
                "painel_min", chave=True,
        fonte=("Painel pequeno não se vende: comprador de dado quer representatividade. 75 mil "
               "usuários consentidos é um piso defensável para um painel nichado — idoso com "
               "doença crônica, monitorado continuamente — que é estreito mas profundo. Abaixo "
               "disso o modelo zera a receita de dados."),
        obs="Abaixo deste número, a receita de dados é zero, mesmo com contrato na tabela."); r += 1
    linha_unica(ws, r, "Custo de processamento do painel (R$/usuário/mês)", 0.05, BRL2,
                "custo_painel",
        obs="Anonimização, agregação e entrega dos dados. Entra na linha de nuvem."); r += 1
    linha_unica(ws, r, "Considerar iniciativas B2B EM ESTUDO?  (1 = Sim · 0 = Não)", 0, '0',
                "inc_estudo", chave=True,
        obs="⚠ Desligado. Liga as linhas marcadas como 'Em estudo' na tabela abaixo."); r += 1

    rotulo(ws, r, "Iniciativas de receita B2B", 0, bold=True); r += 1
    c = ws.cell(r, 1, "Uma linha por iniciativa. As duas primeiras são a venda de dados, já "
                      "planejada. As demais estão em branco para vocês preencherem — recompra de "
                      "medicamento de uso crônico, integrações com farmácia, o que vier. Categoria "
                      "separa o que aparece como 'venda de dados' do que aparece como 'outras "
                      "iniciativas' na DRE.")
    c.font = f(9, False, CINZA, italic=True)
    c.alignment = Alignment(wrap_text=True, vertical="top", indent=1)
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=9)
    ws.row_dimensions[r].height = 30
    r += 1
    for j, h in enumerate(["Iniciativa", "Categoria", "Status", "Modelo", "Valor",
                           "Início", "Fim"]):
        c = ws.cell(r, 1 + j, h)
        c.font = f(9, True, BRANCO); c.fill = FILL_HEADER
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    r += 1
    P["b2b_ini"] = r
    linhas_b2b = [
        ("Venda de dados à indústria farmacêutica", "Dados", "Planejado", "Fixo mensal",
         45000, dt.datetime(2028, 7, 1), dt.datetime(2030, 12, 1)),
        ("Venda de dados a redes de farmácia", "Dados", "Planejado", "Fixo mensal",
         25000, dt.datetime(2029, 1, 1), dt.datetime(2030, 12, 1)),
        ("Recompra de medicamento de uso crônico", "Outras", "Em estudo", "Fixo mensal",
         0, dt.datetime(2029, 1, 1), dt.datetime(2030, 12, 1)),
    ] + [("", "Outras", "Em estudo", "Fixo mensal", 0,
          dt.datetime(2029, 1, 1), dt.datetime(2030, 12, 1)) for _ in range(6)]
    for nome, cat, st, mod, val, ini_d, fim_d in linhas_b2b:
        vazia = not nome
        ws.cell(r, 1, nome)
        ws.cell(r, 2, cat)
        ws.cell(r, 3, st)
        ws.cell(r, 4, mod)
        ws.cell(r, 5, val).number_format = BRL2
        ws.cell(r, 6, ini_d).number_format = MES
        ws.cell(r, 7, fim_d).number_format = MES
        for j in range(1, 8):
            cc = ws.cell(r, j)
            cc.border = BORDA_FINA
            cc.font = f(9, False, CINZA if vazia else AZUL_INPUT)
            cc.alignment = Alignment(horizontal="left" if j == 1 else "center",
                                     indent=1 if j == 1 else 0)
            if not vazia:
                cc.fill = FILL_OK if st == "Planejado" else FILL_ALERTA
        r += 1
    P["b2b_fim"] = r - 1
    for col, opts in (("B", '"Dados,Outras"'), ("C", '"Planejado,Em estudo,Descartado"'),
                      ("D", '"Fixo mensal,Por usuário"')):
        dv = DataValidation(type="list", formula1=opts, allow_blank=True, showDropDown=False)
        ws.add_data_validation(dv)
        dv.add(f"{col}{P['b2b_ini']}:{col}{P['b2b_fim']}")
    nota(ws, f"I{P['b2b_ini']}",
         "Modelo 'Fixo mensal': o valor é R$/mês de contrato. 'Por usuário': o valor é "
         "R$ por usuário do painel por mês.")
    r += 2

    rotulo(ws, r, "Marketplace de dispositivos", 0, bold=True); r += 1
    linha_unica(ws, r, "Comissão média sobre GMV", 0.10, PCT, "com_mp", chave=True,
        obs="Brief: faixa de 5% a 15%. Adotado o ponto médio."); r += 1
    linha_unica(ws, r, "Ticket médio do dispositivo (R$)", 350.00, BRL2, "ticket_mp",
        obs="Oxímetro ~R$150, smartwatch de saúde ~R$400, balança/pressão ~R$250."); r += 1
    linha_input(ws, r, "Taxa de compra mensal (% da base ativa que compra no mês)",
                [0.000, 0.0015, 0.0025, 0.0040, 0.0050], '0.00%', "attach_mp",
        obs="Marketplace só entra no ar em 2027. Compra de device é evento raro e de alto ticket."); r += 2

    # ==================================================================
    secao(ws, r, "6. CUSTOS OPERACIONAIS (exceto pessoas — ver aba Pessoas)", 12); r += 1
    linha_unica(ws, r, "Início dos custos de nuvem", dt.datetime(2026, 5, 1), MES, "cloud_ini",
        obs="Primeiro gasto real registrado: R$ 532,29 em 04/05/2026."); r += 1
    linha_input(ws, r, "Nuvem — custo fixo mensal (R$)", [600, 2500, 8000, 18000, 30000], BRL,
        "cloud_fixo",
        fonte=("Ancorado no gasto real de R$ 532,29 em 04/05/2026 (pré-lançamento). Cresce com "
               "ambientes, observabilidade, banco gerenciado e redundância."),
        obs="Base: gasto real de R$ 532,29 em 04/05/2026."); r += 1
    linha_input(ws, r, "Nuvem — custo variável por usuário PAGANTE (R$/mês)",
                [1.50, 1.35, 1.15, 1.00, 0.90], BRL2, "cloud_pag", chave=True,
        fonte=("Usuário pagante tem histórico: série temporal de sinais vitais armazenada e consultável. "
               "Custo de banco time-series + storage + processamento. Cai com escala."),
        obs="Só o pagante gera histórico — premissa do brief."); r += 1
    linha_input(ws, r, "Nuvem — custo variável por usuário GRATUITO (R$/mês)",
                [0.02, 0.02, 0.018, 0.015, 0.015], BRL2, "cloud_free",
        fonte=("O brief assume custo ~zero para o gratuito, pois não há histórico. Não é exatamente "
               "zero: a ingestão em tempo real e as chamadas de API consomem compute mesmo sem "
               "persistência. R$ 0,02/mês (R$ 0,24/ano) é esse custo residual."),
        obs="Não é zero: ingestão em tempo real e API consomem compute mesmo sem histórico."); r += 1
    linha_input(ws, r, "Suporte — ferramentas e custo variável por pagante (R$/mês)",
                [0.35, 0.35, 0.30, 0.28, 0.25], BRL2, "sup_var",
        obs="Chatbot, help desk e base de conhecimento. As PESSOAS de suporte estão na aba Pessoas."); r += 1

    linha_unica(ws, r, "Início das despesas administrativas", dt.datetime(2026, 4, 1), MES, "adm_ini",
        obs="A partir da constituição da empresa / primeiro aporte (16/04/2026)."); r += 1
    linha_input(ws, r, "Administrativas (contabilidade, jurídico, ferramentas) — R$/mês",
                [1500, 6000, 25000, 60000, 110000], BRL, "admin",
        obs="Não inclui pessoas. Escalonado com o porte da operação."); r += 1
    linha_input(ws, r, "Marketing recorrente — mídia de conteúdo e ferramentas (R$/mês)",
                [3000, 15000, 60000, 140000, 250000], BRL, "mkt_rec",
        fonte=("Reduzido pela entrada da Shaiane, sócia com agência de marketing própria, que absorve "
               "produção criativa e gestão. Não elimina o custo: ferramentas, mídia de conteúdo, "
               "influenciadores e produção seguem sendo desembolso. Pessoas de growth ficam na aba Pessoas."),
        obs="Reduzido pela agência da Shaiane, mas não eliminado."); r += 1
    linha_unica(ws, r, "Início do marketing recorrente", dt.datetime(2026, 10, 1), MES, "mkt_ini"); r += 1
    linha_unica(ws, r, "Propaganda de lançamento — valor único (R$)", 25000, BRL, "mkt_lanc"); r += 1
    linha_unica(ws, r, "Mês da propaganda de lançamento", dt.datetime(2026, 10, 1), MES,
                "mkt_lanc_mes"); r += 2

    # ==================================================================
    secao(ws, r, "7. CAPITAL — O QUE ENTRA NA CONTA", 12); r += 1
    for t in ["Os investidores em si ficam na aba APORTES, em duas tabelas com linhas em branco: "
              "lançamentos de aporte, e investidores estratégicos que além de aportar trazem "
              "receita recorrente — o caso da Boston Scientific. Aqui ficam só os dois interruptores "
              "que decidem o que o modelo considera."]:
        c = ws.cell(r, 1, "•  " + t)
        c.font = f(9, False, CINZA)
        c.alignment = Alignment(wrap_text=True, vertical="top", indent=1)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=9)
        ws.row_dimensions[r].height = 28
        r += 1
    linha_unica(ws, r, "Considerar aportes PREVISTOS?  (1 = Sim · 0 = Não)", 1, '0',
                "inc_previsto", chave=True,
                fonte=("Aportes em negociação adiantada mas sem assinatura — hoje os R$ 120 mil de "
                       "C-ioT, Felipe Martinelli e Shaiane. Estão ligados por padrão porque a "
                       "negociação está em curso. Desligue para ver a Oppa apenas com o capital "
                       "dos fundadores."),
                obs="Hoje: R$ 120 mil de C-ioT, Martinelli e Shaiane. Desligue para ver só os fundadores."); r += 1
    linha_unica(ws, r, "Considerar acordos EM NEGOCIAÇÃO?  (1 = Sim · 0 = Não)", 0, '0',
                "inc_negoc", chave=True,
                fonte=("Vale para os lançamentos marcados como 'Em negociação' e para as parcerias "
                       "estratégicas com status 'Em negociação' — hoje, a Boston Scientific. "
                       "Desligado por padrão para que a viabilidade não dependa de receita nem de "
                       "capital não contratados. Ligue para dimensionar o upside."),
                obs="⚠ Desligado por padrão. Ligue para incluir a Boston Scientific e qualquer "
                    "outra parceria em negociação."); r += 2

    # ==================================================================
    secao(ws, r, "8. NOTAS METODOLÓGICAS", 12); r += 1
    notas = [
        "FCL (Fluxo de Caixa Livre) NÃO inclui aportes de sócios. Aporte é financiamento, não geração "
        "de caixa — incluí-lo infla VPL e TIR. A planilha original somava o aporte à receita; aqui os "
        "dois fluxos estão separados.",
        "VPL, TIR e Payback são calculados sobre o FCL. Caixa Acumulado e Runway usam FCL + aportes.",
        "Rendimento financeiro do caixa não é considerado (premissa conservadora).",
        "Investimentos saem do caixa no mês do desembolso e são amortizados linearmente em 60 meses na "
        "DRE — por isso o EBITDA da DRE difere do fluxo de caixa livre.",
        "Propaganda de lançamento é despesa de marketing, não investimento — diferente da planilha "
        "original, que a colocava em 'Projeto'.",
        "O modelo troca sozinho do Simples para o Lucro Presumido quando o RBT12 passa de R$ 4,8 milhões.",
        "Todo o custo de pessoas — PJ, CLT e pró-labore — está na aba Pessoas, linha a linha.",
    ]
    for n in notas:
        c = ws.cell(r, 1, "•  " + n)
        c.font = f(9, False, CINZA)
        c.alignment = Alignment(wrap_text=True, vertical="top", indent=1)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=9)
        ws.row_dimensions[r].height = 28
        r += 1

    return ws, P
