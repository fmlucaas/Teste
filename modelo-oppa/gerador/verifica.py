# -*- coding: utf-8 -*-
"""Confere os numeros calculados do modelo."""
import sys, openpyxl
from openpyxl.utils import get_column_letter as gcl
p = sys.argv[1]
wb = openpyxl.load_workbook(p, data_only=True)

def g(sheet, ref):
    return wb[sheet][ref].value

def fmt(v):
    if isinstance(v, float):
        return f"{v:,.2f}"
    return str(v)

print("=" * 78)
print("RESUMO — painel executivo")
ws = wb["Resumo"]
for r in range(4, 70):
    lbl = ws.cell(r, 2).value
    if lbl and not str(lbl).startswith("▸"):
        vals = [ws.cell(r, c).value for c in (3, 4, 5)]
        if any(v is not None for v in vals):
            print(f"  {str(lbl)[:52]:<52} | " + " | ".join(f"{fmt(v):>16}" if v is not None else " " * 16 for v in vals))

print()
print("=" * 78)
print("VERIFICAÇÃO DE RECONCILIAÇÃO (Análise Fluxo x Motor) — deve ser 0")
af = wb["Análise Fluxo"]
# localizar linha de check
chk = None
for r in range(1, 120):
    if af.cell(r, 1).value and "Verificação" in str(af.cell(r, 1).value):
        chk = r
mx = 0
if chk:
    for i in range(2, 62):
        v = af.cell(chk, i).value
        if isinstance(v, (int, float)):
            mx = max(mx, abs(v))
    print(f"  linha {chk}: desvio máximo mensal = {mx}")
else:
    print("  linha de verificação não encontrada")

print()
print("=" * 78)
print("MOTOR — cenário PROVÁVEL (bloco linha 53), marcos anuais")
cen = wb["Cenários"]
S = 53
labels = {4: "Base de usuários", 6: "Pagantes", 16: "Receita bruta/mês",
          18: "Alíquota efetiva", 28: "EBITDA/mês", 30: "FCL/mês", 37: "Caixa acumulado"}
marcos = {"dez/26": "M", "dez/27": "Y", "dez/28": "AK", "dez/29": "AW", "dez/30": "BI"}
print(f"  {'':<22}" + "".join(f"{k:>16}" for k in marcos))
for off, lbl in labels.items():
    linha = f"  {lbl:<22}"
    for k, col in marcos.items():
        v = cen[f"{col}{S+off}"].value
        linha += f"{fmt(v):>16}" if v is not None else " " * 16
    print(linha)

print()
print("=" * 78)
print("DRE ANUAL (cenário ativo)")
d = wb["DRE Projetada"]
hdr = None
for r in range(1, 200):
    if d.cell(r, 1).value == "Demonstração do Resultado":
        hdr = r
if hdr:
    print(f"  {'':<40}" + "".join(f"{d.cell(hdr,c).value!s:>15}" for c in range(2, 8)))
    for r in range(hdr + 1, hdr + 19):
        lbl = d.cell(r, 1).value
        if lbl:
            print(f"  {str(lbl)[:40]:<40}" + "".join(f"{fmt(d.cell(r,c).value):>15}"
                  if d.cell(r, c).value is not None else " " * 15 for c in range(2, 8)))

print()
print("=" * 78)
print("APORTES / INVESTIMENTOS")
a = wb["Aportes"]
for r in range(24, 40):
    if a.cell(r, 2).value:
        print(f"  {str(a.cell(r,2).value)[:52]:<52} {fmt(a.cell(r,3).value)}")
inv = wb["Investimentos"]
for r in range(1, 30):
    if inv.cell(r, 3).value == "TOTAL":
        print(f"  {'CAPEX total':<52} {fmt(inv.cell(r,4).value)}")

print()
print("=" * 78)
print("TRIBUTOS — comparativo Anexo III x Anexo V")
t = wb["Tributos"]
for r in range(1, 80):
    if t.cell(r, 1).value and isinstance(t.cell(r, 1).value, (int, float)) and t.cell(r, 5).value is not None and r > 40:
        print(f"  RBT12 {fmt(t.cell(r,1).value):>14} | AIII {fmt(t.cell(r,2).value):>8} | "
              f"AV {fmt(t.cell(r,3).value):>8} | custo extra/ano {fmt(t.cell(r,5).value):>14}")

print()
print("=" * 78)
print("MÉTRICAS DE UNIDADE (Usuários) — dez/26, dez/28, dez/30")
u = wb["Usuários"]
for r in range(1, 60):
    lbl = u.cell(r, 1).value
    if lbl and any(k in str(lbl) for k in ("LTV", "CAC", "ARPPU", "Runway", "Margem de contrib", "Payback do CAC")):
        print(f"  {str(lbl)[:46]:<46} " + " | ".join(
            f"{fmt(u[f'{c}{r}'].value):>12}" if u[f'{c}{r}'].value is not None else " " * 12
            for c in ("M", "AK", "BI")))
