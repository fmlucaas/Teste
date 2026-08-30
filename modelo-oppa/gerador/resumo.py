# -*- coding: utf-8 -*-
"""Aba Resumo: painel executivo com os tres cenarios lado a lado."""
from openpyxl.styles import Alignment
from common import *

LANO = 6   # linha do ano na aba Cenarios
COLS = ["C", "D", "E"]
NOMES = ["CONSERVADOR", "PROVÁVEL", "AGRESSIVO"]
CORES = ["C00000", "1F3864", "1F7040"]


def build(wb, P, A, I, BLOCOS, FL, U):
    ws = wb.create_sheet("Resumo")
    titulo(ws, "OPPA — VIABILIDADE FINANCEIRA · PAINEL EXECUTIVO",
           "Projeção de 60 meses (jan/2026 – dez/2030). Os três cenários são calculados "
           "simultaneamente. Todos os números vêm de fórmulas — nada está digitado à mão.",
           largura=8)
    for col, w in zip("ABCDEFGH", (4, 52, 21, 21, 21, 4, 60, 4)):
        ws.column_dimensions[col].width = w

    def cen(off, S_index, agg="last", ano=None):
        S = BLOCOS[S_index]
        if agg == "last":
            return f"='Cenários'!${L_FIM}${S+off}"
        if agg == "sum":
            return f"=SUM('Cenários'!$B${S+off}:${L_FIM}${S+off})"
        if agg == "ano":
            return (f"=SUMIFS('Cenários'!$B${S+off}:${L_FIM}${S+off},"
                    f"'Cenários'!$B${LANO}:${L_FIM}${LANO},{ano})")
        raise ValueError(agg)

    r = 4
    # ---------------- cabecalho de cenarios ------------------------------
    ws.cell(r, 2, "INDICADOR").font = f(10, True, BRANCO)
    ws.cell(r, 2).fill = FILL_HEADER
    ws.cell(r, 2).alignment = Alignment(indent=1, vertical="center")
    for k, col in enumerate(COLS):
        c = ws[f"{col}{r}"]
        c.value = NOMES[k]
        c.font = f(10, True, BRANCO)
        c.fill = PatternFill("solid", fgColor=CORES[k])
        c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[r].height = 22
    hdr = r
    r += 1
    ws.cell(r, 2, "Cenário em uso nas abas de detalhe").font = f(9, True, VERDE_LINK)
    ws.cell(r, 2).alignment = Alignment(indent=1)
    for k, col in enumerate(COLS):
        c = ws[f"{col}{r}"]
        c.value = f'=IF(Premissas!$B${P["cenario"]}={k+1},"◀ ATIVO","")'
        c.font = f(9, True, VERDE_LINK)
        c.alignment = Alignment(horizontal="center")
    r += 2

    R = {}

    def bloco(nome):
        nonlocal r
        secao(ws, r, nome, 5)
        r += 1

    def linha(lbl, off, agg, fmt, key=None, ano=None, bold=False, obs=None, custom=None):
        nonlocal r
        c = ws.cell(r, 2, lbl)
        c.font = f(10, bold, PRETO)
        c.alignment = Alignment(indent=1, vertical="center")
        for k, col in enumerate(COLS):
            cc = ws[f"{col}{r}"]
            cc.value = custom(k, col) if custom else cen(off, k, agg, ano)
            cc.number_format = fmt
            cc.font = f(10, bold, PRETO)
            cc.border = BORDA_FINA
            cc.alignment = Alignment(horizontal="center")
            if bold:
                cc.fill = FILL_TOTAL
        if obs:
            nota(ws, f"G{r}", obs)
        if key:
            R[key] = r
        r += 1
        return r - 1

    # ---------------- TRAÇÃO ---------------------------------------------
    bloco("TRAÇÃO — USUÁRIOS E RECEITA")
    linha("Usuários na base em dez/2026", 4, "custom", INT, "u26",
          custom=lambda k, col: f"='Cenários'!$M${BLOCOS[k]+4}",
          obs="Meta do brief: 50 mil usuários ao final de 2026.")
    linha("Usuários na base em dez/2030", 4, "last", INT, "u30")
    linha("Assinantes pagantes em dez/2026", 6, "custom", INT, "p26",
          custom=lambda k, col: f"='Cenários'!$M${BLOCOS[k]+6}",
          obs="Meta do brief: 5% da base pagante.")
    linha("Assinantes pagantes em dez/2030", 6, "last", INT, "p30")
    linha("Receita bruta de 2026", 16, "ano", BRL, "rb26", ano=2026)
    linha("Receita bruta de 2030", 16, "ano", BRL, "rb30", ano=2030)
    linha("Receita bruta acumulada em 5 anos", 16, "sum", BRL, "rb_tot", bold=True)
    r += 1

    # ---------------- RESULTADO -------------------------------------------
    bloco("RESULTADO OPERACIONAL")
    linha("EBITDA de 2026", 28, "ano", BRL, "eb26", ano=2026)
    linha("EBITDA de 2030", 28, "ano", BRL, "eb30", ano=2030)
    linha("EBITDA acumulado em 5 anos", 28, "sum", BRL, "eb_tot", bold=True)
    linha("Margem EBITDA em 2030 (% da receita bruta)", 0, "custom", PCT, "mg30",
          custom=lambda k, col: f'=IFERROR({col}{R["eb30"]}/{col}{R["rb30"]},0)')
    linha("Break-even do EBITDA", 0, "custom", TXT, "be",
          custom=lambda k, col: (
              f'=IF(SUM(\'Cenários\'!$B${BLOCOS[k]+34}:${L_FIM}${BLOCOS[k]+34})=0,"não atinge",'
              f'CHOOSE(MOD(60-SUM(\'Cenários\'!$B${BLOCOS[k]+34}:${L_FIM}${BLOCOS[k]+34}),12)+1,'
              f'"jan","fev","mar","abr","mai","jun","jul","ago","set","out","nov","dez")&"/"&'
              f'(2026+INT((60-SUM(\'Cenários\'!$B${BLOCOS[k]+34}:${L_FIM}${BLOCOS[k]+34}))/12)))'),
          obs="Primeiro mês com EBITDA positivo — quando a operação passa a se pagar.")
    r += 1

    # ---------------- CAPITAL ---------------------------------------------
    bloco("CAPITAL E CAIXA")
    linha("Investimento total (CAPEX)", 29, "sum", BRL, "capex",
          custom=lambda k, col: f"=-SUM('Cenários'!$B${BLOCOS[k]+29}:${L_FIM}${BLOCOS[k]+29})")
    linha("Capital requerido (pico de caixa negativo do FCL)", 33, "custom", BRL, "capreq",
          custom=lambda k, col: f"=-MIN('Cenários'!$B${BLOCOS[k]+33}:${L_FIM}${BLOCOS[k]+33})",
          bold=True,
          obs="Quanto a operação precisa de capital antes de se sustentar sozinha.")
    linha("Capital total considerado", 0, "custom", BRL, "cap_disp",
          custom=lambda k, col: f"=SUM(Aportes!$B${A['total']}:${L_FIM}${A['total']})")
    linha("Caixa mínimo ao longo do período", 37, "custom", BRL, "caixa_min",
          custom=lambda k, col: f"=MIN('Cenários'!$B${BLOCOS[k]+37}:${L_FIM}${BLOCOS[k]+37})",
          bold=True,
          obs="Se ficar negativo, o capital planejado não é suficiente — falta captar a diferença.")
    linha("Mês do caixa mínimo", 0, "custom", TXT, "caixa_min_mes",
          custom=lambda k, col: (
              f"=CHOOSE(MOD(MATCH(MIN('Cenários'!$B${BLOCOS[k]+37}:${L_FIM}${BLOCOS[k]+37}),"
              f"'Cenários'!$B${BLOCOS[k]+37}:${L_FIM}${BLOCOS[k]+37},0)-1,12)+1,"
              f'"jan","fev","mar","abr","mai","jun","jul","ago","set","out","nov","dez")&"/"&'
              f"(2026+INT((MATCH(MIN('Cenários'!$B${BLOCOS[k]+37}:${L_FIM}${BLOCOS[k]+37}),"
              f"'Cenários'!$B${BLOCOS[k]+37}:${L_FIM}${BLOCOS[k]+37},0)-1)/12))"),
          obs="Quando o caixa chega no fundo do poço — é o mês para o qual o aporte precisa chegar antes.")
    linha("Caixa em dez/2030", 37, "last", BRL, "caixa_fim")
    r += 1

    # ---------------- VIABILIDADE -----------------------------------------
    bloco("INDICADORES DE VIABILIDADE (sobre o Fluxo de Caixa Livre)")
    linha("VPL — Valor Presente Líquido", 30, "custom", BRL, "vpl",
          custom=lambda k, col: (f"=NPV(Premissas!$B${P['tma_m']},"
                                 f"'Cenários'!$B${BLOCOS[k]+30}:${L_FIM}${BLOCOS[k]+30})"),
          bold=True, obs="Descontado à TMA definida em Premissas. Positivo = o projeto cria valor.")
    linha("TIR mensal", 30, "custom", PCT2, "tir_m",
          custom=lambda k, col: (f"=IFERROR(IRR('Cenários'!$B${BLOCOS[k]+30}:"
                                 f"${L_FIM}${BLOCOS[k]+30}),\"n/d\")"))
    linha("TIR anual", 0, "custom", PCT, "tir_a",
          custom=lambda k, col: f'=IFERROR((1+{col}{R["tir_m"]})^12-1,"n/d")', bold=True,
          obs="Compare com a TMA. Atenção: quando o investimento inicial é pequeno diante do caixa gerado, a TIR fica muito alta e perde poder de comparação — olhe o VPL junto.")
    linha("Payback simples (meses)", 33, "custom", INT, "pb",
          custom=lambda k, col: (
              f'=IF(\'Cenários\'!${L_FIM}${BLOCOS[k]+33}<0,"> 60",'
              f'COUNTIF(\'Cenários\'!$B${BLOCOS[k]+33}:${L_FIM}${BLOCOS[k]+33},"<0")+1)'))
    linha("Payback descontado (meses)", 32, "custom", INT, "pbd",
          custom=lambda k, col: (
              f'=IF(\'Cenários\'!${L_FIM}${BLOCOS[k]+32}<0,"> 60",'
              f'COUNTIF(\'Cenários\'!$B${BLOCOS[k]+32}:${L_FIM}${BLOCOS[k]+32},"<0")+1)'),
          obs="Critério de decisão: payback desejado definido em Premissas.")
    linha("FCL acumulado em 5 anos", 30, "sum", BRL, "fcl_tot")
    linha("ROI sobre o capital requerido", 0, "custom", PCT, "roi",
          custom=lambda k, col: f'=IFERROR({col}{R["fcl_tot"]}/{col}{R["capreq"]},0)')
    r += 1

    # ---------------- CHECKLIST -------------------------------------------
    bloco("TESTE DE DECISÃO — GO / NO-GO")
    testes = [
        ("O caixa nunca fica negativo?",
         lambda k, col: f'=IF({col}{R["caixa_min"]}>=0,"SIM","NÃO")'),
        ("O VPL é positivo?",
         lambda k, col: f'=IF({col}{R["vpl"]}>0,"SIM","NÃO")'),
        ("A TIR anual supera a TMA?",
         lambda k, col: f'=IF(N({col}{R["tir_a"]})>Premissas!$B${P["tma"]},"SIM","NÃO")'),
        ("O payback descontado cabe no prazo desejado?",
         lambda k, col: (f'=IF(AND(ISNUMBER({col}{R["pbd"]}),'
                         f'N({col}{R["pbd"]})<=Premissas!$B${P["payback_alvo"]}),"SIM","NÃO")')),
    ]
    t_ini = r
    for lbl, fn in testes:
        c = ws.cell(r, 2, lbl); c.font = f(10); c.alignment = Alignment(indent=1)
        for k, col in enumerate(COLS):
            cc = ws[f"{col}{r}"]
            cc.value = fn(k, col)
            cc.font = f(10, True); cc.border = BORDA_FINA
            cc.alignment = Alignment(horizontal="center")
        r += 1
    t_fim = r - 1
    c = ws.cell(r, 2, "DIAGNÓSTICO"); c.font = f(12, True, BRANCO); c.fill = FILL_TITULO
    c.alignment = Alignment(indent=1, vertical="center")
    for k, col in enumerate(COLS):
        cc = ws[f"{col}{r}"]
        cc.value = (f'=IF(COUNTIF({col}{t_ini}:{col}{t_fim},"SIM")=4,"GO — VIÁVEL",'
                    f'IF(COUNTIF({col}{t_ini}:{col}{t_fim},"SIM")>=2,"GO COM RESSALVAS",'
                    f'"NO GO — INVIÁVEL"))')
        cc.font = f(11, True, BRANCO)
        cc.fill = PatternFill("solid", fgColor=CORES[k])
        cc.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[r].height = 24
    R["diag"] = r
    nota(ws, f"G{r}", "Critério herdado da planilha de viabilidade da EMS (Go / No-Go).")
    r += 2

    # ---------------- MÉTRICAS DE UNIDADE ---------------------------------
    bloco("MÉTRICAS DE UNIDADE — CENÁRIO ATIVO (dez/2030)")
    for lbl, src, fmt, obs in [
        ("LTV — valor do assinante", U["ltv"], BRL2, ""),
        ("CAC por assinante pagante", U["cac_pag"], BRL2, ""),
        ("LTV ÷ CAC", U["ltv_cac"], MULT, "Referência: ≥ 3,0x é saudável. Abaixo de 1,0x destrói valor."),
        ("Payback do CAC (meses)", U["pb_cac"], '0.0', "Ideal abaixo de 12 meses."),
        ("ARPPU — receita por assinante", U["arppu"], BRL2, ""),
        ("Runway em dez/2026 (meses de caixa)", U["runway"], '0', ""),
        ("Pessoas na operação em dez/2030", U["hc"], INT,
         "Quadro de pessoal projetado. Edite linha a linha na aba Pessoas."),
    ]:
        c = ws.cell(r, 2, lbl); c.font = f(10); c.alignment = Alignment(indent=1)
        cc = ws["C" + str(r)]
        cc.value = (f"=Usuários!$M${src}" if "dez/2026" in lbl else f"=Usuários!${L_FIM}${src}")
        cc.number_format = fmt; cc.font = f(10, True)
        cc.fill = FILL_TOTAL
        cc.alignment = Alignment(horizontal="center")
        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=5)
        for cx in range(3, 6):
            ws.cell(r, cx).border = BORDA_FINA
        if obs:
            nota(ws, f"G{r}", obs)
        r += 1
    r += 1

    # ---------------- ALERTAS ---------------------------------------------
    secao(ws, r, "PONTOS DE ATENÇÃO IDENTIFICADOS NA MODELAGEM", 8); r += 1
    alertas = [
        ("Fator R: o modelo está no Anexo V, e é isso mesmo",
         "A aba Pessoas calcula o Fator R do quadro projetado: fica entre 1% e 6%, muito abaixo dos 28% "
         "exigidos para o Anexo III. Como pagamento a PJ não conta como folha, o Anexo III simplesmente "
         "não está disponível com a estrutura atual — por isso o regime padrão do modelo é o Anexo V. "
         "Migrar parte do time para CLT ou elevar o pró-labore pode reverter isso e economizar até "
         "R$ 135 mil por ano; compare na aba Tributos antes de decidir."),
        ("Dependência dos aportes não assinados",
         "Dos R$ 200 mil do caso-base, R$ 120 mil vêm de C-ioT, Felipe Martinelli e Shaiane — ainda não "
         "formalizados. Sem eles, o capital cai para R$ 80 mil e o lançamento com aquisição paga não se "
         "sustenta. Este é o maior risco de curto prazo do plano."),
        ("Comissão das lojas de aplicativos",
         "Apple e Google retêm 15% da assinatura (Google desde o primeiro dia; Apple pelo Small Business "
         "Program, válido enquanto a receita anual for inferior a US$ 1 milhão). Acima disso a Apple "
         "volta a 30%. É a maior dedução isolada da receita e não estava no brief."),
        ("Conversão de 5% é premissa de quartil superior",
         "Benchmarks de freemium apontam mediana de 2,2% a 2,6% e faixa típica de 2% a 5%. Os 5% do brief "
         "são defensáveis para saúde com feature paga clara (histórico de dados), mas são a premissa mais "
         "sensível do modelo — teste o cenário Conservador (3%)."),
        ("Teto do Simples Nacional",
         "O modelo troca automaticamente para o Lucro Presumido quando o RBT12 ultrapassa R$ 4,8 milhões. "
         "No cenário Provável isso acontece dentro do horizonte e eleva a carga tributária — planeje a "
         "transição societária e contábil com antecedência."),
        ("Só a equipe do MVP está de fato contratada",
         "Hoje não há time. O único custo de pessoal comprometido são os R$ 21 mil da equipe PJ de "
         "set–out/2026 para entregar o app. Todo o quadro projetado a partir de 2027 é plano, não "
         "compromisso: está na aba Pessoas, linha a linha, e pode ser desligado inteiro pelo controle "
         "de Premissas. É a premissa de custo que mais desloca o resultado."),
        ("A conversão de 2026 depende da rampa, não do patamar",
         "O patamar de 5% do brief é o regime de médio prazo. Com a rampa de maturação de 6 meses, "
         "outubro/2026 converte a 0,83%, novembro a 1,67% e dezembro a 2,50% — 2026 fecha bem abaixo "
         "dos 5%. Ajuste a rampa em Premissas conforme a discussão com os sócios."),
        ("O cenário Conservador mantém a estrutura de custos do Provável",
         "É proposital: mede o risco de dimensionar equipe, marketing e infraestrutura para uma tração "
         "que não se confirma. Na prática a gestão cortaria custos — ajuste as linhas de equipe, "
         "marketing recorrente e administrativas em Premissas para ver o Conservador com a operação "
         "redimensionada."),
        ("Cruzamento dos anexos no teto do Simples",
         "Na última faixa (RBT12 entre R$ 3,6 mi e R$ 4,8 mi) o Anexo V passa a ser ligeiramente mais "
         "barato que o Anexo III, porque a alíquota nominal do Anexo III sobe para 33%. Não é erro da "
         "tabela — é o desenho da LC 123/2006. Até essa faixa, o Anexo III é sempre melhor."),
        ("Boston Scientific fora do caso-base",
         "O acordo não está confirmado e por isso está desligado: o interruptor 'Considerar acordos em "
         "negociação' está em 0, na seção 7 de Premissas. Ligue para dimensionar o upside — e "
         "considere que a exclusividade B2B hospitalar limita outras receitas do mesmo canal. Para "
         "incluir outro investidor estratégico, some uma linha na tabela 2 da aba Aportes."),
    ]
    for tit, txt in alertas:
        c = ws.cell(r, 2, "▸  " + tit)
        c.font = f(10, True, "C00000")
        c.alignment = Alignment(indent=1)
        r += 1
        c = ws.cell(r, 2, txt)
        c.font = f(9, False, CINZA)
        c.alignment = Alignment(wrap_text=True, vertical="top", indent=2)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=7)
        ws.row_dimensions[r].height = 42
        r += 1
    ws.sheet_view.showGridLines = False
    return ws, R
