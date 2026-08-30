# -*- coding: utf-8 -*-
import sys, openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as gcl
from openpyxl.comments import Comment
from openpyxl.worksheet.table import Table, TableStyleInfo
from ordena import calcula
from dados import CRITERIOS, RISCOS, LACUNAS, OBS, PESOS

OUT = sys.argv[1]
FONT = "Arial"
PRETO, BRANCO, CINZA, VERM, VERDE = "000000", "FFFFFF", "595959", "C00000", "1F7040"
F_TIT = PatternFill("solid", fgColor="1F3864")
F_SEC = PatternFill("solid", fgColor="2E5C8A")
F_HDR = PatternFill("solid", fgColor="404040")
F_SUB = PatternFill("solid", fgColor="D9E2F3")
F_TOT = PatternFill("solid", fgColor="BDD7EE")
F_ALE = PatternFill("solid", fgColor="FCE4D6")
F_OK  = PatternFill("solid", fgColor="E2EFDA")
THIN = Side(style="thin", color="BFBFBF")
BORDA = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def altura(textos_larguras, minimo=26):
    """Altura da linha em funcao do texto mais alto entre (texto, largura)."""
    linhas = 1
    for txt, larg in textos_larguras:
        if txt:
            linhas = max(linhas, -(-len(str(txt)) // max(10, int(larg * 0.95))))
    return max(minimo, linhas * 13 + 10)


ONDAS = [
    (1, 10, "Onda 1 — Fundação e confiança", "1F3864",
     "Nada aqui aparece numa demonstração, e sem isso nada mais existe. Conta, perfil, aceite "
     "de termos, navegação, quem pode o quê, e as obrigações de LGPD e segurança que precisam "
     "estar de pé antes de o primeiro dado de saúde entrar no sistema. O SOS entra já nesta onda "
     "porque é barato e é a funcionalidade que mais gera confiança no primeiro contato."),
    (11, 19, "Onda 2 — O produto começa a cuidar", "1F7040",
     "É aqui que o app passa a fazer o que promete: recebe os dados do relógio, mostra na tela "
     "inicial e avisa quando algo dá errado. A ingestão vem antes de tudo porque é o gargalo de "
     "onze itens. iOS primeiro, Android na sequência, reaproveitando o mesmo contrato de dados."),
    (20, 25, "Onda 3 — Cobrar e alertar melhor", "BF8F00",
     "Com o produto funcionando, define-se o que é gratuito e o que é pago, e entram os alertas "
     "que sustentam a assinatura: remédio esquecido e cerco virtual. Antes desta onda não havia "
     "o que restringir; depois dela, há motivo para pagar."),
    (26, 35, "Onda 4 — Médico e complementos", "2E5C8A",
     "O médico entra como ator, o histórico clínico ganha corpo e os itens de rotina e "
     "engajamento completam a experiência. Nada aqui é urgente, tudo aqui é o que diferencia "
     "o produto de um rastreador genérico."),
    (36, 40, "Onda 5 — Expansão", "7030A0",
     "Plano Família, casa conectada e versão web do médico. São apostas de maior custo e menor "
     "certeza. A casa conectada depende de aprovação do Google e merece uma prova de conceito "
     "antes de virar compromisso."),
]


def f(sz=10, b=False, c=PRETO, i=False):
    return Font(name=FONT, size=sz, bold=b, color=c, italic=i)


def titulo(ws, txt, sub, larg):
    ws["A1"] = txt
    ws["A1"].font = f(14, True, BRANCO)
    ws["A1"].alignment = Alignment(vertical="center", indent=1)
    for c in range(1, larg + 1):
        ws.cell(1, c).fill = F_TIT
    ws.row_dimensions[1].height = 26
    ws["A2"] = sub
    ws["A2"].font = f(9, False, CINZA, True)
    ws["A2"].alignment = Alignment(vertical="center", indent=1, wrap_text=False)
    ws.row_dimensions[2].height = 15


def cab(ws, linha, cols, larguras):
    for j, h in enumerate(cols):
        c = ws.cell(linha, 1 + j, h)
        c.font = f(9, True, BRANCO); c.fill = F_HDR
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDA
    for j, w in enumerate(larguras):
        ws.column_dimensions[gcl(1 + j)].width = w
    ws.row_dimensions[linha].height = 34


def onda_de(rank):
    for ini, fim, nome, cor, _ in ONDAS:
        if ini <= rank <= fim:
            return nome.split(" — ")[0], cor
    return "", "000000"


def build():
    itens, ordem, ordem_orig = calcula()
    wb = openpyxl.Workbook(); wb.remove(wb.active)

    # ============================================================ 1. PRIORIZAÇÃO
    ws = wb.create_sheet("Priorização")
    titulo(ws, "BACKLOG OPPA — PRIORIZAÇÃO DOS 40 PBIs",
           "Colunas A a E são as da planilha original, com a coluna E preenchida. Da F em diante "
           "está o porquê de cada posição. Aba “Como priorizei” explica o método; aba “Ordem de "
           "execução” mostra a mesma lista já ordenada.", 17)
    colunas = ["ID PBI", "Épico / Módulo", "Funcionalidade (PBI)",
               "O que é e Para que serve? (Linguagem de Negócio)",
               "Sua Ordem de Prioridade (1 a 40)", "Onda", "Custo de atraso",
               "Esforço relativo", "Índice de prioridade", "Valor para\no usuário",
               "Obrigação\nlegal", "Habilita\nmonetização", "Destrava\noutros",
               "Depende de", "Destrava", "Risco a gerenciar", "Por que nesta posição"]
    larg = [10, 21, 33, 62, 11, 10, 11, 11, 12, 9, 9, 10, 9, 20, 20, 46, 78]
    cab(ws, 4, colunas, larg)
    ws.freeze_panes = "F5"

    r = 5
    for pid in ordem_orig:
        i = itens[pid]
        onda, cor = onda_de(i["rank"])
        vals = [pid, i["epico"], i["func"], i["desc"], i["rank"], onda,
                i["cod"], i["ESF"], round(i["wsjf_cadeia"], 1),
                i["VU"], i["OBR"], i["MON"], i["DEP"],
                ", ".join(i["pre"]) or "—",
                ", ".join(i["bloqueia"]) or "—",
                RISCOS.get(pid, ""), i["just"]]
        for j, v in enumerate(vals):
            c = ws.cell(r, 1 + j, v)
            c.border = BORDA
            c.font = f(9)
            c.alignment = Alignment(vertical="top",
                                    wrap_text=j in (1, 2, 3, 13, 14, 15, 16),
                                    horizontal="center" if j in (0, 4, 5, 6, 7, 8, 9, 10, 11, 12) else "left",
                                    indent=0 if j in (0, 4, 5, 6, 7, 8, 9, 10, 11, 12) else 1)
        ws.cell(r, 5).font = f(11, True, BRANCO)
        ws.cell(r, 5).fill = PatternFill("solid", fgColor=cor)
        ws.cell(r, 6).font = f(9, True, cor)
        ws.cell(r, 7).font = f(9, True)
        ws.cell(r, 9).font = f(9, True)
        if pid in RISCOS:
            ws.cell(r, 16).fill = F_ALE
            ws.cell(r, 16).font = f(9, False, VERM)
        ws.cell(r, 17).font = f(9, False, CINZA)
        ws.row_dimensions[r].height = altura(
            [(i["desc"], larg[3]), (RISCOS.get(pid, ""), larg[15]), (i["just"], larg[16]),
             (", ".join(i["bloqueia"]), larg[14])], 46)
        r += 1
    ws.auto_filter.ref = f"A4:Q{r-1}"

    # ============================================================ 2. ORDEM
    ws2 = wb.create_sheet("Ordem de execução")
    titulo(ws2, "ORDEM DE EXECUÇÃO — DO PRIMEIRO AO QUADRAGÉSIMO",
           "A mesma lista, ordenada e agrupada em ondas. Cada onda é um bloco que faz sentido "
           "entregar junto; a divisão é de sequência, não de prazo.", 8)
    cab(ws2, 4, ["#", "ID PBI", "Épico / Módulo", "Funcionalidade",
                 "Custo de atraso", "Esforço", "Índice", "Depende de"],
        [6, 10, 22, 40, 12, 9, 9, 24])
    ws2.freeze_panes = "A5"
    r = 5
    for ini, fim, nome, cor, obj in ONDAS:
        c = ws2.cell(r, 1, nome)
        c.font = f(11, True, BRANCO)
        for cc in range(1, 9):
            ws2.cell(r, cc).fill = PatternFill("solid", fgColor=cor)
        ws2.row_dimensions[r].height = 20
        r += 1
        c = ws2.cell(r, 1, obj)
        c.font = f(9, False, CINZA, True)
        c.alignment = Alignment(wrap_text=True, vertical="top", indent=1)
        ws2.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8)
        ws2.row_dimensions[r].height = 46
        r += 1
        for pid in ordem[ini-1:fim]:
            i = itens[pid]
            for j, v in enumerate([i["rank"], pid, i["epico"], i["func"], i["cod"],
                                   i["ESF"], round(i["wsjf_cadeia"], 1),
                                   ", ".join(i["pre"]) or "—"]):
                c = ws2.cell(r, 1 + j, v)
                c.border = BORDA; c.font = f(9)
                c.alignment = Alignment(vertical="center", wrap_text=j in (2, 3, 7),
                                        horizontal="center" if j in (0, 4, 5, 6) else "left",
                                        indent=0 if j in (0, 4, 5, 6) else 1)
            ws2.cell(r, 1).font = f(10, True, cor)
            if pid in RISCOS:
                ws2.cell(r, 2).fill = F_ALE
                ws2.cell(r, 2).comment = Comment(RISCOS[pid], "Análise", height=130, width=360)
            ws2.row_dimensions[r].height = 26
            r += 1
        r += 1
    return wb, itens, ordem





# ================================================================== 3. MÉTODO
def aba_metodo(wb, itens, ordem):
    ws = wb.create_sheet("Como priorizei")
    ws.sheet_view.showGridLines = False
    titulo(ws, "COMO PRIORIZEI", "O método, os critérios e o que fazer se você discordar.", 6)
    for col, w in zip("ABCDEF", (3, 34, 96, 12, 12, 3)):
        ws.column_dimensions[col].width = w
    r = 4

    def sec(t):
        nonlocal r
        r += 1
        c = ws.cell(r, 2, t); c.font = f(12, True, BRANCO)
        c.alignment = Alignment(indent=1, vertical="center")
        for cc in (2, 3, 4, 5):
            ws.cell(r, cc).fill = F_SEC
        ws.row_dimensions[r].height = 22; r += 1

    def par(txt, cor=CINZA, b=False, alt=None):
        nonlocal r
        c = ws.cell(r, 2, txt); c.font = f(10, b, cor)
        c.alignment = Alignment(wrap_text=True, vertical="top", indent=1)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)
        ws.row_dimensions[r] .height = alt or max(16, (len(txt)//118 + 1) * 14 + 4)
        r += 1

    def item(t, d, cor=PRETO):
        nonlocal r
        c = ws.cell(r, 2, t); c.font = f(10, True, cor)
        c.alignment = Alignment(indent=1, vertical="top", wrap_text=True)
        c2 = ws.cell(r, 3, d); c2.font = f(10, False, "262626")
        c2.alignment = Alignment(vertical="top", wrap_text=True)
        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=5)
        ws.row_dimensions[r].height = max(17, (len(d)//105 + 1) * 14 + 6)
        ws.cell(r, 2).border = Border(top=Side(style="thin", color="D9D9D9"))
        ws.cell(r, 3).border = Border(top=Side(style="thin", color="D9D9D9"))
        r += 1

    sec("A IDEIA EM UMA FRASE")
    par("Primeiro o que entrega mais valor por unidade de esforço; depois, corrigindo essa ordem "
        "sempre que um item só puder ser feito depois de outro. Nenhuma posição foi escolhida a "
        "dedo: todas saem da mesma conta, e a conta está aberta na aba Priorização para você "
        "conferir ou discordar item a item.", PRETO, alt=44)

    sec("OS QUATRO CRITÉRIOS")
    par("Cada PBI recebe nota de 1 a 5 em quatro critérios. Os pesos estão na última coluna.", PRETO)
    for sigla, nome, peso, desc in CRITERIOS:
        item(f"{nome}  (peso {str(peso).replace('.',',')})", desc)

    sec("A CONTA")
    item("Custo de atraso", "É o que se perde por adiar o item: "
         "3 × valor  +  2,5 × obrigação  +  2 × monetização  +  2 × destrava outros. "
         "Quanto maior, mais caro deixar para depois.")
    item("Índice de prioridade", "Custo de atraso dividido pelo esforço relativo. É o WSJF, "
         "usado em Scaled Agile: entrega mais cedo o que dá mais retorno por unidade de trabalho. "
         "Dois itens de mesmo valor, o menor sai primeiro — porque libera o time antes.")
    item("Correção de cadeia", "Um índice calculado só com o próprio item é míope: um habilitador "
         "que pontua baixo sozinho, mas destrava um bloco inteiro de alto valor, ficaria "
         "injustamente para trás. O índice aqui considera a melhor sequência alcançável a partir "
         "de cada item — foi o que trouxe a Conexão Rápida dos Relógios (PBI-26) da 34ª para a "
         "11ª posição: sozinha ela vale pouco, mas onze itens esperam por ela.")
    item("Trava de dependência", "Nenhum item aparece antes de algo de que ele dependa, por mais "
         "alto que seja seu índice. As colunas “Depende de” e “Destrava” mostram cada uma dessas "
         "amarras.")

    sec("ESFORÇO NÃO É CUSTO")
    par("A coluna de esforço é tamanho relativo, de 1 (muito pequeno) a 5 (muito grande), e serve "
        "apenas para comparar os itens entre si dentro desta lista. Não há nenhuma conversão para "
        "reais, horas ou prazo aqui — conforme combinado, orçamento é assunto de outra conversa.",
        PRETO, alt=44)

    sec("O QUE ESTA ORDEM DELIBERADAMENTE NÃO FAZ")
    item("Não é ordem de valor puro", "Um item grande e valioso pode aparecer depois de três itens "
         "pequenos e medianos, porque os três juntos saem antes dele. Se a decisão for maximizar "
         "valor entregue cedo em vez de vazão, ordene pela coluna Custo de atraso — é uma "
         "leitura legítima e diferente da mesma tabela.")
    item("Não é cronograma", "São posições relativas e ondas de sequência, sem datas. Duas pessoas "
         "trabalhando em paralelo podem tocar itens de posições distantes ao mesmo tempo, desde "
         "que as dependências sejam respeitadas.")
    item("Não substitui conversa", "As notas são julgamento explícito, não verdade. O ganho está "
         "em estarem escritas: dá para discordar de uma nota específica em vez de discutir a "
         "lista inteira. Mude a nota, refaça a conta, e a ordem se ajusta sozinha.")

    sec("SE VOCÊ DISCORDAR DE ALGUMA POSIÇÃO")
    par("Vá até a linha do item na aba Priorização e olhe as quatro notas e o esforço. Em geral a "
        "divergência está em uma nota só. Os casos que mais costumam gerar debate nesta lista "
        "estão marcados na coluna “Risco a gerenciar”.", PRETO)

    sec("MAIOR CUSTO DE ATRASO — O QUE MAIS DÓI ADIAR")
    par("Estes são os itens mais caros de deixar para depois, independentemente do tamanho. "
        "Se o objetivo for impacto máximo cedo, é por aqui que se começa a discussão.", PRETO)
    top = sorted(itens.values(), key=lambda i: -i["cod"])[:8]
    ws.cell(r, 2, "PBI").font = f(9, True, BRANCO)
    ws.cell(r, 3, "Funcionalidade").font = f(9, True, BRANCO)
    ws.cell(r, 4, "Custo de atraso").font = f(9, True, BRANCO)
    ws.cell(r, 5, "Posição").font = f(9, True, BRANCO)
    for cc in (2, 3, 4, 5):
        ws.cell(r, cc).fill = F_HDR
        ws.cell(r, cc).alignment = Alignment(horizontal="center")
    r += 1
    for i in top:
        ws.cell(r, 2, i["id"]).font = f(9, True)
        ws.cell(r, 3, i["func"]).font = f(9)
        ws.cell(r, 4, i["cod"]).font = f(9, True)
        ws.cell(r, 5, i["rank"]).font = f(9)
        for cc in (2, 3, 4, 5):
            ws.cell(r, cc).border = BORDA
            ws.cell(r, cc).alignment = Alignment(horizontal="center" if cc != 3 else "left",
                                                 indent=1 if cc == 3 else 0)
        r += 1
    r += 1
    par("Repare no Cerco Virtual (PBI-36): é o maior custo de atraso da lista inteira e ainda assim "
        "aparece em 22º, só porque é grande. Para uma família que lida com demência, é a "
        "funcionalidade que define o produto. Se for para promover um único item contra a ordem "
        "calculada, é este.", VERM, b=True, alt=44)
    return ws


# ================================================================== 4. LACUNAS
def aba_lacunas(wb):
    ws = wb.create_sheet("Lacunas do backlog")
    titulo(ws, "LACUNAS — O QUE FALTA NO BACKLOG",
           "Dezesseis itens que não estão na lista dos 40 e que outros itens pressupõem. "
           "Seis são bloqueadores: há PBIs que não podem ser construídos sem eles.", 6)
    cab(ws, 4, ["Código", "O que falta", "Tipo", "Urgência",
                "Por que é uma lacuna", "Afeta", "Onde entraria"],
        [10, 38, 16, 11, 88, 34, 30])
    ws.freeze_panes = "A5"
    r = 5
    for cod, nome, tipo, urg, desc, afeta, onde in LACUNAS:
        for j, v in enumerate([cod, nome, tipo, urg, desc, afeta, onde]):
            c = ws.cell(r, 1 + j, v)
            c.border = BORDA; c.font = f(9)
            c.alignment = Alignment(vertical="top", wrap_text=j in (1, 4, 5),
                                    horizontal="center" if j in (0, 2, 3) else "left",
                                    indent=0 if j in (0, 2, 3) else 1)
        ws.cell(r, 1).font = f(9, True)
        ws.cell(r, 2).font = f(10, True, PRETO)
        if tipo == "Bloqueador":
            ws.cell(r, 3).fill = F_ALE; ws.cell(r, 3).font = f(9, True, VERM)
        if urg == "Crítica":
            ws.cell(r, 4).fill = F_ALE; ws.cell(r, 4).font = f(9, True, VERM)
        ws.cell(r, 5).font = f(9, False, "262626")
        ws.cell(r, 7).font = f(9, True, "1F3864")
        ws.row_dimensions[r].height = altura([(desc, 88), (afeta, 34), (onde, 30)], 46)
        r += 1
    r += 1
    c = ws.cell(r, 1, "Sugestão: transformar as seis lacunas bloqueadoras em PBIs antes de começar a "
                      "Onda 1. Cinco delas — vínculo, grupo, medicamentos, localização e push — são "
                      "infraestrutura de produto que aparece em quase toda tela; a sexta, cobrança, "
                      "é o que separa ter usuários de ter receita.")
    c.font = f(10, True, VERM)
    c.alignment = Alignment(wrap_text=True, vertical="top", indent=1)
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=7)
    ws.row_dimensions[r].height = 34
    return ws


# ================================================================== 5. OBSERVAÇÕES
def aba_obs(wb):
    ws = wb.create_sheet("Riscos e observações")
    ws.sheet_view.showGridLines = False
    titulo(ws, "RISCOS E OBSERVAÇÕES",
           "Pontos que apareceram na leitura do backlog e que valem uma decisão consciente.", 5)
    for col, w in zip("ABCD", (3, 4, 118, 3)):
        ws.column_dimensions[col].width = w
    r = 4
    for cod, tit, txt in OBS:
        c = ws.cell(r, 2, cod); c.font = f(9, True, BRANCO)
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.fill = F_SEC
        c2 = ws.cell(r, 3, tit); c2.font = f(11, True, "1F3864")
        c2.alignment = Alignment(vertical="center", indent=1)
        ws.row_dimensions[r].height = 20
        r += 1
        c = ws.cell(r, 3, txt); c.font = f(10, False, "262626")
        c.alignment = Alignment(wrap_text=True, vertical="top", indent=1)
        ws.row_dimensions[r].height = altura([(txt, 118)], 30)
        r += 2
    return ws


if __name__ == "__main__":
    wb, itens, ordem = build()
    aba_metodo(wb, itens, ordem)
    aba_lacunas(wb)
    aba_obs(wb)
    wb._sheets = [wb[n] for n in ["Priorização", "Ordem de execução", "Como priorizei",
                                  "Lacunas do backlog", "Riscos e observações"]]
    for n, c in {"Priorização": "1F3864", "Ordem de execução": "1F7040",
                 "Como priorizei": "BF8F00", "Lacunas do backlog": "C00000",
                 "Riscos e observações": "7030A0"}.items():
        wb[n].sheet_properties.tabColor = c
    wb.active = 0
    wb.save(OUT)
    print("salvo:", OUT)
