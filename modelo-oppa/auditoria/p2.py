# -*- coding: utf-8 -*-
"""Auditoria 2 — reimplementa o modelo em Python, do zero, e compara com a planilha."""
import sys, math, datetime as dt, openpyxl

def xlround(x, casas=0):
    """ROUND do Excel: metade sempre para longe do zero (Python arredonda para o par)."""
    f = 10 ** casas
    return math.floor(abs(x) * f + 0.5) / f * (1 if x >= 0 else -1)
from openpyxl.utils import get_column_letter as gcl

ARQ = sys.argv[1]
V = openpyxl.load_workbook(ARQ, data_only=True)
P, CE, PE, IN, AP, TR = V["Premissas"], V["Cenários"], V["Pessoas"], V["Investimentos"], V["Aportes"], V["Tributos"]
N = 60

def lin(ws, txt, col=1, exato=False):
    for r in range(1, ws.max_row + 1):
        v = ws.cell(r, col).value
        if isinstance(v, str) and ((v.strip() == txt) if exato else txt in v):
            return r
    raise KeyError(f"{ws.title}: '{txt}'")

def b(ws, txt):   return ws.cell(lin(ws, txt), 2).value
def anos(ws, txt): 
    r = lin(ws, txt); return [ws.cell(r, 3 + i).value for i in range(5)]

def ano_de(i):    return 2026 + (i - 1) // 12
def idx(i):       return ano_de(i) - 2026

def blocos(CE):
    """Descobre as linhas iniciais dos tres blocos de cenario na aba Cenarios."""
    m = {}
    for r in range(1, CE.max_row + 1):
        v = CE.cell(r, 1).value
        if isinstance(v, str) and v.startswith("CENÁRIO "):
            m[int(v.split()[1])] = r
    assert len(m) == 3, m
    return m

BLOCOS_CEN = blocos(CE)

# ---------------------------------------------------------------- premissas
cen        = b(P, "Cenário ativo  (1")
rampa      = b(P, "Rampa de maturação")
lanc_idx   = b(P, "Mês nº do lançamento")
tma_m      = P.cell(lin(P, "TMA equivalente mensal"), 2).value
p_ess      = b(P, "Plano Essencial (R$/mês)")
p_pre      = b(P, "Plano Premium (R$/mês)")
p_fam      = b(P, "Plano Família (R$/mês)")
reaj       = b(P, "Reajuste anual de preços")
mix        = [anos(P, "Plano Essencial", ), anos(P, "Plano Premium"), anos(P, "Plano Família")]
# as linhas de mix sao as segundas ocorrencias (as primeiras sao os precos, col B)
def anos_mix(txt):
    achados = [r for r in range(1, P.max_row+1)
               if isinstance(P.cell(r,1).value, str) and P.cell(r,1).value.strip() == txt]
    r = achados[-1]
    return [P.cell(r, 3+i).value for i in range(5)]
mix_e, mix_p, mix_f = anos_mix("Plano Essencial"), anos_mix("Plano Premium"), anos_mix("Plano Família")

def drv(rot, ordem):
    """Pega a linha do cenário `ordem` (1,2,3) dentro do bloco do driver `rot`."""
    base = lin(P, rot)
    nomes = {1: "Conservador", 2: "Provável", 3: "Agressivo"}[ordem]
    for r in range(base + 1, base + 5):
        if isinstance(P.cell(r, 1).value, str) and nomes in P.cell(r, 1).value:
            return [P.cell(r, 3 + i).value for i in range(5)]
    raise KeyError(rot)

CEN_TESTE = int(sys.argv[2]) if len(sys.argv) > 2 else cen
novos  = drv("Novos usuários captados por mês", CEN_TESTE)
churn  = drv("Churn mensal da base", CEN_TESTE)
conv   = drv("Taxa de conversão (pagantes", CEN_TESTE)
cac    = drv("CAC — custo de aquisição", CEN_TESTE)
organ  = anos(P, "Aquisição orgânica")
tx_loja  = b(P, "Taxa das lojas de aplicativos")
sh_loja  = anos(P, "% da receita de assinaturas cobrada")
tx_pgto  = b(P, "Taxa de meio de pagamento")
com_mp   = b(P, "Comissão média sobre GMV")
tick_mp  = b(P, "Ticket médio do dispositivo")
att_mp   = anos(P, "Taxa de compra mensal")
cloud_ini= b(P, "Início dos custos de nuvem")
cl_fix   = anos(P, "Nuvem — custo fixo mensal")
cl_pag   = anos(P, "Nuvem — custo variável por usuário PAGANTE")
cl_free  = anos(P, "Nuvem — custo variável por usuário GRATUITO")
sup_var  = anos(P, "Suporte — ferramentas e custo variável")
adm_ini  = b(P, "Início das despesas administrativas")
admin    = anos(P, "Administrativas (contabilidade")
mkt_rec  = anos(P, "Marketing recorrente")
mkt_ini  = b(P, "Início do marketing recorrente")
mkt_lanc = b(P, "Propaganda de lançamento")
mkt_lmes = b(P, "Mês da propaganda de lançamento")
regime   = b(P, "Regime  (1 = Simples")
lp       = b(P, "Alíquota Lucro Presumido")
teto     = b(P, "Teto do Simples Nacional")
lancto   = b(P, "Data de lançamento (calculada)")

def mes_data(i): return dt.datetime(ano_de(i), (i - 1) % 12 + 1, 1)

# ---------------------------------------------------------------- tributos
a3 = lin(TR, "ANEXO III"); a5 = lin(TR, "ANEXO V")
tab3 = [(TR.cell(a3+2+k,2).value, TR.cell(a3+2+k,4).value, TR.cell(a3+2+k,5).value) for k in range(6)]
tab5 = [(TR.cell(a5+2+k,2).value, TR.cell(a5+2+k,4).value, TR.cell(a5+2+k,5).value) for k in range(6)]
tab  = tab5 if regime == 2 else tab3

def aliq(rbt12):
    if regime == 3 or rbt12 > teto:
        return lp
    if rbt12 == 0:
        return tab[0][1]
    nom, ded = tab[0][1], tab[0][2]
    for de, n_, d_ in tab:
        if rbt12 >= de:
            nom, ded = n_, d_
    return max(tab[0][1], (rbt12 * nom - ded) / rbt12)

# ---------------------------------------------------------------- pessoas
pi = lin(PE, "Função / paphel") if False else None
hdr = lin(PE, "Função / papel")
pes_ini = hdr + 1
pes_fim = pes_ini
while PE.cell(pes_fim + 1, 7).value is not None and isinstance(PE.cell(pes_fim+1,7).value,(int,float)):
    pes_fim += 1
plano_on = b(P, "Considerar o plano de contratações")
def pessoal(i, classe):
    d = mes_data(i); tot = 0.0
    for r in range(pes_ini, pes_fim + 1):
        st, ini, fim, cls = PE.cell(r,3).value, PE.cell(r,8).value, PE.cell(r,9).value, PE.cell(r,10).value
        custo = PE.cell(r,7).value or 0
        if cls != classe or ini is None or fim is None: continue
        ativo = (st == "Contratado") + (st == "Plano") * plano_on
        if ini <= d <= fim:
            tot += ativo * custo
    return tot

# ---------------------------------------------------------------- capex
ci = lin(IN, "Item", 2) + 1
cf = ci
while IN.cell(cf + 1, 4).value is not None and isinstance(IN.cell(cf+1,4).value,(int,float)):
    cf += 1
def capex(i):
    d = mes_data(i)
    ult = dt.datetime(d.year + (d.month == 12), d.month % 12 + 1, 1)
    tot = sum(IN.cell(r,4).value for r in range(ci, cf+1)
              if IN.cell(r,1).value and d <= IN.cell(r,1).value < ult)
    return tot + pessoal(i, "Investimento")

# ---------------------------------------------------------------- aportes
ai = lin(AP, "Sócio / Investidor", 2) + 1
af = ai
while AP.cell(af + 1, 3).value is not None: af += 1
ei = lin(AP, "Investidor", 1) + 1
ef = ei
while AP.cell(ef + 1, 3).value is not None and AP.cell(ef+1,2).value: ef += 1
inc_prev = b(P, "Considerar aportes PREVISTOS")
inc_neg  = b(P, "Considerar acordos EM NEGOCIAÇÃO")
def gate_st(st):
    return {"Realizado":1,"Compromisso":1}.get(st, inc_prev if st=="Previsto" else (inc_neg if st=="Em negociação" else 0))
def aportes(i):
    d = mes_data(i); ult = dt.datetime(d.year + (d.month==12), d.month%12+1, 1)
    tot = sum(gate_st(AP.cell(r,4).value) * (AP.cell(r,3).value or 0)
              for r in range(ai, af+1)
              if AP.cell(r,1).value and d <= AP.cell(r,1).value < ult)
    for r in range(ei, ef+1):
        st, mes, val = AP.cell(r,2).value, AP.cell(r,4).value, AP.cell(r,3).value or 0
        g = (st=="Confirmado") + (st=="Em negociação")*inc_neg
        if mes and d <= mes < ult: tot += g*val
    return tot
def b2b(i):
    d = mes_data(i); tot = 0.0
    for r in range(ei, ef+1):
        st, rec, i_r, f_r = AP.cell(r,2).value, AP.cell(r,5).value or 0, AP.cell(r,6).value, AP.cell(r,7).value
        g = (st=="Confirmado") + (st=="Em negociação")*inc_neg
        if i_r and f_r and i_r <= d <= f_r: tot += g*rec
    return tot

# ---------------------------------------------------------------- modelo
res = {k: [0.0]*(N+1) for k in
       ["novos","base","pag","rec_ass","rec_mp","rec_b2b","bruta","rbt12","aliq","imp",
        "loja","pgto","liq","nuvem","sup","pes","mkt","adm","ebitda","capex","fcl",
        "fcld","fcld_ac","fcl_ac","aporte","caixa"]}
base_ant = 0.0
for i in range(1, N+1):
    a, d = idx(i), mes_data(i)
    nv = 0 if d < lancto else novos[a]
    ch = base_ant * churn[a]
    base = base_ant - ch + nv
    cv = 0 if d < lancto else conv[a] * min(1, max(0, (i - lanc_idx + 1) / rampa))
    pag = xlround(base * cv)
    pe_ = xlround(pag * mix_e[a]); pp = xlround(pag * mix_p[a]); pf = pag - pe_ - pp
    infl = (1 + reaj) ** a
    ass = (pe_*p_ess + pp*p_pre + pf*p_fam) * infl
    mp  = base * att_mp[a] * tick_mp * com_mp
    bb  = b2b(i)
    bruta = ass + mp + bb
    res["novos"][i], res["base"][i], res["pag"][i] = nv, base, pag
    res["rec_ass"][i], res["rec_mp"][i], res["rec_b2b"][i], res["bruta"][i] = ass, mp, bb, bruta
    if i == 1:   rbt = bruta*12
    elif i <= 12: rbt = sum(res["bruta"][1:i])/(i-1)*12
    else:        rbt = sum(res["bruta"][i-12:i])
    al = aliq(rbt)
    imp  = -bruta*al
    loja = -ass*tx_loja*sh_loja[a]
    pgto = -(ass*(1-sh_loja[a]) + mp)*tx_pgto
    liq  = bruta + imp + loja + pgto
    nuv  = -(pag*cl_pag[a] + (base-pag)*cl_free[a] + (cl_fix[a] if d >= cloud_ini else 0))
    sup  = -pag*sup_var[a]
    pes  = -pessoal(i, "Despesa")
    mkt  = -(nv*(1-organ[a])*cac[a] + (mkt_rec[a] if d >= mkt_ini else 0)
             + (mkt_lanc if d == dt.datetime(mkt_lmes.year, mkt_lmes.month, 1) else 0))
    adm  = -(admin[a] if d >= adm_ini else 0)
    eb   = liq + nuv + sup + pes + mkt + adm
    cpx  = -capex(i)
    fcl  = eb + cpx
    for k, v in dict(rbt12=rbt, aliq=al, imp=imp, loja=loja, pgto=pgto, liq=liq, nuvem=nuv,
                     sup=sup, pes=pes, mkt=mkt, adm=adm, ebitda=eb, capex=cpx, fcl=fcl).items():
        res[k][i] = v
    res["fcld"][i]    = fcl/(1+tma_m)**i
    res["fcld_ac"][i] = res["fcld_ac"][i-1] + res["fcld"][i]
    res["fcl_ac"][i]  = res["fcl_ac"][i-1] + fcl
    res["aporte"][i]  = aportes(i)
    res["caixa"][i]   = res["caixa"][i-1] + fcl + res["aporte"][i]
    base_ant = base

# ---------------------------------------------------------------- comparação
S = BLOCOS_CEN[CEN_TESTE]
MAPA = [("novos",1),("base",4),("pag",6),("rec_ass",13),("rec_mp",14),("rec_b2b",15),
        ("bruta",16),("rbt12",17),("aliq",18),("imp",19),("loja",20),("pgto",21),("liq",22),
        ("nuvem",23),("sup",24),("pes",25),("mkt",26),("adm",27),("ebitda",28),("capex",29),
        ("fcl",30),("fcld",31),("fcld_ac",32),("fcl_ac",33),("aporte",35),("caixa",37)]
nome_cen = {1:"CONSERVADOR",2:"PROVÁVEL",3:"AGRESSIVO"}[CEN_TESTE]
print(f"CENÁRIO {CEN_TESTE} — {nome_cen}")
print(f"{'linha do modelo':<22}{'maior desvio':>16}{'mês':>6}   veredito")
print("-"*74)
falhas = 0
for chave, off in MAPA:
    pior, mes = 0.0, 0
    for i in range(1, N+1):
        xl = CE.cell(S+off, 1+i).value or 0
        py = res[chave][i]
        dif = abs(xl - py)
        tol = max(0.02, abs(py)*1e-9)
        if dif > tol and dif > pior:
            pior, mes = dif, i
    ok = pior <= 0.02
    if not ok: falhas += 1
    print(f"{chave:<22}{pior:>16,.4f}{mes:>6}   {'ok' if ok else '<<< DIVERGE'}")
print("-"*74)
print(f"{'LINHAS DIVERGENTES:':<22}{falhas:>16}")
