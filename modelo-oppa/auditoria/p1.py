# -*- coding: utf-8 -*-
"""Auditoria 1 — mecanica: intervalos, totais, referencias, formatos."""
import re, sys, openpyxl
from openpyxl.utils import get_column_letter as gcl, column_index_from_string as cidx

ARQ = sys.argv[1]
F = openpyxl.load_workbook(ARQ)                 # formulas
V = openpyxl.load_workbook(ARQ, data_only=True) # valores
ach = []
def flag(sev, aba, cel, msg):
    ach.append((sev, aba, cel, msg))

REF = re.compile(r"(?:'([^']+)'|(\b[A-Za-zÀ-ÿ][A-Za-zÀ-ÿ0-9 ]*))!\$?([A-Z]{1,2})\$?(\d+)")
RANGE = re.compile(r"\$?([A-Z]{1,2})\$?(\d+):\$?([A-Z]{1,2})\$?(\d+)")

# ---------------------------------------------------------------- A) referências
print("A) Referências entre abas apontando para célula vazia")
n = 0
for ws in F.worksheets:
    for row in ws.iter_rows():
        for c in row:
            if not (isinstance(c.value, str) and c.value.startswith("=")):
                continue
            for m in REF.finditer(c.value):
                alvo = m.group(1) or m.group(2)
                if alvo not in F.sheetnames:
                    continue
                col, lin = m.group(3), int(m.group(4))
                tgt = F[alvo][f"{col}{lin}"]
                if tgt.value is None:
                    flag("ALTA", ws.title, c.coordinate,
                         f"referencia {alvo}!{col}{lin}, que está vazia")
                    n += 1
print(f"   {n} ocorrência(s)")

# ---------------------------------------------------------------- B) totais BJ
print("B) Coluna TOTAL (BJ) conferindo com a soma de B:BI")
n = 0
for ws in F.worksheets:
    wv = V[ws.title]
    for r in range(1, ws.max_row + 1):
        cel = ws.cell(r, 62)
        if not (isinstance(cel.value, str) and cel.value.startswith("=SUM(B")):
            continue
        soma = sum(x for x in (wv.cell(r, c).value for c in range(2, 62))
                   if isinstance(x, (int, float)))
        tot = wv.cell(r, 62).value
        if isinstance(tot, (int, float)) and abs(tot - soma) > 0.01:
            flag("ALTA", ws.title, f"BJ{r}",
                 f"total {tot:,.2f} ≠ soma dos meses {soma:,.2f} (dif {tot-soma:,.2f})")
            n += 1
print(f"   {n} divergência(s)")

# ---------------------------------------------------------------- C) linhas de total
print("C) Linhas rotuladas como TOTAL conferindo com suas parcelas")
n = 0
for ws in F.worksheets:
    wv = V[ws.title]
    for r in range(1, ws.max_row + 1):
        rot = ws.cell(r, 1).value
        if not (isinstance(rot, str) and re.search(r"^\(?=?\)?\s*TOTAL|^TOTAL", rot.strip(), re.I)):
            continue
        f0 = ws.cell(r, 2).value
        m = RANGE.search(str(f0) or "")
        if not m or not str(f0).startswith("=SUM("):
            continue
        r1, r2 = int(m.group(2)), int(m.group(4))
        for c in range(2, 62):
            tot = wv.cell(r, c).value
            parc = sum(x for x in (wv.cell(rr, c).value for rr in range(r1, r2 + 1))
                       if isinstance(x, (int, float)))
            if isinstance(tot, (int, float)) and abs(tot - parc) > 0.01:
                flag("ALTA", ws.title, f"{gcl(c)}{r}",
                     f"'{rot[:34]}' = {tot:,.2f} mas as parcelas somam {parc:,.2f}")
                n += 1
                break
print(f"   {n} divergência(s)")

# ---------------------------------------------------------------- D) hardcodes
print("D) Números fixos no meio de linhas de fórmula (possível sobrescrita)")
n = 0
for ws in F.worksheets:
    for r in range(1, ws.max_row + 1):
        tipos = []
        for c in range(2, 62):
            v = ws.cell(r, c).value
            if v is None: tipos.append(None)
            elif isinstance(v, str) and v.startswith("="): tipos.append("f")
            elif isinstance(v, (int, float)): tipos.append("n")
            else: tipos.append("t")
        forms = tipos.count("f")
        nums = [i for i, t in enumerate(tipos) if t == "n"]
        if forms >= 40 and nums:
            for i in nums:
                flag("MÉDIA", ws.title, f"{gcl(2+i)}{r}",
                     f"valor fixo numa linha com {forms} fórmulas — '{str(ws.cell(r,1).value)[:34]}'")
                n += 1
print(f"   {n} ocorrência(s)")

# ---------------------------------------------------------------- E) percentuais
print("E) Células formatadas como % com valor fora da faixa plausível")
n = 0
for ws in V.worksheets:
    for row in ws.iter_rows():
        for c in row:
            if c.number_format and "%" in c.number_format and isinstance(c.value, (int, float)):
                if abs(c.value) > 5:
                    flag("MÉDIA", ws.title, c.coordinate,
                         f"formato % com valor {c.value:,.2f} (=" f"{c.value*100:,.0f}%)")
                    n += 1
print(f"   {n} ocorrência(s)")

# ---------------------------------------------------------------- F) sobreposição
print("F) Somas com intervalo que inclui a própria linha (referência circular velada)")
n = 0
for ws in F.worksheets:
    for row in ws.iter_rows():
        for c in row:
            if not (isinstance(c.value, str) and c.value.startswith("=")):
                continue
            for m in RANGE.finditer(c.value):
                if "!" in c.value[:m.start()].split("(")[-1]:
                    continue
                r1, r2 = int(m.group(2)), int(m.group(4))
                c1, c2 = cidx(m.group(1)), cidx(m.group(3))
                if r1 <= c.row <= r2 and c1 <= c.column <= c2:
                    flag("ALTA", ws.title, c.coordinate, f"intervalo {m.group(0)} contém a própria célula")
                    n += 1
print(f"   {n} ocorrência(s)")

# ---------------------------------------------------------------- G) vazios
print("G) Fórmulas cujo resultado ficou vazio ou texto inesperado")
n = 0
for ws in F.worksheets:
    wv = V[ws.title]
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith("="):
                r = wv[c.coordinate].value
                if r is None:
                    flag("MÉDIA", ws.title, c.coordinate, "fórmula sem valor calculado")
                    n += 1
print(f"   {n} ocorrência(s)")

print("\n" + "=" * 96)
if not ach:
    print("AUDITORIA 1: nenhum achado.")
else:
    print(f"AUDITORIA 1 — {len(ach)} achado(s):\n")
    for sev, aba, cel, msg in sorted(ach, key=lambda x: (x[0] != "ALTA", x[1], x[2]))[:60]:
        print(f"  [{sev:<5}] {aba:<16} {cel:<8} {msg}")
    if len(ach) > 60:
        print(f"  ... e mais {len(ach)-60}")
