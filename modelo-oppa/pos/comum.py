# -*- coding: utf-8 -*-
"""Utilitarios para PATCHAR uma planilha existente, sem regerar do zero.

O arquivo do usuario tem edicoes dele — aportes, projecoes, precos. Regerar
pelo gerador apagaria tudo isso. Entao estes scripts abrem o arquivo dele,
localizam as linhas PELO ROTULO e acrescentam o que falta.

Nunca localize linha por numero aqui: o arquivo que chega pode ter sido mexido.
"""
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as gcl

FONT = "Arial"
BRL   = '"R$ "#,##0;("R$ "#,##0);"–"'
BRL2  = '"R$ "#,##0.00;("R$ "#,##0.00);"–"'
INT   = '#,##0;(#,##0);"–"'
PCT   = '0.0%;(0.0%);"–"'
PCT2  = '0.00%;(0.00%);"–"'
MULT  = '0.0"x";(0.0"x");"–"'
MES   = 'mmm/yy'
TXT   = '@'

AZUL_INPUT = "0000FF"
PRETO      = "000000"
VERDE_LINK = "008000"
CINZA      = "595959"
BRANCO     = "FFFFFF"
VERMELHO   = "C00000"

FILL_TITULO = PatternFill("solid", fgColor="1F3864")
FILL_SECAO  = PatternFill("solid", fgColor="2E5C8A")
FILL_SUB    = PatternFill("solid", fgColor="D9E2F3")
FILL_TOTAL  = PatternFill("solid", fgColor="BDD7EE")
FILL_INPUT  = PatternFill("solid", fgColor="FFFF00")
FILL_OK     = PatternFill("solid", fgColor="E2EFDA")
FILL_ALERTA = PatternFill("solid", fgColor="FCE4D6")
FILL_HEADER = PatternFill("solid", fgColor="404040")

THIN = Side(style="thin", color="BFBFBF")
BORDA = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

N_MESES = 60
COL0 = 2
L_FIM = gcl(COL0 + N_MESES - 1)      # BI


def f(size=10, bold=False, color=PRETO, italic=False):
    return Font(name=FONT, size=size, bold=bold, color=color, italic=italic)


def linha_por_rotulo(ws, texto, col=1, inicio=1, fim=None, exato=False):
    """Primeira linha cujo rotulo casa. Lanca erro se nao achar — melhor
    quebrar aqui do que escrever formula apontando para o lugar errado."""
    fim = fim or ws.max_row
    for r in range(inicio, fim + 1):
        v = ws.cell(r, col).value
        if isinstance(v, str):
            v = v.strip()
            if (v == texto) if exato else (texto in v):
                return r
    raise KeyError(f"{ws.title}: rótulo não encontrado — {texto!r}")


def blocos_cenario(ws):
    """Linhas iniciais dos tres blocos da aba Cenarios, em ordem."""
    out = []
    for r in range(1, ws.max_row + 1):
        v = ws.cell(r, 1).value
        if isinstance(v, str) and v.strip().startswith("CENÁRIO "):
            out.append(r)
    assert len(out) == 3, f"esperava 3 blocos de cenário, achei {len(out)}"
    return out


def linhas_do_motor(wb, rotulos):
    """{rotulo: [linha no bloco 1, bloco 2, bloco 3]} — por rotulo, nao por offset."""
    CE = wb["Cenários"]
    B = blocos_cenario(CE)
    out = {}
    for rot in rotulos:
        linhas = []
        for k, S in enumerate(B):
            fim = B[k + 1] - 1 if k + 1 < len(B) else CE.max_row
            linhas.append(linha_por_rotulo(CE, rot, inicio=S + 1, fim=fim))
        out[rot] = linhas
    return out


def titulo(ws, texto, subtitulo=None, largura=20):
    ws["A1"] = texto
    ws["A1"].font = f(14, True, BRANCO)
    ws["A1"].alignment = Alignment(vertical="center", indent=1)
    for c in range(1, largura + 1):
        ws.cell(1, c).fill = FILL_TITULO
    ws.row_dimensions[1].height = 26
    if subtitulo:
        ws["A2"] = subtitulo
        ws["A2"].font = f(9, False, CINZA, italic=True)
        ws["A2"].alignment = Alignment(vertical="center", indent=1)
        ws.row_dimensions[2].height = 16


def secao(ws, linha, texto, largura=20):
    ws.cell(linha, 1, texto).font = f(10, True, BRANCO)
    for c in range(1, largura + 1):
        ws.cell(linha, c).fill = FILL_SECAO
    ws.row_dimensions[linha].height = 18


def nota(ws, linha, texto, col=7, ate=14):
    c = ws.cell(linha, col, texto)
    c.font = f(8, False, CINZA, italic=True)
    c.alignment = Alignment(wrap_text=True, vertical="center", indent=1)
    ws.merge_cells(start_row=linha, start_column=col, end_row=linha, end_column=ate)
    alt = 13 + 10 * (len(texto) // 78)
    atual = ws.row_dimensions[linha].height or 0
    ws.row_dimensions[linha].height = max(atual, alt, 15)
